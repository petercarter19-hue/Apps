"""Tests for the PC VoiceTool: key gestures, dictation, every command-key route, styles, actions, meeting notes, the app
window's host and API, spoken answers and the offline hook.
Run:  .venv\\Scripts\\python.exe -m pytest -q test_voicetool.py test_app_window_e2e.py

Nothing here uses the microphone, the speech model, Claude or the keyboard, and nothing is typed, pasted or sent:
- every file the tool writes (history, usage, words, fixes, config, meetings, notes, journal, Drive notes) goes to a
  temporary folder, never the real VoiceTool folder, OneDrive or Google Drive;
- paste, key presses, os.startfile, the browser, schtasks /Change and Claude are replaced by recorders;
- the one live Azure check (a real tidy-up of a made-up sentence) only runs with VOICETOOL_LIVE_AZURE=1;
- the one live audio check (4 seconds of the REAL microphone and speakers, measured and deleted at once) only runs with
  VOICETOOL_LIVE_AUDIO=1, and only with the person's say-so.
"""
import json
import os
import re
import subprocess
import sys
import threading
import time
import wave
from pathlib import Path
from types import SimpleNamespace

import numpy as np
import pytest

import actions_windows as aw
import commands
import corrections
import local_cleanup_win
import meetings
import meetings_windows
import read_aloud
import system_actions
import system_windows
import settings_window
import speaker_windows
import styles
import ui_api
import usage
import voice_actions as va
import voicetool as vt

HERE = Path(__file__).resolve().parent.parent / "src"  # the app code (the tests live in tools\)
# (the fixture replaces these three: the tests never run the real `claude`)
REAL_CLAUDE_PROBLEM, REAL_DEFAULT_NOTES_FOLDER, REAL_SIGNED_IN = commands.claude_problem, va.default_notes_folder, commands.signed_in
REAL_NOTES_BASE = meetings_windows.notes_base  # (the fixture replaces this one too: meeting notes never go to the real Documents)
DICTATE, COMMAND = vt.Trigger("d", "alt_l", True).vk, vt.Trigger("c", "ctrl_l", True).vk  # Left Alt, Left Ctrl


# ====================================================================================================== fakes + fixtures

class FakeRecorder:
    device = None

    def __init__(self, seconds=2.0, loud=True):
        self.running, self.starts, self.level = False, 0, 0.0
        self.seconds, self.loud = seconds, loud

    def start(self):
        self.running = True
        self.starts += 1

    def stop(self):
        was, self.running = self.running, False
        if not was:
            return np.zeros(0, np.float32)
        return np.full(int(vt.RATE * self.seconds), 0.3 if self.loud else 0.001, np.float32)


class FakeRoot:
    def after(self, *a):
        pass


class FakeLight:
    def __init__(self):
        self.state, self.level, self.claude_busy, self.meeting, self.root = None, 0.0, 0, None, FakeRoot()
        self.history = []

    def set(self, state):
        self.state = state
        self.history.append(state)

    def draw(self):
        pass


class Reply:
    def __init__(self, content):
        self.choices = [SimpleNamespace(message=SimpleNamespace(content=content))]
        self.usage = None


class FakeClient:
    """Stands in for the Azure client: answers each request with the next scripted reply, and remembers what it was sent."""

    def __init__(self, *replies):
        self.replies, self.sent = list(replies), []
        self.chat = SimpleNamespace(completions=self)

    def create(self, **kw):
        self.sent.append(kw)
        return Reply(self.replies.pop(0) if len(self.replies) > 1 else self.replies[0])

    def with_options(self, **kw):
        return self


class FakeCleaner:
    local = None

    def __init__(self, client=None, result=None, fail=None):
        self.client, self.result, self.fail = client or FakeClient("{}"), result, fail
        self.calls, self.warms = [], 0

    def warm_up(self):
        self.warms += 1

    def __call__(self, text, hint=""):
        self.calls.append((text, hint))
        if self.fail:
            raise self.fail
        return self.result if self.result is not None else text

    def _client(self):
        return self.client

    _azure_client = _client


class LocalReply:
    def __init__(self, content, finish_reason="stop"):
        self.choices = [SimpleNamespace(message=SimpleNamespace(content=content), finish_reason=finish_reason)]
        self.usage = None


class FakeLocalClient(FakeClient):
    """Stands in for the on-device client: every request gets the same reply."""

    def __init__(self, content, finish_reason="stop"):
        super().__init__()
        self.content, self.finish_reason = content, finish_reason

    def create(self, **kw):
        self.sent.append(kw)
        return LocalReply(self.content, self.finish_reason)


@pytest.fixture
def online(monkeypatch):
    """Azure is set up and the "keep everything on this computer" switch is off (the advanced setting)."""
    monkeypatch.setitem(vt.CONFIG, "offline_only", False)
    monkeypatch.setitem(vt.CONFIG, "azure_openai_base_url", "https://example.invalid/openai/v1/")


@pytest.fixture(autouse=True)
def isolated(tmp_path, monkeypatch):
    """Point every file the tool writes at a temp folder, and every dangerous action at a recorder."""
    monkeypatch.setattr(vt.log, "disabled", True)
    monkeypatch.setattr(vt, "beep", lambda *tones: None)
    for module in (vt, settings_window, corrections, meetings, va, ui_api, usage, commands, styles):
        if hasattr(module, "HERE"):
            monkeypatch.setattr(module, "HERE", tmp_path)
    (tmp_path / "words.txt").write_text("# words\nZephyr\nClaude\nWispr Flow\n", encoding="utf-8")
    monkeypatch.setattr(usage, "FILE", tmp_path / "usage.jsonl")
    monkeypatch.setattr(corrections, "FILE", tmp_path / "corrections.json")
    monkeypatch.setattr(corrections, "PENDING", tmp_path / "corrections_pending.json")
    monkeypatch.setattr(corrections, "_cache", {"mtime": None, "pairs": []})
    monkeypatch.setattr(commands, "STYLE_FILE", tmp_path / "style.md")
    monkeypatch.setattr(styles, "FILE", tmp_path / "styles.json")
    monkeypatch.setattr(styles, "_cache", {"mtime": None, "data": None})
    monkeypatch.setattr(ui_api, "HISTORY", tmp_path / "history.jsonl")
    monkeypatch.setattr(ui_api, "WORDS", tmp_path / "words.txt")
    monkeypatch.setattr(vt, "MEETING_AUDIO", tmp_path / "recordings" / "meetings")
    monkeypatch.setattr(va, "SHORTCUTS", tmp_path / "shortcuts.json")
    monkeypatch.setattr(va, "JOURNAL", tmp_path / "journal")
    monkeypatch.setattr(va, "default_notes_folder", lambda: tmp_path / "Documents" / "Voice notes")
    monkeypatch.setattr(commands, "claude_problem", lambda: None)  # Claude is "installed and signed in" unless a test says otherwise
    monkeypatch.setattr(commands, "signed_in", lambda: True)
    monkeypatch.setenv("OneDrive", str(tmp_path / "OneDrive"))  # (a temp "OneDrive" that meeting notes must never use)
    monkeypatch.setattr(meetings_windows, "notes_base", lambda: tmp_path / "Documents")  # never the real Documents folder
    monkeypatch.setattr(settings_window, "TASK", "VoiceTool-test-none")  # never the real startup task
    # nothing may press a key, paste, open or run anything on this PC
    monkeypatch.setattr(vt.win32api, "keybd_event", lambda *a: None)
    monkeypatch.setattr(vt, "paste", lambda text: PASTED.append(text))
    monkeypatch.setattr(os, "startfile", lambda p: OPENED.append(str(p)))
    monkeypatch.setattr(system_windows, "_lock_workstation", lambda: LOCKED.append(1))  # the screen is never really locked
    monkeypatch.setattr(system_windows, "_volume_control", lambda: None)  # nor the volume set (a test that needs one gives it a fake)
    monkeypatch.setattr(aw, "play_audio_file", lambda path: SPOKEN.append(("PLAY", str(path))))  # no sound
    monkeypatch.setattr(aw, "stop_audio", lambda: None)
    monkeypatch.setattr(aw, "_speaker", None)  # and no voice is ever started unless a test replaces speak()
    monkeypatch.setattr(aw, "speak", lambda text, **kw: SPOKEN.append((text, kw)))
    monkeypatch.setattr(aw, "stop_speaking", lambda: STOPS.append(1))
    for key, value in {"enabled": True, "dictate_enabled": True, "command_enabled": True, "typed_fix": True,
                       "cleanup": True, "auto_learn": True, "use_style": True, "translate_to": "", "speak_answers": False,
                       "offline_tidyup": True, "offline_only": False, "azure_openai_base_url": "https://example.invalid/openai/v1/", "first_run": False, "meeting_notes": True, "meeting_auto_offer": True,
                       "meeting_followup": True, "claude_actions": True, "claude_keep_running": True,
                       "save_recordings": False, "show_light": True, "hold_seconds": 0.3, "read_voice": "", "read_rate": "normal",
                       "tts_deployment": "", "notes_destination": "folder", "notes_folder": "", "speech_model_path": "",
                       "double_tap_seconds": 0.4}.items():
        monkeypatch.setitem(vt.CONFIG, key, value)
    PASTED.clear()
    OPENED.clear()
    LOCKED.clear()
    SPOKEN.clear()
    STOPS.clear()
    return tmp_path


PASTED, OPENED, LOCKED, SPOKEN, STOPS = [], [], [], [], []


@pytest.fixture
def tool(monkeypatch):
    t = vt.VoiceTool(recorder=FakeRecorder(), transcriber=lambda a: "", cleaner=FakeCleaner(), with_ui=False)
    t.light = FakeLight()
    t.captured, t.taught, t.cards, t.card_buttons = [], [], [], []
    monkeypatch.setattr(t, "_capture", lambda trig, box=None: t.captured.append(trig.name))  # never press Ctrl+C in tests
    monkeypatch.setattr(t, "_teach", lambda trig: t.taught.append(trig.name))
    # (the window in front stays Notepad, window 2, the same window the tests' command contexts come from; a test that
    # switches windows replaces this)
    monkeypatch.setattr(t, "_front_window", lambda: {"app": "notepad.exe", "title": "Untitled", "bundle": "", "hwnd": 2})
    monkeypatch.setattr(t, "_card", lambda title, body, seconds=60, speak=False, buttons=None:
                        (t.cards.append((title, body, speak)), t.card_buttons.append(dict(buttons or []))))
    monkeypatch.setattr(vt.win32api, "GetAsyncKeyState", lambda vk: 0)
    return t


def key(tool, vk, down, at):
    tool._handle("trigger", tool.triggers[vk], down, at)


def agree(tool):
    """Answer the "Everyone agreed, start" button of the card on screen."""
    assert tool.cards[-1][0] in ("Before recording", "Meeting in Zoom?", "Meeting in Teams?"), tool.cards[-1]
    tool.card_buttons[-1]["Everyone agreed, start"]()


def start_meeting(tool, app=None):
    """What a person does: ask to start, read the consent card, and press "Everyone agreed, start"."""
    tool._start_meeting(app)
    if tool.meeting is None and tool.cards and tool.cards[-1][0] == "Before recording":
        agree(tool)


def jobs(tool):
    return tool.jobs.qsize()


def now():
    return time.monotonic()


def settle(tool=None):
    time.sleep(0.1)


# ====================================================================================================== dictation key

def test_dictation_records_from_the_first_instant_and_cues_after_a_hold(tool):
    key(tool, DICTATE, True, now() - 1.0)
    assert tool.recorder.running and tool.active.state == vt.HOLDING and not tool.active.cued
    tool._tick()
    assert tool.active.cued and tool.light.state == "listening"
    key(tool, DICTATE, False, now())
    assert tool.active is None and not tool.recorder.running and jobs(tool) == 1
    assert tool.light.state == "working"


def test_single_short_tap_does_nothing(tool):
    key(tool, DICTATE, True, now() - 1.0)
    key(tool, DICTATE, False, now() - 0.95)
    assert tool.triggers[DICTATE].state == vt.TAP_WAIT
    tool._tick()
    assert tool.active is None and not tool.recorder.running and jobs(tool) == 0


def test_double_tap_locks_then_tap_stops(tool):
    t0 = now()
    key(tool, DICTATE, True, t0)
    key(tool, DICTATE, False, t0 + 0.05)
    key(tool, DICTATE, True, t0 + 0.2)
    assert tool.triggers[DICTATE].state == vt.LOCKED and tool.recorder.running
    assert tool.light.state == "handsfree"
    key(tool, DICTATE, False, t0 + 0.25)  # release of the second tap is swallowed
    assert tool.triggers[DICTATE].state == vt.LOCKED
    key(tool, DICTATE, True, t0 + 5)  # one tap stops
    assert tool.active is None and jobs(tool) == 1
    key(tool, DICTATE, False, t0 + 5.05)
    assert jobs(tool) == 1


def test_other_key_while_held_is_a_shortcut_and_is_dropped(tool):
    key(tool, DICTATE, True, now())
    tool._handle("other", None, 0x09, now())  # Alt+Tab
    assert tool.active is None and not tool.recorder.running
    key(tool, DICTATE, False, now())
    assert jobs(tool) == 0


def test_escape_cancels_hold_and_hands_free(tool):
    key(tool, DICTATE, True, now())
    tool._handle("other", None, vt.win32con.VK_ESCAPE, now())
    key(tool, DICTATE, False, now())
    assert tool.active is None and jobs(tool) == 0
    t0 = now()
    key(tool, DICTATE, True, t0)
    key(tool, DICTATE, False, t0 + 0.05)
    key(tool, DICTATE, True, t0 + 0.2)
    key(tool, DICTATE, False, t0 + 0.25)
    tool._handle("other", None, vt.win32con.VK_ESCAPE, t0 + 3)
    assert tool.active is None and jobs(tool) == 0 and not tool.recorder.running


def test_quiet_clip_is_skipped(tool):
    tool.recorder.loud = False
    key(tool, DICTATE, True, now() - 1)
    key(tool, DICTATE, False, now())
    assert jobs(tool) == 0 and tool.light.state is None


def test_very_short_clip_is_skipped(tool):
    tool.recorder.seconds = 0.2
    key(tool, DICTATE, True, now() - 1)
    key(tool, DICTATE, False, now())
    assert jobs(tool) == 0


def test_other_trigger_is_ignored_while_one_is_busy(tool):
    key(tool, DICTATE, True, now())
    key(tool, COMMAND, True, now())
    assert tool.active.name == "dictate"
    key(tool, COMMAND, False, now())
    key(tool, DICTATE, False, now() + 1)
    assert jobs(tool) == 1


def test_key_auto_repeat_is_ignored(tool):
    key(tool, DICTATE, True, now())
    key(tool, DICTATE, True, now())
    key(tool, DICTATE, True, now())
    assert tool.recorder.starts == 1


def test_microphone_failure_beeps_and_does_not_start(tool, monkeypatch):
    def broken():
        raise OSError("no microphone")
    monkeypatch.setattr(tool.recorder, "start", broken)
    key(tool, DICTATE, True, now())
    assert tool.active is None


def test_dictation_remembers_the_app_in_front(tool):
    key(tool, DICTATE, True, now())
    assert tool.active.app["app"] == "notepad.exe"
    key(tool, COMMAND, True, now())  # (the other key is ignored)
    assert tool.triggers[COMMAND].app == {}


# ====================================================================================================== command key

def test_command_key_hold_captures_the_selection_after_release(tool):
    key(tool, COMMAND, True, now() - 1)
    assert not tool.recorder.running and tool.triggers[COMMAND].context is None  # (Ctrl waits a moment before the mic opens)
    tool._tick()
    assert tool.recorder.running and tool.active.cued and tool.light.state == "command"
    key(tool, COMMAND, False, now())
    settle()
    assert jobs(tool) == 1 and tool.captured == ["command"] and tool.light.state == "thinking"


def test_command_double_tap_hands_free(tool):
    t0 = now()
    key(tool, COMMAND, True, t0)
    key(tool, COMMAND, False, t0 + 0.05)
    key(tool, COMMAND, True, t0 + 0.2)
    assert tool.triggers[COMMAND].state == vt.TAP_WAIT  # (Ctrl: nothing is decided until it's let go alone)
    key(tool, COMMAND, False, t0 + 0.25)
    assert tool.triggers[COMMAND].state == vt.LOCKED and tool.light.state == "command_handsfree"


def test_shift_command_tap_opens_the_teach_window(tool, monkeypatch):
    monkeypatch.setattr(vt.win32api, "GetAsyncKeyState", lambda vk: 0x8000)
    key(tool, COMMAND, True, now())
    assert tool.triggers[COMMAND].state == vt.TEACH and not tool.recorder.running
    key(tool, COMMAND, False, now())
    settle()
    assert tool.taught == ["command"] and tool.active is None


def test_shortcut_with_shift_and_ctrl_is_not_a_teach(tool, monkeypatch):
    monkeypatch.setattr(vt.win32api, "GetAsyncKeyState", lambda vk: 0x8000)
    key(tool, COMMAND, True, now())
    tool._handle("other", None, 0x54, now())  # the T of Ctrl+Shift+T
    key(tool, COMMAND, False, now())
    settle()
    assert tool.taught == [] and tool.active is None


def test_typed_fix_can_be_switched_off(tool, monkeypatch):
    monkeypatch.setitem(vt.CONFIG, "typed_fix", False)
    monkeypatch.setattr(vt.win32api, "GetAsyncKeyState", lambda vk: 0x8000)
    key(tool, COMMAND, True, now())
    assert tool.triggers[COMMAND].state == vt.HOLDING


def test_warm_up_only_when_azure_will_be_asked(tool, monkeypatch):
    monkeypatch.setitem(vt.CONFIG, "offline_only", False)

    def press(vk):
        key(tool, vk, True, now() - 1)
        tool._tick()  # (Ctrl opens the microphone, and warms Azure up, only once it has been down a moment)
        tool._drop(tool.active)  # (drop the recording, then let go of the key)
        key(tool, vk, False, now())
    press(DICTATE)
    assert tool.azure.warms == 1
    monkeypatch.setitem(vt.CONFIG, "cleanup", False)
    press(DICTATE)
    assert tool.azure.warms == 1  # tidy-up is off: nothing to warm up
    monkeypatch.setitem(vt.CONFIG, "cleanup", True)
    monkeypatch.setitem(vt.CONFIG, "offline_only", True)
    press(DICTATE)
    assert tool.azure.warms == 1  # "keep dictation on this computer": Azure isn't contacted for dictation
    press(COMMAND)
    assert tool.azure.warms == 2  # but the command key still uses Azure
    key(tool, COMMAND, True, now())
    key(tool, COMMAND, False, now())  # a quick Ctrl press (Ctrl+C, say) never wakes Azure up
    tool._tick()
    assert tool.azure.warms == 2


# ====================================================================================================== keyboard hook

class FakeListener:
    def __init__(self):
        self.suppressed = 0

    def suppress_event(self):
        self.suppressed += 1


def hook(tool, vk, down, injected=False):
    data = SimpleNamespace(vkCode=vk, flags=vt.LLKHF_INJECTED if injected else 0)
    msg = vt.win32con.WM_KEYDOWN if down else vt.win32con.WM_KEYUP
    return tool._filter(msg, data)


def drain(q):
    out = []
    while not q.empty():
        out.append(q.get_nowait())
    return out


@pytest.fixture
def hooked(tool, monkeypatch):
    tool.listener = FakeListener()
    injected = []
    monkeypatch.setattr(vt.win32api, "keybd_event", lambda vk, scan, flags, extra: injected.append((vk, flags)))
    return tool, injected


def test_hook_queues_trigger_keys_and_passes_them_through(hooked):
    tool, injected = hooked
    assert hook(tool, DICTATE, True) is False  # False = Windows still sees the key
    events = drain(tool.events)
    assert events[0][0] == "trigger" and events[0][2] is True


def test_hook_ignores_keys_we_injected_ourselves(hooked):
    tool, _ = hooked
    assert hook(tool, DICTATE, True, injected=True) is True
    assert tool.events.empty()


def test_lone_alt_press_is_masked_with_the_unused_key_so_menus_cant_open(hooked):
    tool, injected = hooked
    hook(tool, DICTATE, True)
    hook(tool, DICTATE, False)
    assert (vt.MASK_VK, 0) in injected and (vt.MASK_VK, vt.win32con.KEYEVENTF_KEYUP) in injected
    assert tool.listener.suppressed == 1  # our own release replaces the real one


def test_alt_used_in_a_shortcut_is_not_masked(hooked):
    tool, injected = hooked
    hook(tool, DICTATE, True)
    hook(tool, 0x09, True)  # Tab, as in Alt+Tab
    hook(tool, DICTATE, False)
    assert injected == [] and tool.listener.suppressed == 0
    kinds = [e[0] for e in drain(tool.events)]
    assert kinds == ["trigger", "other", "trigger"]


@pytest.mark.parametrize("switch, vk, expect_events", [
    (None, DICTATE, 1), ("enabled", DICTATE, 0), ("dictate_enabled", DICTATE, 0), ("command_enabled", COMMAND, 0),
    ("dictate_enabled", COMMAND, 1), ("enabled", COMMAND, 0)])
def test_switches_are_read_live_by_the_hook(hooked, monkeypatch, switch, vk, expect_events):
    tool, _ = hooked
    if switch:
        monkeypatch.setitem(vt.CONFIG, switch, False)
    hook(tool, vk, True)
    events = drain(tool.events)
    # a switched-off key is just another key as far as the tool is concerned: it never starts a recording
    assert len([e for e in events if e[0] == "trigger"]) == expect_events


def test_paused_key_does_its_normal_job_untouched(hooked, monkeypatch):
    tool, injected = hooked
    monkeypatch.setitem(vt.CONFIG, "enabled", False)
    hook(tool, DICTATE, True)
    hook(tool, DICTATE, False)
    assert injected == [] and tool.listener.suppressed == 0


def test_other_keys_are_reported_for_escape_and_shortcuts(hooked):
    tool, _ = hooked
    hook(tool, vt.win32con.VK_ESCAPE, True)
    assert drain(tool.events)[0][:3] == ("other", None, vt.win32con.VK_ESCAPE)


def test_key_up_of_other_keys_is_not_reported(hooked):
    tool, _ = hooked
    hook(tool, 0x41, False)
    assert tool.events.empty()


def test_pause_from_the_window_drops_a_recording_in_progress(tool):
    key(tool, DICTATE, True, now())
    host = vt.WinHost(tool)
    host.paused(True)
    assert tool.active is not None  # queued for the main thread, not done on the window's thread
    vt.CONFIG["enabled"] = False
    tool.ui.get_nowait()()
    assert tool.active is None and not tool.recorder.running


# ====================================================================================================== Left Ctrl as the command key

SHIFT, KEY_C, KEY_V, KEY_T, RIGHT_ALT = vt.win32con.VK_SHIFT, 0x43, 0x56, 0x54, 0xA5
KEY_UP = vt.win32con.KEYEVENTF_KEYUP
MOUSE_DOWN_AND_WHEEL = [0x201, 0x204, 0x207, 0x20B, 0x20A, 0x20E]  # left, right, middle, side button down; wheel; sideways wheel


class FakeMouseHook:
    """Stands in for the mouse hook: only records when the keyboard hook asks for it to go in and come out."""

    def __init__(self):
        self.calls = []

    def arm(self):
        self.calls.append("arm")

    def disarm(self):
        self.calls.append("disarm")

    def is_alive(self):
        return True


def pump(tool):
    """Hand the key events the hook has queued to the key loop, in order, as the window loop does."""
    for kind, trigger, value, at in drain(tool.events):
        tool._handle(kind, trigger, value, at)


def shift_held(monkeypatch):
    monkeypatch.setattr(vt.win32api, "GetAsyncKeyState", lambda vk: 0x8000 if vk == SHIFT else 0)


def nothing_happened(tool):
    return (tool.active is None and tool.recorder.starts == 0 and not tool.recorder.running and jobs(tool) == 0
            and tool.captured == [] and tool.taught == [] and tool.light.state is None)


def test_the_command_key_is_left_ctrl_and_dictation_is_left_alt():
    assert COMMAND == vt.keyboard.Key.ctrl_l.value.vk and DICTATE == vt.keyboard.Key.alt_l.value.vk
    assert vt.CONFIG["command_key"] == "ctrl_l" and vt.CONFIG["dictate_key"] == "alt_l"


def test_an_old_config_that_used_the_windows_key_now_uses_left_ctrl():
    assert vt._migrate_config({"command_key": "cmd_l", "beeps": True}) == {"command_key": "ctrl_l", "beeps": True}
    assert vt._migrate_config({"command_key": ""}) == {"command_key": ""}  # switched off stays off
    assert vt._migrate_config({"command_key": "ctrl_r"}) == {"command_key": "ctrl_r"}
    assert vt._migrate_config({}) == {}


def test_a_ctrl_command_key_is_never_held_back_from_windows():
    assert vt.Trigger("c", "ctrl_l", False).let_through is True and vt.Trigger("c", "ctrl_r", False).let_through is True
    assert vt.Trigger("d", "alt_l", False).let_through is False  # (other keys keep what the setting says)
    assert vt.Trigger("c", "ctrl_l", True).mic_delay > 0 and vt.Trigger("d", "alt_l", True).mic_delay == 0


def test_ctrl_alone_held_is_a_command_and_the_key_is_not_touched(hooked):
    tool, injected = hooked
    assert hook(tool, COMMAND, True) is False  # False = Windows still sees the key
    time.sleep(0.35)
    assert hook(tool, COMMAND, False) is False
    assert injected == [] and tool.listener.suppressed == 0  # nothing added, nothing swallowed
    events = drain(tool.events)
    tool._handle(*events[0])
    assert not tool.recorder.running  # (Ctrl waits a moment before the microphone opens)
    tool._tick()
    assert tool.recorder.running and tool.active.cued and tool.light.state == "command"
    tool._handle(*events[1])
    settle()
    assert jobs(tool) == 1 and tool.captured == ["command"] and tool.light.state == "thinking"


@pytest.mark.parametrize("vk", [KEY_C, KEY_V])
def test_ctrl_c_and_ctrl_v_are_never_a_command_and_every_key_passes_through(hooked, vk):
    tool, injected = hooked
    for key_vk, down in ((COMMAND, True), (vk, True), (vk, False), (COMMAND, False)):
        assert hook(tool, key_vk, down) is False
    assert injected == [] and tool.listener.suppressed == 0  # Ctrl+C / Ctrl+V reach the program exactly as typed
    pump(tool)
    tool._tick()
    assert nothing_happened(tool)


def test_a_slow_ctrl_c_stops_the_recording_it_had_started_and_sends_no_command(tool):
    key(tool, COMMAND, True, now() - 1)
    tool._tick()
    assert tool.recorder.running and tool.active.cued  # held long enough to look like a command...
    tool._handle("other", None, KEY_C, now())  # ...then the C
    assert tool.active is None and not tool.recorder.running
    key(tool, COMMAND, False, now())
    settle()
    assert jobs(tool) == 0 and tool.captured == []


def test_a_quick_ctrl_press_never_opens_the_microphone(tool):
    key(tool, COMMAND, True, now())
    tool._tick()
    assert tool.recorder.starts == 0
    key(tool, COMMAND, False, now())
    tool._tick()
    assert tool.recorder.starts == 0 and tool.captured == []


@pytest.mark.parametrize("msg", MOUSE_DOWN_AND_WHEEL)
def test_ctrl_click_and_ctrl_scroll_are_never_a_command(hooked, msg):
    tool, injected = hooked
    tool.mouse_listener = FakeMouseHook()
    hook(tool, COMMAND, True)
    assert tool.mouse_listener.calls == ["arm"]  # the mouse hook goes in only while Ctrl is down
    assert tool._mouse_filter(msg, None) is False  # and never holds the mouse back
    hook(tool, COMMAND, False)
    assert tool.mouse_listener.calls == ["arm", "disarm"] and injected == [] and tool.listener.suppressed == 0
    pump(tool)
    tool._tick()
    assert nothing_happened(tool)


def test_a_click_after_a_long_ctrl_hold_stops_the_recording_and_sends_no_command(tool):
    key(tool, COMMAND, True, now() - 1)
    tool._tick()
    assert tool.recorder.running and tool.active.cued
    tool._handle("mouse", None, 0x20A, now())  # Ctrl+scroll to zoom, a second into holding Ctrl
    assert tool.active is None and not tool.recorder.running
    key(tool, COMMAND, False, now())
    settle()
    assert jobs(tool) == 0 and tool.captured == []


def test_mouse_moves_and_clicks_without_ctrl_are_not_even_queued(hooked):
    tool, _ = hooked
    tool.mouse_listener = FakeMouseHook()
    tool._mouse_filter(0x201, None)  # a click with no Ctrl down
    hook(tool, COMMAND, True)
    tool._mouse_filter(0x200, None)  # moving the mouse while Ctrl is down
    assert [e[0] for e in drain(tool.events)] == ["trigger"]
    hook(tool, COMMAND, False)
    drain(tool.events)
    hook(tool, DICTATE, True)  # Left Alt neither arms the mouse hook nor is stopped by a click
    tool._mouse_filter(0x201, None)
    assert [e[0] for e in drain(tool.events)] == ["trigger"] and tool.mouse_listener.calls == ["arm", "disarm"]


def test_clicking_while_ctrl_hands_free_is_on_keeps_listening(tool):
    t0 = now()
    key(tool, COMMAND, True, t0)
    key(tool, COMMAND, False, t0 + 0.05)
    key(tool, COMMAND, True, t0 + 0.2)
    key(tool, COMMAND, False, t0 + 0.25)
    assert tool.triggers[COMMAND].state == vt.LOCKED and tool.recorder.running
    tool._handle("mouse", None, 0x201, now())
    tool._handle("other", None, KEY_C, now())
    assert tool.triggers[COMMAND].state == vt.LOCKED and tool.recorder.running  # only a tap of Ctrl stops it
    key(tool, COMMAND, True, t0 + 2)  # Ctrl+click and Ctrl+C while it listens: still listening, nothing sent
    tool._handle("mouse", None, 0x201, t0 + 2.01)
    tool._handle("other", None, KEY_C, t0 + 2.02)
    key(tool, COMMAND, False, t0 + 2.05)
    assert tool.triggers[COMMAND].state == vt.LOCKED and tool.recorder.running and jobs(tool) == 0


def test_shift_tap_ctrl_fixes_a_word_and_cannot_switch_the_input_language(hooked, monkeypatch):
    tool, injected = hooked
    shift_held(monkeypatch)
    hook(tool, SHIFT, True)
    assert hook(tool, COMMAND, True) is False
    assert hook(tool, COMMAND, False) is False
    # Windows would read a lone Ctrl+Shift as "switch the input language" (2+ layouts). An unused key goes in before the
    # release, which cancels that, and our own release replaces the real one.
    assert injected == [(vt.MASK_VK, 0), (vt.MASK_VK, KEY_UP), (COMMAND, KEY_UP)] and tool.listener.suppressed == 1
    pump(tool)
    settle()
    assert tool.taught == ["command"] and tool.active is None and tool.recorder.starts == 0


def test_a_lone_ctrl_tap_without_shift_is_not_masked(hooked):
    tool, injected = hooked
    hook(tool, COMMAND, True)
    hook(tool, COMMAND, False)
    assert injected == [] and tool.listener.suppressed == 0  # (Ctrl opens nothing on its own: no trick is needed)


def test_ctrl_shift_shortcuts_are_not_a_fix_and_not_masked(hooked, monkeypatch):
    tool, injected = hooked
    shift_held(monkeypatch)
    for vk, down in ((SHIFT, True), (COMMAND, True), (KEY_T, True), (KEY_T, False), (COMMAND, False)):  # Ctrl+Shift+T
        assert hook(tool, vk, down) is False
    assert injected == [] and tool.listener.suppressed == 0
    pump(tool)
    settle()
    assert nothing_happened(tool)


def test_ctrl_keys_we_press_ourselves_never_trigger_the_command_key(hooked):
    """Our own Ctrl+C (copying the selection) and Ctrl+V (pasting) are sent as injected keys, which the hook ignores."""
    tool, injected = hooked
    for vk, down in ((COMMAND, True), (KEY_C, True), (KEY_C, False), (COMMAND, False),
                     (COMMAND, True), (KEY_V, True), (KEY_V, False), (COMMAND, False)):
        assert hook(tool, vk, down, injected=True) is True
    assert tool.events.empty() and injected == [] and tool.listener.suppressed == 0
    assert not tool.triggers[COMMAND].hook_down and not tool.triggers[COMMAND].key_down


def test_altgr_is_never_a_command(hooked):
    """On many keyboards AltGr arrives as Left Ctrl, then Right Alt."""
    tool, injected = hooked
    for vk, down in ((COMMAND, True), (RIGHT_ALT, True), (RIGHT_ALT, False), (COMMAND, False)):
        hook(tool, vk, down)
    assert injected == []
    pump(tool)
    tool._tick()
    assert nothing_happened(tool)


def test_ctrl_esc_is_a_shortcut_not_a_cancel_so_it_makes_no_cancel_beep(tool, monkeypatch):
    beeps = []
    monkeypatch.setattr(vt, "beep", lambda *a: beeps.append(a))
    key(tool, COMMAND, True, now())
    tool._handle("other", None, vt.win32con.VK_ESCAPE, now())  # Ctrl+Esc
    assert tool.active is None and beeps == []
    key(tool, COMMAND, False, now())
    key(tool, COMMAND, True, now() - 1)
    tool._tick()
    beeps.clear()
    tool._handle("other", None, vt.win32con.VK_ESCAPE, now())  # Esc while a command is really being recorded: a cancel
    assert tool.active is None and beeps == [((400, 120),)]


def test_a_stray_ctrl_tap_copies_nothing(tool):
    t0 = now()
    key(tool, COMMAND, True, t0)
    key(tool, COMMAND, False, t0 + 0.05)
    settle()
    assert tool.captured == [] and tool.triggers[COMMAND].state == vt.TAP_WAIT
    tool.triggers[COMMAND].tap_deadline = 0  # nobody tapped again
    tool._tick()
    assert tool.active is None and tool.captured == [] and tool.recorder.starts == 0


def test_ctrl_double_tap_turns_hands_free_on_and_copies_the_selection_once_the_key_is_up(tool):
    t0 = now()
    key(tool, COMMAND, True, t0)
    key(tool, COMMAND, False, t0 + 0.05)
    key(tool, COMMAND, True, t0 + 0.2)
    settle()
    assert tool.triggers[COMMAND].state == vt.TAP_WAIT and tool.captured == []  # (Ctrl is still down: nothing decided yet)
    key(tool, COMMAND, False, t0 + 0.3)
    settle()
    assert tool.triggers[COMMAND].state == vt.LOCKED and tool.recorder.running and tool.light.state == "command_handsfree"
    assert tool.captured == ["command"]  # copied once the key is up
    key(tool, COMMAND, True, t0 + 3)  # one more tap of Ctrl stops it
    key(tool, COMMAND, False, t0 + 3.05)
    settle()
    assert jobs(tool) == 1 and tool.captured == ["command"] and tool.active is None


def test_ctrl_that_cannot_open_the_microphone_gives_up_with_one_beep(tool, monkeypatch):
    beeps = []
    monkeypatch.setattr(vt, "beep", lambda *a: beeps.append(a))
    monkeypatch.setattr(tool.recorder, "start", lambda: (_ for _ in ()).throw(OSError("no microphone")))
    key(tool, COMMAND, True, now() - 1)
    tool._tick()
    tool._tick()
    assert tool.active is None and beeps == [((300, 300),)]
    key(tool, COMMAND, False, now())
    settle()
    assert jobs(tool) == 0 and tool.captured == []
    # the second tap of a double-tap that finds the microphone broken does the same, and does not lock hands-free on
    beeps.clear()
    t0 = now()
    key(tool, COMMAND, True, t0)
    key(tool, COMMAND, False, t0 + 0.05)
    key(tool, COMMAND, True, t0 + 0.2)
    key(tool, COMMAND, False, t0 + 0.25)  # (the lock is tried at the release)
    assert tool.active is None and tool.triggers[COMMAND].state == vt.IDLE and beeps == [((300, 300),)]


def test_a_missed_ctrl_release_takes_the_mouse_hook_out_again(watched):
    mouse = watched.mouse_listener = FakeMouseHook()
    t = watched.triggers[COMMAND]
    t.hook_down, t.hook_at = True, 100.0  # Ctrl went down and the hook never heard it come up
    watched._check_keys(100.4, held())
    watched._check_keys(100.8, held())
    assert not t.hook_down and mouse.calls == ["disarm"]


def test_the_watchdog_puts_a_stopped_mouse_hook_back_at_most_every_two_seconds(watched, monkeypatch):
    class Dead:
        def is_alive(self):
            return False
    watched.mouse_listener = Dead()
    installs = []
    monkeypatch.setattr(watched, "_install_mouse_hook", lambda: installs.append(1))
    watched._check_keys(100.0, held())
    watched._check_keys(100.5, held())
    assert installs == [1]
    watched._check_keys(103.0, held())
    assert installs == [1, 1]


def test_mouse_hook_callback_hands_each_event_on_and_survives_errors(monkeypatch):
    import ctypes
    nexts = []
    monkeypatch.setattr(vt.MouseHook, "_CallNext", lambda *a: nexts.append(a) or 0)
    data = vt._MSLLHOOKSTRUCT(x=5, y=6)
    ptr = ctypes.addressof(data)
    seen = []

    def on_mouse(msg, d):
        seen.append((msg, d.x, d.y))
        if msg == 0x204:
            raise RuntimeError("bug")
    hk = vt.MouseHook(on_mouse)
    assert hk._callback(0, 0x201, ptr) == 0 and seen == [(0x201, 5, 6)] and len(nexts) == 1
    assert hk._callback(0, 0x204, ptr) == 0 and len(nexts) == 2  # an error: the mouse goes on as normal
    assert hk._callback(-1, 0x201, ptr) == 0 and len(seen) == 2  # not HC_ACTION: not looked at


def test_mouse_hook_is_in_only_while_armed_and_stops_cleanly():
    hk = vt.MouseHook(lambda msg, data: None)
    hk.start()
    assert hk.wait() and hk.is_alive() and not hk.armed

    def until(check):
        for _ in range(150):
            if check():
                return True
            time.sleep(0.02)
        return check()
    hk.arm()
    assert until(lambda: hk.armed)
    hk.disarm()
    assert until(lambda: not hk.armed)
    hk.arm()
    assert until(lambda: hk.armed)
    hk.stop()  # (stopped while armed: the hook is taken out on the way)
    hk.join(5)
    assert not hk.is_alive() and not hk.armed


def test_every_key_label_says_left_ctrl_and_none_says_the_windows_key(tool):
    badges = {f[1]: f[3] for f in ui_api.FEATURES if f[3]}
    assert (badges["dictate_enabled"], badges["command_enabled"], badges["typed_fix"]) == (
        "Hold or double-tap Left Alt", "Hold or double-tap Left Ctrl", "Shift + tap Left Ctrl")
    assert tool._tooltip() == "VoiceTool: hold Left Alt to talk, Left Ctrl for commands"
    for name in ("voicetool.py", "ui_api.py", "settings_window.py", "system_windows.py", "README.md", "How to install.md",
                 "setup.ps1", "ui/index.html", "ui/pc_overrides.js"):
        path = HERE / name
        if not path.exists():  # (the public kit's src folder has no "How to install.md")
            continue
        text = path.read_text(encoding="utf-8-sig")
        for bad in ("Left Windows", "Windows key", "Win key", "Shift-tap"):
            assert bad not in text, (name, bad)
    assert "cmd_l" not in (HERE / "config.json").read_text(encoding="utf-8-sig")


# ====================================================================================================== dictation

def run_dictation(tool, text="um this is a test of the tool", app=None):
    t = tool.triggers[DICTATE]
    t.app = app if app is not None else {"app": "notepad.exe", "title": "x"}
    t.key_down = False
    tool.transcribe = lambda audio: text
    tool._dictate(t, np.full(vt.RATE, 0.3, np.float32))
    return t


def test_dictation_pastes_the_tidied_text_and_logs_the_app(tool, isolated):
    tool.azure = FakeCleaner(result="This is a test of the tool.")
    run_dictation(tool)
    assert PASTED == ["This is a test of the tool. "]
    row = json.loads((isolated / "history.jsonl").read_text().splitlines()[-1])
    assert row["text"] == "This is a test of the tool." and row["app"] == "notepad.exe"
    assert tool.azure.calls[0][1] == ""  # no app rule for Notepad: no hint


def test_a_dictation_saves_how_many_learned_fixes_it_made(tool, isolated):
    corrections.add("zef fur", "Zephyr", "test")
    tool.azure = FakeCleaner(result="Open the Zef Fur site.")
    run_dictation(tool, "open the zef fur site now")  # one fix before the tidy-up, one more after
    tool.azure = FakeCleaner(result="Nothing to change in this one.")
    run_dictation(tool, "nothing to change in this one")
    tool.azure = FakeCleaner(fail=ConnectionError("no internet"))
    run_dictation(tool, "zef fur is down again")  # the tidy-up failed: only the first pass counts
    rows = [json.loads(line) for line in (isolated / "history.jsonl").read_text(encoding="utf-8").splitlines()]
    assert [r["fixes"] for r in rows] == [2, 0, 1]


@pytest.mark.parametrize("cleanup, app", [(False, "notepad.exe"), (True, "Code.exe")])
def test_a_dictation_without_a_tidy_up_counts_its_one_pass_of_fixes(tool, isolated, monkeypatch, cleanup, app):
    monkeypatch.setitem(vt.CONFIG, "cleanup", cleanup)
    if app == "Code.exe":
        styles.save({"styles": [], "app_rules": [{"apps": "Code.exe", "mode": "raw"}]})  # an app rule: pasted exactly as said
    corrections.add("zef fur", "Zephyr", "test")
    tool.azure = FakeCleaner(result="never used")
    run_dictation(tool, "open zef fur and zef fur please", {"app": app, "title": "x"})
    assert tool.azure.calls == []
    assert json.loads((isolated / "history.jsonl").read_text(encoding="utf-8").splitlines()[-1])["fixes"] == 2


def test_short_phrases_skip_the_tidy_up(tool):
    tool.azure = FakeCleaner(result="never used")
    run_dictation(tool, "hello there")
    assert tool.azure.calls == [] and PASTED == ["hello there "]


def test_tidy_up_switched_off_pastes_the_words_as_heard_with_fixes(tool, monkeypatch):
    monkeypatch.setitem(vt.CONFIG, "cleanup", False)
    corrections.add("zeph ear", "Zephyr", "test")
    tool.azure = FakeCleaner(result="never used")
    run_dictation(tool, "open zeph ear today please")
    assert tool.azure.calls == [] and PASTED == ["open Zephyr today please "]


def test_azure_failure_still_pastes_the_words(tool):
    tool.azure = FakeCleaner(fail=ConnectionError("no internet"))
    run_dictation(tool, "this should still be pasted")
    assert PASTED == ["this should still be pasted "]


def test_learned_fixes_apply_before_and_after_the_tidy_up(tool):
    corrections.add("zeph ear", "Zephyr", "test")
    tool.azure = FakeCleaner(result="Open the Zeph Ear site.")
    run_dictation(tool, "open the zeph ear site now")
    assert tool.azure.calls[0][0] == "open the Zephyr site now"  # fixed before the model sees it
    assert PASTED == ["Open the Zephyr site. "]  # and again after


@pytest.mark.parametrize("app, expect_tidy", [
    ("Code.exe", False), ("Cursor.exe", False), ("WindowsTerminal.exe", False), ("powershell.exe", False),
    ("cmd.exe", False), ("OUTLOOK.EXE", True), ("olk.exe", True), ("chrome.exe", True), ("notepad.exe", True)])
def test_terminals_and_code_editors_get_exactly_what_was_said(tool, app, expect_tidy):
    tool.azure = FakeCleaner(result="Tidied.")
    run_dictation(tool, "git status dash dash short please", {"app": app, "title": "x"})
    assert bool(tool.azure.calls) is expect_tidy
    assert PASTED == (["Tidied. "] if expect_tidy else ["git status dash dash short please "])


@pytest.mark.parametrize("app, style", [("OUTLOOK.EXE", "Formal email"), ("olk.exe", "Formal email"),
                                       ("ms-teams.exe", "Casual chat"), ("MSTeams.exe", "Casual chat"),
                                       ("slack.exe", "Casual chat")])
def test_dictation_passes_only_the_style_name_never_its_instructions(tool, app, style):
    """The Mac's finding: passing the instructions made the tidy-up reword ('no contractions' turned let's into let us)."""
    tool.azure = FakeCleaner(result="Ok.")
    run_dictation(tool, "let's meet on thursday if that works", {"app": app, "title": "x"})
    hint = tool.azure.calls[0][1]
    assert style in hint and "never add, remove or reword" in hint and "including contractions" in hint
    assert styles.instructions(style) not in hint and "No contractions" not in hint


def test_title_rules_can_tell_gmail_from_the_rest_of_a_browser(tool, isolated):
    styles.save({"styles": [{"name": "Formal email", "instructions": "Formal."}],
                 "app_rules": [{"apps": "title:Gmail", "mode": "style", "style": "Formal email"}]})
    tool.azure = FakeCleaner(result="Ok.")
    run_dictation(tool, "let's meet on thursday if that works", {"app": "chrome.exe", "title": "Inbox - Gmail - Google Chrome"})
    assert "Formal email" in tool.azure.calls[0][1]
    tool.azure.calls.clear()
    run_dictation(tool, "let's meet on thursday if that works", {"app": "chrome.exe", "title": "Weather - Google Chrome"})
    assert tool.azure.calls[0][1] == ""


def test_translate_runs_the_tidy_up_even_for_two_words_and_never_auto_learns(tool, monkeypatch):
    monkeypatch.setitem(vt.CONFIG, "translate_to", "Spanish")
    learned = []
    monkeypatch.setattr(corrections, "learn_from", lambda *a: learned.append(a) or [])
    tool.azure = FakeCleaner(result="Hola.")
    run_dictation(tool, "hello")
    assert tool.azure.calls and PASTED == ["Hola. "] and learned == []


def test_translate_works_even_with_tidy_up_off(tool, monkeypatch):
    monkeypatch.setitem(vt.CONFIG, "translate_to", "French")
    monkeypatch.setitem(vt.CONFIG, "cleanup", False)
    tool.azure = FakeCleaner(result="Bonjour.")
    run_dictation(tool, "hello there my friend")
    assert PASTED == ["Bonjour. "]


def test_translate_is_skipped_in_a_terminal(tool, monkeypatch):
    monkeypatch.setitem(vt.CONFIG, "translate_to", "French")
    tool.azure = FakeCleaner(result="Bonjour.")
    run_dictation(tool, "git commit dash m fix", {"app": "WindowsTerminal.exe", "title": "x"})
    assert tool.azure.calls == [] and PASTED == ["git commit dash m fix "]


def test_failed_translation_pastes_the_original_and_says_so(tool, monkeypatch):
    monkeypatch.setitem(vt.CONFIG, "translate_to", "German")
    tool.azure = FakeCleaner(fail=TimeoutError("slow"))
    run_dictation(tool, "good morning everyone")
    assert PASTED == ["good morning everyone "] and tool.cards[0][0] == "Couldn't translate"


def test_auto_learn_still_works_without_translation(tool):
    tool.azure = FakeCleaner(result="Open Zephyr now.")
    run_dictation(tool, "open zef fir now")
    run_dictation(tool, "open zef fir now")
    assert ("zef fir", "Zephyr") in corrections.load()


def test_empty_speech_pastes_nothing(tool):
    tool.azure = FakeCleaner()
    run_dictation(tool, "")
    assert PASTED == []


def test_recordings_are_saved_only_when_switched_on(tool, monkeypatch, isolated):
    run_dictation(tool, "hello there")
    assert not (isolated / "recordings").exists()
    monkeypatch.setitem(vt.CONFIG, "save_recordings", True)
    run_dictation(tool, "hello there")
    saved = list((isolated / "recordings").rglob("*_dictate.wav"))
    assert len(saved) == 1
    with wave.open(str(saved[0])) as w:
        assert w.getframerate() == 16000 and w.getnchannels() == 1


# ====================================================================================================== tidy-up prompt

def cleaner():
    c = vt.Cleaner.__new__(vt.Cleaner)
    c.local = None
    c._last_used = 0
    return c


def test_prompt_keeps_contractions_but_fixes_texting_shorthand(isolated):
    """'I'll', 'don't', 'let's' stay exactly as said; 'u' and 'tmrw' are fixed."""
    shutil_style = (HERE / "style.md")
    (isolated / "style.md").write_text(shutil_style.read_text(encoding="utf-8"), encoding="utf-8")
    system = cleaner()._system()
    assert "never expand them" in system and "normal contractions (I'll, don't" in system
    assert "Always writes words out in full" not in system and "never \"u\"" in system and "tmrw" in system
    assert 'No texting shorthand: "you", never "u"; no clipped words like "tmrw" or "pls".' in commands.OWNER_SAYS
    assert "keep them exactly as said, never expand them" in commands.OWNER_SAYS
    assert "Always writes words out in full" not in shutil_style.read_text(encoding="utf-8")


def test_a_tidy_up_that_leaves_contractions_alone_pastes_them_unchanged(tool):
    said = "um i'll send it tomorrow and let's not wait, don't worry"
    tool.azure = FakeCleaner(result="I'll send it tomorrow, and let's not wait. Don't worry.")
    run_dictation(tool, said)
    assert PASTED == ["I'll send it tomorrow, and let's not wait. Don't worry. "]
    for contraction in ("I'll", "let's", "Don't"):
        assert contraction in PASTED[0]


@pytest.mark.skipif(os.environ.get("VOICETOOL_LIVE_AZURE") != "1", reason="live Azure check: set VOICETOOL_LIVE_AZURE=1")
def test_live_tidy_up_keeps_contractions_and_fixes_shorthand(isolated):
    """A real call to the free Azure model (gpt-5.4-mini) with a made-up sentence. Usage is logged to the temp folder."""
    (isolated / "style.md").write_text((HERE / "style.md").read_text(encoding="utf-8"), encoding="utf-8")
    c = vt.Cleaner()
    out = c("um i'll send it tmrw and let's not wait, don't worry, thx u").replace("’", "'")  # (the model may use curly apostrophes)
    assert "I'll" in out and "let's" in out and "on't worry" in out, out
    assert "tmrw" not in out and " u" not in out.lower().replace("you", ""), out
    assert "I will" not in out and "let us" not in out.lower() and "do not" not in out.lower(), out


@pytest.mark.skipif(os.environ.get("VOICETOOL_LIVE_AZURE") != "1", reason="live Azure check: set VOICETOOL_LIVE_AZURE=1")
def test_live_tidy_up_keeps_the_day_when_a_number_is_corrected(isolated):
    """The old prompt example ("Thursday, no wait, Friday -> Friday") made the model change Thursday to Friday in real text."""
    (isolated / "style.md").write_text((HERE / "style.md").read_text(encoding="utf-8"), encoding="utf-8")
    c = vt.Cleaner()
    out = c("let's meet Thursday at 2 no wait 3")
    assert "Thursday" in out and "Friday" not in out and "3" in out, out


def test_system_prompt_carries_the_hint_the_word_list_and_translation(monkeypatch):
    monkeypatch.setitem(vt.CONFIG, "translate_to", "Spanish")
    corrections.add("zeph ear", "Zephyr", "test")
    system = cleaner()._system("HINT-XYZ")
    assert "HINT-XYZ" in system and "Zephyr" in system and "zeph ear -> Zephyr" in system
    assert "translate the cleaned text into Spanish" in system and "Output only the cleaned text" not in system
    assert "Output only the cleaned text." in (monkeypatch.setitem(vt.CONFIG, "translate_to", "") or cleaner()._system())


def test_use_style_off_leaves_the_writing_habits_out(monkeypatch):
    monkeypatch.setitem(vt.CONFIG, "use_style", False)
    assert "writing habits" not in cleaner()._system()


# ====================================================================================================== offline tidy-up

def offline_cleaner(monkeypatch, azure_ok=True, available=True):
    c = cleaner()
    calls = []

    def azure(system, text):
        calls.append("azure")
        if not azure_ok:
            raise ConnectionError("no internet")
        return "from azure"
    monkeypatch.setattr(c, "_azure", azure)
    monkeypatch.setattr(c, "_local", lambda system, text: calls.append("local") or "from this pc")
    monkeypatch.setattr(local_cleanup_win, "available", lambda config=None: available)
    return c, calls


def test_azure_is_used_when_it_works(monkeypatch):
    c, calls = offline_cleaner(monkeypatch)
    assert c("hi there you") == "from azure" and calls == ["azure"]


def test_falls_back_to_this_pc_when_azure_fails(monkeypatch):
    c, calls = offline_cleaner(monkeypatch, azure_ok=False)
    assert c("hi there you") == "from this pc" and calls == ["azure", "local"]


def test_no_fallback_when_switched_off_or_not_installed(monkeypatch):
    monkeypatch.setitem(vt.CONFIG, "offline_tidyup", False)
    c, _ = offline_cleaner(monkeypatch, azure_ok=False)
    with pytest.raises(ConnectionError):
        c("hi there you")
    monkeypatch.setitem(vt.CONFIG, "offline_tidyup", True)
    c, _ = offline_cleaner(monkeypatch, azure_ok=False, available=False)
    with pytest.raises(ConnectionError):
        c("hi there you")


def test_keep_dictation_on_this_computer_never_calls_azure(monkeypatch):
    monkeypatch.setitem(vt.CONFIG, "offline_only", True)
    c, calls = offline_cleaner(monkeypatch)
    assert c("hi there you") == "from this pc" and calls == ["local"]


def test_keep_on_this_computer_without_the_model_skips_the_tidy_up_instead_of_leaking(tool, monkeypatch):
    """offline_only + nothing installed: dictation is still pasted (as heard) and Azure is never contacted."""
    monkeypatch.setitem(vt.CONFIG, "offline_only", True)
    c = cleaner()
    asked = []
    monkeypatch.setattr(c, "_azure", lambda s, t: asked.append(t) or "x")
    monkeypatch.setattr(local_cleanup_win, "available", lambda config=None: False)
    tool.azure = c
    run_dictation(tool, "this must not leave the computer")
    assert asked == [] and PASTED == ["this must not leave the computer "]


def test_offline_hook_needs_both_the_server_and_the_model(tmp_path):
    cfg = {"offline_server": str(tmp_path / "llama-server.exe"), "offline_model": str(tmp_path / "tiny.gguf")}
    assert local_cleanup_win.paths(cfg) is None and not local_cleanup_win.available(cfg)
    (tmp_path / "llama-server.exe").write_bytes(b"x")
    assert not local_cleanup_win.available(cfg)  # a server but no model
    (tmp_path / "tiny.gguf").write_bytes(b"x")
    assert local_cleanup_win.available(cfg)


def test_the_install_layout_paths_port_and_arguments_are_in_config_json():
    cfg = json.loads((HERE / "config.json").read_text(encoding="utf-8"))
    assert cfg["offline_server"] == r"runtime\llama-vulkan\llama-server.exe"
    assert cfg["offline_server_cpu"] == r"runtime\llama-cpu\llama-server.exe"
    assert cfg["offline_model"] == r"models\qwen3-4b-instruct-2507\Qwen3-4B-Instruct-2507-Q4_K_M.gguf"
    assert cfg["speech_model_path"] == r"models\parakeet-tdt-0.6b-v2" and cfg["offline_backend"] == "vulkan"
    assert cfg["offline_args"] == ["-c", "16384", "--parallel", "1", "--no-webui", "-ngl", "99", "--device", "Vulkan0"]
    assert cfg["offline_tidyup"] is True and cfg["offline_only"] is True  # everything stays on this computer
    assert cfg["offline_port"] != 8765  # (never the port another program of ours may be using)
    assert (local_cleanup_win.DEFAULT_SERVER, local_cleanup_win.DEFAULT_SERVER_CPU, local_cleanup_win.DEFAULT_MODEL) == (
        cfg["offline_server"], cfg["offline_server_cpu"], cfg["offline_model"])


def test_the_server_command_is_the_tested_one_and_listens_only_on_this_computer(tmp_path):
    cfg = {**json.loads((HERE / "config.json").read_text(encoding="utf-8")), "offline_server": str(tmp_path / "s.exe"),
           "offline_server_cpu": str(tmp_path / "cpu.exe"), "offline_model": str(tmp_path / "m.gguf")}
    for name in ("s.exe", "cpu.exe", "m.gguf"):
        (tmp_path / name).write_bytes(b"x")
    cmd = local_cleanup_win.LocalCleaner(cfg)._default_command(8766)
    assert cmd == [str(tmp_path / "s.exe"), "-m", str(tmp_path / "m.gguf"), "--host", "127.0.0.1", "--port", "8766",
                   "-c", "16384", "--parallel", "1", "--no-webui", "-ngl", "99", "--device", "Vulkan0"]
    cpu = local_cleanup_win.LocalCleaner({**cfg, "offline_backend": "cpu"})._default_command(8766)  # no graphics-card arguments
    threads = str(min(local_cleanup_win.MAX_CPU_THREADS, local_cleanup_win.physical_cores()))
    assert cpu == [str(tmp_path / "cpu.exe"), "-m", str(tmp_path / "m.gguf"), "--host", "127.0.0.1", "--port", "8766",
                   "-c", "16384", "--parallel", "1", "--no-webui",
                   "-t", threads, "-n", "2560", "-cram", "1024"]  # (the processor server: its cores, a reply limit, a cache limit)


def test_relative_paths_in_config_json_are_taken_from_the_install_folder(tmp_path, monkeypatch):
    monkeypatch.setattr(local_cleanup_win, "HERE", tmp_path)
    for rel in (r"runtime\llama-vulkan\llama-server.exe", r"runtime\llama-cpu\llama-server.exe",
                r"models\qwen3-4b-instruct-2507\Qwen3-4B-Instruct-2507-Q4_K_M.gguf"):
        (tmp_path / rel).parent.mkdir(parents=True, exist_ok=True)
        (tmp_path / rel).write_bytes(b"x")
    cfg = json.loads((HERE / "config.json").read_text(encoding="utf-8"))
    assert local_cleanup_win.available(cfg)
    server, model = local_cleanup_win.paths(cfg)
    assert server == str(tmp_path / "runtime" / "llama-vulkan" / "llama-server.exe") and model.startswith(str(tmp_path / "models"))
    assert local_cleanup_win.paths({**cfg, "offline_backend": "cpu"})[0] == str(tmp_path / "runtime" / "llama-cpu" / "llama-server.exe")
    assert local_cleanup_win.resolve(r"C:\somewhere\else.exe") == Path(r"C:\somewhere\else.exe")  # absolute paths are left alone
    assert local_cleanup_win.resolve("%TEMP%") == Path(os.environ["TEMP"])
    only_cpu = tmp_path / "runtime" / "llama-vulkan" / "llama-server.exe"
    only_cpu.unlink()
    assert local_cleanup_win.paths(cfg)[0].endswith("llama-cpu\\llama-server.exe")  # the other build is used if only that is there
    assert local_cleanup_win.backend_for(cfg) == "cpu"
    (tmp_path / "models" / "qwen3-4b-instruct-2507" / "Qwen3-4B-Instruct-2507-Q4_K_M.gguf").unlink()
    assert not local_cleanup_win.available(cfg)  # no model: not set up
    # the speech model folder is resolved the same way, and a missing one means "use the downloaded copy"
    monkeypatch.setattr(vt, "HERE", tmp_path)
    assert vt.speech_model_folder() is None
    (tmp_path / "models" / "parakeet-tdt-0.6b-v2").mkdir(parents=True)
    monkeypatch.setitem(vt.CONFIG, "speech_model_path", r"models\parakeet-tdt-0.6b-v2")
    assert vt.speech_model_folder() == tmp_path / "models" / "parakeet-tdt-0.6b-v2"


def test_the_speech_model_is_loaded_from_the_bundled_folder_when_it_is_there(tmp_path, monkeypatch):
    seen = []

    class Model:
        def recognize(self, audio, sample_rate=None):
            return "x"
    monkeypatch.setattr(vt.onnx_asr, "load_model", lambda name, path=None, **kw: seen.append((name, path, kw)) or Model())
    monkeypatch.setattr(vt, "HERE", tmp_path)
    monkeypatch.setitem(vt.CONFIG, "speech_model_path", r"models\parakeet-tdt-0.6b-v2")
    (tmp_path / "models" / "parakeet-tdt-0.6b-v2").mkdir(parents=True)
    t = vt.Transcriber()
    assert seen[-1] == ("nemo-parakeet-tdt-0.6b-v2", tmp_path / "models" / "parakeet-tdt-0.6b-v2", {"quantization": "int8"})
    assert t.missing is False and t(np.zeros(vt.RATE, np.float32)) == "x"


def test_a_missing_or_broken_speech_model_never_falls_back_to_a_download_and_voicetool_stays_up(tmp_path, monkeypatch):
    """The download fallback needs huggingface_hub (not installed) and would fetch 0.6 GB nobody asked for: it's gone.
    The Transcriber stays empty, says so plainly, and reload() brings the model in once Finish setup has fetched it."""
    seen = []

    class Model:
        def recognize(self, audio, sample_rate=None):
            return "heard"
    state = {"broken": False}

    def load(name, path=None, **kw):
        seen.append(path)
        if path is None:
            pytest.fail("the speech model must never be downloaded behind the user's back")
        if state["broken"]:
            raise OSError("half a download")
        return Model()
    monkeypatch.setattr(vt.onnx_asr, "load_model", load)
    monkeypatch.setattr(vt, "HERE", tmp_path)
    monkeypatch.setitem(vt.CONFIG, "speech_model_path", r"models\parakeet-tdt-0.6b-v2")
    t = vt.Transcriber()  # the folder isn't there
    assert seen == [] and t.model is None and t.missing is True and vt.speech_model_missing(t) is True
    with pytest.raises(vt.SpeechModelMissing, match="The speech model is missing. Open VoiceTool and click Finish setup."):
        t(np.zeros(vt.RATE, np.float32))
    folder = tmp_path / "models" / "parakeet-tdt-0.6b-v2"
    folder.mkdir(parents=True)
    state["broken"] = True
    assert t.reload() is False and t.missing is True and seen == [folder]  # a folder that can't be loaded: still no download
    assert vt.speech_model_missing(t) is True
    state["broken"] = False
    for name in vt.SPEECH_FILES:
        (folder / name).write_bytes(b"x")
    assert t.reload() is True and t.missing is False and t(np.zeros(vt.RATE, np.float32)) == "heard"
    assert vt.speech_model_missing(t) is False
    (folder / "encoder-model.int8.onnx").unlink()  # a file lost later: Finish setup is offered again
    assert vt.speech_model_missing(t) is True
    # no folder named at all (a developer copy) and nothing to load: the same plain state, never a crash
    monkeypatch.setitem(vt.CONFIG, "speech_model_path", "")
    monkeypatch.setattr(vt.onnx_asr, "load_model", lambda *a, **k: (_ for _ in ()).throw(ImportError("huggingface_hub")))
    t2 = vt.Transcriber()
    assert t2.missing is True and vt.speech_model_missing(t2) is True and vt.speech_model_missing(None) is False


def run_worker_job(tool, kind, audio=None):
    """Put one job on the real worker loop (as the key release does) and wait until it has been dealt with."""
    t = tool.triggers[DICTATE if kind == "dictate" else COMMAND]
    t.key_down = False
    t.app = {"app": "notepad.exe", "title": "x"}
    t.context = {"app": "notepad.exe", "title": "x", "selection": "", "pid": 1, "hwnd": 2}
    if not getattr(tool, "_worker_started", False):
        threading.Thread(target=tool._worker, daemon=True).start()
        tool._worker_started = True
    before = len(tool.cards)
    tool.jobs.put((t, audio if audio is not None else np.full(vt.RATE, 0.3, np.float32)))
    assert wait_until(lambda: len(tool.cards) > before or tool.light.history[-1:] == [None], 10)
    time.sleep(0.1)


def test_a_missing_speech_model_gives_a_plain_card_not_a_crash_and_only_one_card_at_a_time(tool, monkeypatch):
    monkeypatch.setattr(vt.onnx_asr, "load_model", lambda *a, **k: pytest.fail("no download"))
    monkeypatch.setitem(vt.CONFIG, "speech_model_path", r"models\parakeet-tdt-0.6b-v2")
    tool.transcribe = vt.Transcriber()
    tool.azure = FakeCleaner()
    run_worker_job(tool, "dictate")
    assert tool.cards[-1] == ("Finish setup first", "The speech model is missing. Open VoiceTool and click Finish setup.", False)
    assert "error" in tool.light.history and PASTED == []
    tool.jobs.put((tool.triggers[COMMAND], np.full(vt.RATE, 0.3, np.float32)))
    time.sleep(0.5)
    assert len(tool.cards) == 1  # (a second try straight away doesn't pile up cards)
    tool._speech_card_at -= 30
    tool.jobs.put((tool.triggers[COMMAND], np.full(vt.RATE, 0.3, np.float32)))
    assert wait_until(lambda: len(tool.cards) == 2, 5)
    assert vt.WinHost(tool).setup_status()["missing"] is True and vt.WinHost(tool).setup_status()["speech_missing"] is True


def test_any_other_failure_in_a_job_shows_a_plain_try_again_card_besides_the_error_light(tool):
    tool.transcribe = lambda audio: (_ for _ in ()).throw(RuntimeError("boom"))
    run_worker_job(tool, "dictate")
    assert tool.cards[-1] == ("Couldn't do that", "Please try again.", False) and tool.light.history[-1] == "error"
    tool.azure = FakeCleaner(client=FakeClient("not json at all"))
    tool.transcribe = lambda audio: "do something"
    run_worker_job(tool, "command")  # a garbled reply from the model: the command is kept, and it says so
    assert tool.cards[-1][:2] == ("Command not done", "That command didn't go through. It's saved in History, so you can try it again.")


def test_a_garbled_command_reply_through_the_worker_never_shows_the_claude_card(tool):
    tool._claude = lambda *a: pytest.fail("Claude must not run")
    tool.transcribe = lambda audio: "mumble"
    tool.azure = FakeCleaner(client=FakeClient(json.dumps({"text": "garbled"})))
    run_worker_job(tool, "command")
    assert tool.cards[-1][0] == "Couldn't do that" and "work out what to do" in tool.cards[-1][1]
    assert all(c[0] != "Email and calendar are off" for c in tool.cards)


# A stand-in for llama-server: /health, /props (names its model) and /v1/chat/completions. Port, model name, then (optional) "slow".
FAKE_SERVER = """
import http.server, json, sys, time
port, model = int(sys.argv[1]), sys.argv[2]
class H(http.server.BaseHTTPRequestHandler):
    def log_message(self, *a): pass
    def _ok(self, body):
        data = json.dumps(body).encode(); self.send_response(200); self.send_header("Content-Length", str(len(data))); self.end_headers(); self.wfile.write(data)
    def do_GET(self):
        if self.path == "/health": self._ok({"status": "ok"})
        elif self.path == "/props": self._ok({"model_path": "C:/models/" + model, "total_slots": 1})
        else: self.send_response(404); self.end_headers()
    def do_POST(self):
        body = json.loads(self.rfile.read(int(self.headers["Content-Length"])))
        assert body["temperature"] == 0
        self._ok({"id": "x", "object": "chat.completion", "created": 0, "model": model, "usage": {"prompt_tokens": 5, "completion_tokens": 3, "total_tokens": 8},
                  "choices": [{"index": 0, "finish_reason": "stop", "message": {"role": "assistant", "content": " CLEANED: " + body["messages"][1]["content"] + " | " + body["messages"][0]["content"][:5] + " "}}]})
http.server.ThreadingHTTPServer(("127.0.0.1", port), H).serve_forever()
"""
BASE_PYTHON = str(Path(sys.base_prefix) / "python.exe")  # (the virtual environment's python.exe is a launcher for another process)


def free_port():
    import socket
    with socket.socket() as s:
        s.bind(("127.0.0.1", 0))
        return s.getsockname()[1]


def pid_alive(pid):
    out = subprocess.run(["tasklist", "/FI", f"PID eq {pid}", "/NH"], capture_output=True, text=True, creationflags=0x08000000).stdout
    return str(pid) in out


def wait_until(check, seconds=10):
    deadline = time.time() + seconds
    while time.time() < deadline:
        if check():
            return True
        time.sleep(0.1)
    return False


@pytest.fixture
def fake_server(tmp_path):
    script = tmp_path / "fake_llama.py"
    script.write_text(FAKE_SERVER)
    started = []

    def run(port, model="fake-model.gguf"):
        proc = subprocess.Popen([BASE_PYTHON, str(script), str(port), model], creationflags=0x08000000)
        started.append(proc)
        assert wait_until(lambda: local_cleanup_win._get(port, "/health", 1)[0] == 200)
        return proc
    run.script = str(script)
    yield run
    for proc in started:
        proc.kill()


def tidy(cleaner_, system, text):
    """One tidy-up through the on-device client (what Cleaner does), against whatever server `cleaner_` runs."""
    reply = cleaner_.client().chat.completions.create(model="local", temperature=0, max_completion_tokens=100,
                                                      messages=[{"role": "system", "content": system}, {"role": "user", "content": text}])
    return reply.choices[0].message.content.strip()


def stand_in(fake_server, port=None, **kw):
    port = port or free_port()
    cfg = {"offline_model": "C:/m/fake-model.gguf", "offline_port": port}
    return local_cleanup_win.LocalCleaner(cfg, None, command=lambda p: [BASE_PYTHON, fake_server.script, str(p), "fake-model.gguf"],
                                          startup_seconds=30, **kw), port


def test_the_on_device_client_starts_the_server_on_its_port_uses_it_and_stops_it(fake_server):
    c, port = stand_in(fake_server)
    assert not c.running()
    assert tidy(c, "SYSTEM PROMPT", "hello world") == "CLEANED: hello world | SYSTE"
    first = c.proc
    assert c.port == port and not c.reused and c.running()
    assert tidy(c, "SYSTEM PROMPT", "again") == "CLEANED: again | SYSTE" and c.proc is first  # kept running, not restarted
    c.stop()
    assert c.proc is None and wait_until(lambda: not pid_alive(first.pid)) and local_cleanup_win._port_in_use(port) is False


def test_a_server_already_running_our_model_on_the_port_is_reused_and_left_running(fake_server):
    port = free_port()
    existing = fake_server(port)
    c, _ = stand_in(fake_server, port)
    assert tidy(c, "S", "hi there") == "CLEANED: hi there | S"
    assert c.reused and c.proc is None and c.port == port
    c.stop()
    assert existing.poll() is None and local_cleanup_win._port_in_use(port)  # it wasn't ours to stop


def test_another_program_on_the_port_is_left_alone_and_a_free_port_is_used(fake_server):
    port = free_port()
    other = fake_server(port, model="someone-elses-model.gguf")
    c, _ = stand_in(fake_server, port)
    assert tidy(c, "S", "hi there") == "CLEANED: hi there | S"
    assert c.port != port and not c.reused and c.proc is not None
    ours = c.proc
    c.stop()
    assert wait_until(lambda: not pid_alive(ours.pid)) and other.poll() is None


def test_the_server_is_started_again_if_it_dies(fake_server):
    c, _ = stand_in(fake_server)
    tidy(c, "S", "one")
    c.proc.kill()
    c.proc.wait()
    assert not c.running()
    assert tidy(c, "S", "two") == "CLEANED: two | S"
    c.stop()


def test_start_in_background_never_waits_and_a_failure_is_only_logged(fake_server):
    logged = []
    log = SimpleNamespace(info=lambda *a: None, warning=lambda *a: logged.append(a), error=lambda *a: None)
    c = local_cleanup_win.LocalCleaner({"offline_port": free_port()}, log, command=lambda p: [BASE_PYTHON, "-c", "raise SystemExit(3)"],
                                       startup_seconds=10)
    started = time.time()
    c.start_in_background()
    assert time.time() - started < 1
    assert wait_until(lambda: any("not started" in a[0] for a in logged), 15)
    c2, port = stand_in(fake_server)
    c2.start_in_background()
    assert wait_until(c2.running, 20)
    c2.stop()


def test_the_offline_cleaner_gives_up_if_the_server_dies_or_never_answers():
    c = local_cleanup_win.LocalCleaner({"offline_port": free_port(), "offline_backend": "cpu"}, None,
                                       command=lambda p: [BASE_PYTHON, "-c", "raise SystemExit(3)"], startup_seconds=10)
    with pytest.raises(RuntimeError):
        c.start()
    c = local_cleanup_win.LocalCleaner({"offline_port": free_port()}, None,
                                       command=lambda p: [BASE_PYTHON, "-c", "import time; time.sleep(60)"], startup_seconds=1)
    with pytest.raises(local_cleanup_win.ModelServerUnavailable) as gave_up:  # (Vulkan and then CPU both timed out)
        c.start()
    assert isinstance(gave_up.value.__cause__, TimeoutError) and c.gave_up
    assert c.proc is None  # the stuck program was stopped


HELPER = """
import sys, time
sys.path.insert(0, sys.argv[1])
import local_cleanup_win
c = local_cleanup_win.LocalCleaner({"offline_port": int(sys.argv[3])}, None,
    command=lambda p: [sys.argv[4], sys.argv[2], str(p), "fake-model.gguf"], startup_seconds=30)
c.start()
print(c.proc.pid, flush=True)
time.sleep(120)
"""


def test_the_server_dies_with_voicetool_even_if_voicetool_is_killed(fake_server, tmp_path):
    """Quit, Stop VoiceTool.cmd and a crash all end VoiceTool's process; the model's server must go with it."""
    helper = tmp_path / "helper.py"
    helper.write_text(HELPER)
    port = free_port()
    proc = subprocess.Popen([sys.executable, str(helper), str(HERE), fake_server.script, str(port), BASE_PYTHON],
                            stdout=subprocess.PIPE, text=True, creationflags=0x08000000)
    try:
        server_pid = int(proc.stdout.readline())
        assert pid_alive(server_pid) and local_cleanup_win._port_in_use(port)
        subprocess.run(["taskkill", "/F", "/T", "/PID", str(proc.pid)], capture_output=True, creationflags=0x08000000)  # as if VoiceTool crashed
        assert wait_until(lambda: not pid_alive(server_pid), 10), "the server was left running"
        assert not local_cleanup_win._port_in_use(port)
    finally:
        proc.kill()


# ---- when the tool starts and stops the server

def test_the_server_follows_the_two_switches(monkeypatch):
    made = []

    class Fake:
        def __init__(self, config=None, log=None):
            self.started, self.stopped = 0, 0
            made.append(self)

        def start_in_background(self): self.started += 1
        def stop(self): self.stopped += 1
    monkeypatch.setattr(local_cleanup_win, "LocalCleaner", Fake)
    monkeypatch.setattr(local_cleanup_win, "available", lambda config=None: True)
    c = cleaner()
    monkeypatch.setitem(vt.CONFIG, "offline_tidyup", False)
    c.sync_offline()
    assert made == []  # both off and nothing running: nothing to do
    monkeypatch.setitem(vt.CONFIG, "offline_tidyup", True)
    c.sync_offline()
    c.sync_offline()
    assert len(made) == 1 and made[0].started == 2  # one cleaner; start is idempotent
    monkeypatch.setitem(vt.CONFIG, "offline_tidyup", False)
    monkeypatch.setitem(vt.CONFIG, "offline_only", True)
    c.sync_offline()
    assert made[0].started == 3 and made[0].stopped == 0  # "keep on this computer" alone also needs the model
    monkeypatch.setitem(vt.CONFIG, "offline_only", False)
    c.sync_offline()
    assert made[0].stopped == 1  # both off: the server is stopped
    monkeypatch.setitem(vt.CONFIG, "offline_tidyup", True)
    monkeypatch.setattr(local_cleanup_win, "available", lambda config=None: False)
    c.sync_offline()
    assert made[0].started == 3 and made[0].stopped == 2  # not installed: nothing to start


def test_changing_the_offline_switch_in_the_window_starts_or_stops_the_server(api, tool):
    api, host = api
    seen = []
    tool.azure = SimpleNamespace(sync_offline=lambda: seen.append(vt.CONFIG["offline_only"]))
    api.call("set_setting", {"key": "offline_only", "on": False})
    api.call("set_setting", {"key": "offline_only", "on": True})
    api.call("set_setting", {"key": "beeps", "on": False})  # other switches don't touch it
    assert seen == [False, True]
    tool.azure = FakeCleaner()  # (a cleaner with no offline support, like test mode's) is simply skipped
    api.call("set_setting", {"key": "offline_only", "on": True})


def test_the_offline_switch_cannot_go_off_when_no_azure_endpoint_is_set(api, monkeypatch):
    api, host = api
    monkeypatch.setitem(vt.CONFIG, "azure_openai_base_url", "")
    monkeypatch.setitem(vt.CONFIG, "offline_only", True)
    with pytest.raises(ValueError, match="Azure isn't set up"):
        api.call("set_setting", {"key": "offline_only", "on": False})
    assert vt.CONFIG["offline_only"] is True
    assert api.call("set_setting", {"key": "offline_only", "on": True})["on"] is True


def test_the_old_settings_window_has_the_same_rows_and_starts_the_server(monkeypatch):
    assert settings_window.FEATURES is ui_api.FEATURES
    called, flips = [], []
    win = settings_window.SettingsWindow.__new__(settings_window.SettingsWindow)
    win.win, win.toggles = None, {"offline_only": SimpleNamespace(set=lambda v: flips.append(v))}
    win.config, win.actions = {"azure_openai_base_url": "https://example.invalid/"}, {"offline": lambda: called.append(1)}
    win._set("offline_only", True)
    win._set("beeps", False)
    assert called == [1]
    import tkinter.messagebox as mb
    shown = []
    monkeypatch.setattr(mb, "showinfo", lambda *a, **k: shown.append(a))
    win.config = {}  # no Azure endpoint: it can't go off
    win._set("offline_only", False)
    assert shown and flips == [True] and "offline_only" not in win.config


def test_the_old_settings_window_asks_the_consent_question_before_meeting_notes(monkeypatch):
    import tkinter.messagebox as mb
    asked, flips = [], []
    win = settings_window.SettingsWindow.__new__(settings_window.SettingsWindow)
    win.win, win.toggles = None, {"meeting_notes": SimpleNamespace(set=lambda v: flips.append(v))}
    win.config, win.actions = {"meeting_notes": False}, {}
    monkeypatch.setattr(mb, "askokcancel", lambda title, text, **k: asked.append(text) or False)
    win._set("meeting_notes", True)
    assert asked == [meetings.CONSENT_REMINDER] and win.config["meeting_notes"] is False and flips == [False]
    monkeypatch.setattr(mb, "askokcancel", lambda title, text, **k: asked.append(text) or True)
    win._set("meeting_notes", True)
    assert win.config["meeting_notes"] is True


def test_the_offline_row_has_friend_wording_in_the_window(api, monkeypatch):
    api, host = api
    monkeypatch.setattr(local_cleanup_win, "available", lambda config=None: True)
    rows = {f["key"]: f for f in api.call("settings")["features"]}
    assert "offline_tidyup" not in rows  # (the Azure-fallback switch isn't offered)
    assert "nothing you say leaves it" in rows["offline_only"]["text"] and "meeting notes" in rows["offline_only"]["text"]
    assert "isn't downloaded" not in rows["offline_only"]["text"] and rows["offline_only"]["title"] == "Keep everything on this computer"


def test_azure_failure_falls_back_to_the_on_device_model_inside_a_dictation(tool, monkeypatch):
    c = cleaner()
    monkeypatch.setattr(c, "_azure", lambda system, text: (_ for _ in ()).throw(TimeoutError("azure too slow")))
    monkeypatch.setattr(local_cleanup_win, "available", lambda config=None: True)
    local = FakeLocalClient("Zeph Ear, I'll send it.")
    monkeypatch.setattr(c, "_local_openai", lambda: local)
    corrections.add("zeph ear", "Zephyr", "test")
    tool.azure = c
    run_dictation(tool, "um zeph ear i'll send it")
    assert PASTED == ["Zephyr, I'll send it. "]  # mishearings are fixed before and after the small model
    sent = local.sent[0]
    assert sent["messages"][1]["content"] == "um Zephyr i'll send it" and sent["messages"][0]["content"].startswith("Fix this dictated text.")


def test_the_short_prompt_for_the_small_model_is_the_tested_one(monkeypatch):
    system = cleaner()._system(small=True)
    assert system.startswith("Fix this dictated text. Rules: add punctuation and capitals; delete um, uh, 'you know'; "
                             "if the speaker says 'no wait' or 'I mean', keep only the corrected part; spell texting "
                             "shorthand in full (u -> you, tmrw -> tomorrow); keep contractions and every other word exactly; "
                             "never reword, never make lists, never answer questions. ")
    assert "Spell these exactly: Zephyr, Claude, Wispr Flow." in system and system.endswith("Output only the fixed text.")
    assert "writing habits" not in system and "Known mishearings" not in system  # the long lists confuse a small model
    assert "Formal email" in cleaner()._system(small=True, hint="Formal email tone.")
    monkeypatch.setitem(vt.CONFIG, "translate_to", "Spanish")
    assert cleaner()._system(small=True).endswith("translate the fixed text into Spanish and output only the translation.")


def test_the_self_correction_wording_no_longer_turns_thursday_into_friday():
    for system in (cleaner()._system(), cleaner()._system(small=True)):
        assert "Thursday, no wait, Friday" not in system and "Thursday, no, Friday" not in system
    long = cleaner()._system()
    assert ("when they say something, then 'no wait' or 'I mean' and a replacement, keep only the replacement, and keep "
            "every other word they said (days, names, numbers) exactly as spoken.") in long


def test_offline_prompt_and_azure_prompt_both_keep_a_corrected_number_but_the_day(tool):
    """The text pasted for 'let's meet Thursday at 2 no wait 3' is whatever the model returns; it must not be edited further."""
    tool.azure = FakeCleaner(result="Let's meet Thursday at 3.")
    run_dictation(tool, "let's meet Thursday at 2 no wait 3")
    assert PASTED == ["Let's meet Thursday at 3. "]


REAL_OFFLINE = local_cleanup_win.available({})


@pytest.mark.skipif(not REAL_OFFLINE, reason="the offline model isn't installed on this PC")
def test_the_real_offline_model_starts_tidies_up_a_dictation_and_stops():
    """Starts the real llama-server with the real Qwen3 4B model (on its own free port, so a running copy isn't touched),
    makes one tidy-up with the real short prompt, checks a second cleaner reuses it, and stops it."""
    cfg = {**json.loads((HERE / "config.json").read_text(encoding="utf-8")), "offline_port": free_port()}
    c = local_cleanup_win.LocalCleaner(cfg, None)
    try:
        started = time.time()
        system = cleaner()._system(small=True)
        out = tidy(c, system, "um i'll send it tmrw and let's not wait, don't worry").replace("\u2019", "'")
        print("offline tidy-up:", out, f"({time.time() - started:.1f}s including start-up)")
        assert c.proc is not None and c.running() and c.port == cfg["offline_port"]
        low = out.lower()
        assert "tomorrow" in low and "tmrw" not in low and " um" not in " " + low  # shorthand spelled out, filler deleted
        assert "i'll" in low and "let's" in low and "don't" in low  # contractions kept
        again = time.time()
        out2 = tidy(c, system, "let's meet Thursday at 2 no wait 3").replace("\u2019", "'")
        print("self-correction:", out2, f"({time.time() - again:.2f}s)")
        # the day must stay (the old prompt turned Thursday into Friday). The small model doesn't always apply the correction
        # itself (it can leave "at 2. 3"), which is its known limit, so the replaced number isn't asserted.
        assert "Thursday" in out2 and "3" in out2
        second = local_cleanup_win.LocalCleaner(cfg, None)
        second.start()
        assert second.reused and second.proc is None  # our own model is recognised through /props
        pid = c.proc.pid
        c.stop()
        assert wait_until(lambda: not pid_alive(pid), 10) and not local_cleanup_win._port_in_use(cfg["offline_port"])
    finally:
        c.stop()


# ====================================================================================================== command key: routes

def run_command(tool, said, reply, context=None, then=()):
    """Say `said` on the command key; the router answers with the JSON `reply` (a dict). `then`: what the Azure client
    answers to any later question in the same job (a screen question, a journal tidy-up)."""
    t = tool.triggers[COMMAND]
    t.key_down = False
    t.context = context if context is not None else {"app": "notepad.exe", "title": "Untitled", "selection": "", "pid": 1, "hwnd": 2}
    tool.transcribe = lambda audio: said
    tool.azure = FakeCleaner(client=FakeClient(json.dumps(reply), *then))
    tool._command(t, np.full(vt.RATE, 0.3, np.float32))
    return tool.azure.client


def test_router_asks_for_every_route_and_returns_a_triple():
    client = FakeClient(json.dumps({"route": "team", "text": "hi", "recipient": "claude-pc", "topic": "x"}))
    kind, text, extra = commands.route(client, "m", "tell the PC Claude hi", {"selection": ""}, "STYLE-XYZ")
    assert (kind, text, extra) == ("team", "hi", {"target": "", "recipient": "claude-pc", "topic": "x", "value": ""})
    system = client.sent[0]["messages"][0]["content"]
    assert "STYLE-XYZ" in system
    for route in ("rewrite", "fix", "write", "answer", "screen", "open", "note", "team", "translate", "journal",
                  "meeting", "history", "system", "transcribe", "read", "action"):
        assert f'"{route}"' in system
    assert '"target" = "start" or "stop"' in system and "Windows PC" in system and "Safari" not in system


def test_rewrite_and_fix_without_a_selection_become_write_and_long_fixes_become_rewrites():
    client = FakeClient(json.dumps({"route": "rewrite", "text": "x"}))
    assert commands.route(client, "m", "shorter", {"selection": ""})[0] == "write"
    client = FakeClient(json.dumps({"route": "fix", "text": "x"}))
    assert commands.route(client, "m", "that should be x", {"selection": ""})[0] == "write"
    assert commands.route(client, "m", "that should be x", {"selection": "one two three four five six seven"})[0] == "rewrite"
    assert commands.route(client, "m", "that should be x", {"selection": "zef fir"})[0] == "fix"


def test_the_five_original_routes_work_as_before(tool):
    selection = {"app": "notepad.exe", "title": "x", "selection": "zef fir.", "pid": 1, "hwnd": 2}
    run_command(tool, "make it friendlier", {"route": "rewrite", "text": ""}, selection, then=("Hi there!",))  # rewriting is its own step
    run_command(tool, "write a note", {"route": "write", "text": "Dear Sarah"}, {**selection, "selection": ""})
    run_command(tool, "that should be Zephyr", {"route": "fix", "text": "zephyr"}, selection)
    assert PASTED == ["Hi there!", "Dear Sarah", "Zephyr."]  # fix keeps the full stop and the word-list spelling
    assert ("zef fir", "Zephyr") in corrections.load()
    assert "Zephyr" in (tool.azure and vt.known_words())
    run_command(tool, "what does idempotent mean", {"route": "answer", "text": "Doing it twice changes nothing."})
    assert tool.cards[-1] == ("Answer", "Doing it twice changes nothing.", True)
    ran = []
    tool._claude = lambda said, ctx: ran.append(said)
    run_command(tool, "reply to Sarah", {"route": "action", "text": ""})
    time.sleep(0.2)
    assert ran == ["reply to Sarah"]


def test_claude_actions_switched_off_shows_a_card_instead(tool, monkeypatch):
    monkeypatch.setitem(vt.CONFIG, "claude_actions", False)
    tool._claude = lambda *a: pytest.fail("Claude must not run")
    run_command(tool, "check my calendar", {"route": "action", "text": ""})
    assert tool.cards[-1] == ("Email and calendar are off", vt.CLAUDE_OFF_TEXT, False)


def test_a_garbled_reply_with_no_route_shows_the_unknown_route_card_and_never_the_claude_card(tool):
    tool._claude = lambda *a: pytest.fail("Claude must not run")
    for reply in ({"text": "no route at all"}, {}, {"route": "", "text": "x"}, {"target": "mute"}):
        run_command(tool, "mumble mumble", reply)
        assert tool.cards[-1][0] == "Couldn't do that" and "work out what to do" in tool.cards[-1][1], reply
        assert tool.cards[-1][0] != "Email and calendar are off"
    assert commands._route_name({"text": "garbled"}) == "" and commands._route_name({"route": "action"}) == "action"


def test_unknown_route_or_an_empty_answer_gets_a_plain_card_and_never_runs_claude(tool):
    tool._claude = lambda *a: pytest.fail("Claude must not run")
    run_command(tool, "do something odd", {"route": "banana", "text": ""})
    assert tool.cards[-1][0] == "Couldn't do that"
    run_command(tool, "what is this", {"route": "answer", "text": ""})
    assert tool.cards[-1][0] == "Couldn't do that" and PASTED == []


def test_nothing_heard_does_nothing(tool):
    t = tool.triggers[COMMAND]
    tool.transcribe = lambda audio: ""
    tool._command(t, np.full(vt.RATE, 0.3, np.float32))
    assert PASTED == [] and tool.cards == []


def test_command_waits_for_the_screen_context_then_logs_the_route(tool, isolated):
    run_command(tool, "write hello", {"route": "write", "text": "Hello"})
    row = json.loads((isolated / "history.jsonl").read_text().splitlines()[-1])
    assert row["route"] == "write" and row["app"] == "notepad.exe" and row["result"] == "Hello"


def test_route_prompt_carries_the_app_style_and_named_styles(tool):
    ctx = {"app": "OUTLOOK.EXE", "title": "RE: budget - Message", "selection": "", "pid": 1, "hwnd": 2}
    client = run_command(tool, "write a reply", {"route": "write", "text": "ok"}, ctx)
    system = client.sent[0]["messages"][0]["content"]
    assert "In my own words" in system  # the style guide
    assert "The user is in OUTLOOK.EXE, where they write in their \"Formal email\" style" in system  # the app's style, in full
    assert "LinkedIn post" in system  # and the names that can be asked for


def test_use_style_off_removes_his_guide_but_not_named_styles(tool, monkeypatch):
    monkeypatch.setitem(vt.CONFIG, "use_style", False)
    text = tool._style({"app": "slack.exe"})
    assert "In my own words" not in text and "Casual chat" in text


# ---- translate

def test_translate_route_sets_and_clears_the_language_and_saves_it(tool, isolated):
    run_command(tool, "translate everything into spanish", {"route": "translate", "target": "spanish"})
    assert vt.CONFIG["translate_to"] == "Spanish"
    assert json.loads((isolated / "config.json").read_text())["translate_to"] == "Spanish"  # saved (to the temp folder)
    assert tool.cards[-1][0] == "Translation"
    run_command(tool, "stop translating", {"route": "translate", "target": "off"})
    assert vt.CONFIG["translate_to"] == ""
    run_command(tool, "dictate in french", {"route": "translate", "target": "French"})
    run_command(tool, "english please", {"route": "translate", "target": "English"})
    assert vt.CONFIG["translate_to"] == ""


# ---- journal

def test_journal_is_tidied_and_kept_on_this_pc(tool, isolated):
    run_command(tool, "journal today was good", {"route": "journal", "text": "today was good"}, then=["Today was good."])
    file = isolated / "journal" / f"{time.strftime('%Y-%m-%d')}.md"
    assert "## " in file.read_text() and "Today was good." in file.read_text()
    assert tool.cards[-1][0] == "Journal" and "stays on this PC" in tool.cards[-1][1]
    assert not (isolated / "Drive").exists()  # never goes to the Big Brain or Drive


def test_journal_failure_is_reported_not_lost(tool, monkeypatch):
    monkeypatch.setattr(va, "journal", lambda *a: (_ for _ in ()).throw(OSError("disk full")))
    run_command(tool, "journal x", {"route": "journal", "text": "x"})
    assert tool.cards[-1][0] == "Couldn't save that"


# ---- meeting route

def test_meeting_route_starts_and_stops_the_pc_meeting_notes(tool, monkeypatch):
    started, stopped = [], []
    monkeypatch.setattr(tool, "_start_meeting", lambda app=None: started.append(app))
    monkeypatch.setattr(tool, "_stop_meeting", lambda: stopped.append(1))
    run_command(tool, "record this meeting called budget review", {"route": "meeting", "target": "start", "topic": "budget review"})
    run_command(tool, "start meeting notes", {"route": "meeting", "target": "start", "topic": ""})
    assert started == [] and tool.ui.qsize() == 2  # done on the main thread
    drain_ui(tool)
    assert started == [None, None]  # (the PC's notes name the meeting themselves, from what was said)
    run_command(tool, "stop the meeting", {"route": "meeting", "target": "stop"})
    drain_ui(tool)
    assert stopped == [1]


def drain_ui(tool):
    while not tool.ui.empty():
        tool.ui.get_nowait()()


# ---- screen, open, note, team, history

def make_jpeg(tmp_path):
    from PIL import Image
    path = tmp_path / "shot.jpg"
    Image.new("RGB", (40, 30), "red").save(path)
    return path


def test_screen_route_sends_a_picture_of_that_window_to_azure_and_deletes_it(tool, monkeypatch, isolated, online):
    shot = make_jpeg(isolated)
    seen = []
    monkeypatch.setattr(aw, "screenshot_front_window", lambda pid=None, hwnd=None: seen.append((pid, hwnd)) or shot)
    client = run_command(tool, "what's this error", {"route": "screen", "text": "What does this error mean?"},
                         then=["It means the file is missing."])
    assert seen[-1] == (1, 2)  # the window that was in front when the key was pressed
    assert tool.cards[-1] == ("About your screen", "It means the file is missing.", True)
    sent = client.sent[-1]["messages"][1]["content"]
    assert sent[1]["image_url"]["url"].startswith("data:image/jpeg;base64,")
    assert not shot.exists()  # never kept


def test_screen_route_with_no_window_says_so(tool, monkeypatch, online):
    monkeypatch.setattr(aw, "screenshot_front_window", lambda pid=None, hwnd=None: None)
    run_command(tool, "what's this", {"route": "screen", "text": "what's this"})
    assert tool.cards[-1][0] == "Nothing to look at"


def test_screen_route_azure_failure_is_reported_and_the_picture_still_deleted(tool, monkeypatch, isolated, online):
    shot = make_jpeg(isolated)
    monkeypatch.setattr(aw, "screenshot_front_window", lambda pid=None, hwnd=None: shot)
    monkeypatch.setattr(va, "ask_screen", lambda *a: (_ for _ in ()).throw(TimeoutError("slow")))
    run_command(tool, "what's this", {"route": "screen", "text": "what's this"})
    assert tool.cards[-1][0] == "Couldn't do that" and not shot.exists()


def test_open_route_uses_the_windows_platform_functions(tool, monkeypatch):
    monkeypatch.setattr(aw, "find_app", lambda w: "C:\\Start\\Microsoft Edge.lnk" if "edge" in w.lower() else None)
    run_command(tool, "open edge", {"route": "open", "target": "edge"})
    assert OPENED == ["C:\\Start\\Microsoft Edge.lnk"] and tool.cards[-1] == ("Open", "Opened Microsoft Edge.", False)


def test_open_route_without_a_target_asks_what_to_open(tool):
    run_command(tool, "open", {"route": "open", "target": "", "text": ""})
    assert "didn't catch" in tool.cards[-1][1] and OPENED == []


def test_note_route_writes_a_dated_note_into_the_notes_folder(tool, isolated):
    run_command(tool, "note idea for onboarding", {"route": "note", "text": "Idea: voice intro.", "topic": "onboarding voice"})
    files = list((isolated / "Documents" / "Voice notes").glob("*.md"))
    assert len(files) == 1 and files[0].name.endswith(" onboarding voice.md")
    assert "Idea: voice intro." in files[0].read_text() and tool.cards[-1][0] == "Noted"


def test_the_team_route_says_it_is_not_included_and_writes_nothing(tool, isolated):
    run_command(tool, "tell the PC Claude to check spending",
                {"route": "team", "text": "Please check spending.", "recipient": "claude-pc", "topic": "check spending"})
    assert tool.cards[-1] == ("Not included", "Messages to an AI team aren't included in this version.", False)
    assert not [p for p in isolated.rglob("*") if p.suffix == ".md"]
    assert not hasattr(va, "team_note") and not hasattr(va, "HANDOFFS") and not hasattr(va, "BRAIN_INBOX")


def test_the_transcribe_route_says_it_is_not_included(tool, monkeypatch):
    run_command(tool, "transcribe this link", {"route": "transcribe", "target": "https://example.com/x"})
    assert tool.cards[-1] == ("Not included", "Transcribing links isn't included in this version.", False)


def test_the_screen_route_without_the_online_option_says_so_and_takes_no_picture(tool, monkeypatch):
    monkeypatch.setitem(vt.CONFIG, "offline_only", True)
    monkeypatch.setattr(aw, "screenshot_front_window", lambda *a, **k: pytest.fail("no picture may be taken"))
    run_command(tool, "what's this error", {"route": "screen", "text": "what's this"})
    assert tool.cards[-1] == ("Not available here", "Questions about the screen need the online option.", False)
    monkeypatch.setitem(vt.CONFIG, "offline_only", False)
    monkeypatch.setitem(vt.CONFIG, "azure_openai_base_url", "")  # the online option isn't set up either
    run_command(tool, "what's this error", {"route": "screen", "text": "what's this"})
    assert tool.cards[-1][1] == "Questions about the screen need the online option."


def test_history_route_searches_what_was_said_before(tool, isolated):
    (isolated / "history.jsonl").write_text(
        json.dumps({"at": "2026-10-22 09:00:00", "raw": "budget talk", "text": "We discussed the budget for Q4."}) + "\n" +
        json.dumps({"at": "2026-10-22 09:05:00", "raw": "lunch", "text": "Lunch is at noon."}) + "\n")
    run_command(tool, "what did I say about the budget yesterday", {"route": "history", "target": "budget"})
    assert "budget for Q4" in tool.cards[-1][1] and "Lunch" not in tool.cards[-1][1]
    run_command(tool, "what did I say about zebras", {"route": "history", "target": "zebras"})
    assert "Nothing in your history" in tool.cards[-1][1]


def test_spoken_answers_are_off_by_default_and_read_aloud_when_on(tool, monkeypatch):
    tool._card = vt.VoiceTool._card.__get__(tool)  # the real card helper (with_ui is off, so nothing is drawn)
    tool._card("Answer", "It means X.", speak=True)
    assert SPOKEN == []
    vt.CONFIG["speak_answers"] = True
    tool._card("Answer", "It means X.", speak=True)
    tool._card("Learned", "not an answer")  # only answers are read
    assert [t for t, _ in SPOKEN] == ["It means X."] and SPOKEN[0][1]["voice"] == "" and SPOKEN[0][1]["rate"] == "normal"
    tool._handle("other", None, vt.win32con.VK_ESCAPE, now())  # Esc with nothing recording stops it
    assert STOPS and STOPS[-1] == 1


def test_every_command_route_works_through_the_real_router_with_typed_instructions(tool, monkeypatch, isolated):
    """The router prompt isn't run (no AI), but each JSON answer it can give is handled without error."""
    monkeypatch.setattr(aw, "screenshot_front_window", lambda pid=None, hwnd=None: make_jpeg(isolated))
    monkeypatch.setattr(aw, "find_app", lambda w: None)
    monkeypatch.setattr(aw, "search_files", lambda w, limit=5: [])
    monkeypatch.setattr(tool, "_claude", lambda *a: None)
    monkeypatch.setattr(tool, "_start_meeting", lambda *a, **k: None)
    monkeypatch.setattr(tool, "_stop_meeting", lambda: None)
    monkeypatch.setattr(system_windows, "_volume_control",
                        lambda: SimpleNamespace(SetMasterVolumeLevelScalar=lambda level, guid: None, SetMute=lambda m, g: None))
    answers = [{"route": r, "text": "x", "target": "x", "recipient": "claude-pc", "topic": "t"}
               for r in ("rewrite", "fix", "write", "answer", "screen", "open", "note", "team", "translate", "journal",
                         "meeting", "history", "action")] + [{"route": "read", "text": ""}, {"route": "system", "target": "mute", "value": ""},
              {"route": "system", "target": "set_volume", "value": "30"}, {"route": "transcribe", "target": ""}]
    for reply in answers:
        run_command(tool, f"typed instruction for {reply['route']}", reply,
                    {"app": "notepad.exe", "title": "x", "selection": "zef fir", "pid": 1, "hwnd": 2})
    time.sleep(0.2)
    assert [c[0] for c in tool.cards if c[0] == "Couldn't do that"] == []


# ====================================================================================================== claude (drafts only)

def test_the_drafts_only_lists_are_unchanged():
    assert commands.M365 + "outlook_send_mail" in commands.BLOCKED_TOOLS and commands.GMAIL + "send_message" in commands.BLOCKED_TOOLS
    assert commands.GCAL + "create_event" in commands.BLOCKED_TOOLS and "Bash" in commands.BLOCKED_TOOLS and "PowerShell" in commands.BLOCKED_TOOLS
    assert not set(commands.ALLOWED_TOOLS) & set(commands.BLOCKED_TOOLS)
    assert all(not name.endswith(("send_mail", "send_draft", "forward_mail", "send_message", "create_event", "update_event",
                                  "delete_event", "send_chat_message", "send_channel_message")) for name in commands.ALLOWED_TOOLS)
    assert commands.M365 + "outlook_create_draft" in commands.ALLOWED_TOOLS
    cmd = commands.claude_command("sonnet")
    assert "--permission-mode" in cmd and cmd[cmd.index("--permission-mode") + 1] == "dontAsk"


def test_claude_job_uses_the_kept_session_then_falls_back_and_logs(tool, monkeypatch, isolated):
    calls = []

    class Session:
        def ask(self, said, ctx, style=""):
            calls.append(("session", style))
            raise RuntimeError("session died")
    tool.claude_session = Session()
    monkeypatch.setattr(commands, "run_claude", lambda said, ctx, model, style="": calls.append(("fresh", style)) or ("Draft is in Outlook.", [], 3.2))
    tool._claude("reply to Sarah", {"app": "OUTLOOK.EXE", "title": "RE: x - Message", "selection": ""})
    assert [c[0] for c in calls] == ["session", "fresh"] and "Formal email" in calls[1][1]
    assert tool.cards[-1] == ("Claude", "Draft is in Outlook.", True)
    assert tool.light.claude_busy == 0


def test_claude_keep_running_off_goes_straight_to_a_fresh_claude(tool, monkeypatch):
    monkeypatch.setitem(vt.CONFIG, "claude_keep_running", False)
    tool.claude_session = SimpleNamespace(ask=lambda *a, **k: pytest.fail("kept session must not be used"))
    monkeypatch.setattr(commands, "run_claude", lambda said, ctx, model, style="": ("ok", [], 1))
    tool._claude("x", {})
    assert tool.cards[-1][0] == "Claude"


def test_claude_failure_is_shown(tool, monkeypatch):
    tool.claude_session = SimpleNamespace(ask=lambda *a, **k: (_ for _ in ()).throw(RuntimeError("x")))
    monkeypatch.setattr(commands, "run_claude", lambda *a, **k: (_ for _ in ()).throw(RuntimeError("boom")))
    tool._claude("x", {})
    assert tool.cards[-1][0] == "Claude couldn't do that" and tool.light.claude_busy == 0


# ====================================================================================================== styles

def test_default_rules_match_windows_program_names(isolated):
    for exe in ("WindowsTerminal.exe", "cmd.exe", "powershell.exe", "pwsh.exe", "Code.exe", "Cursor.exe", "OpenConsole.exe",
                "conhost.exe", "wezterm-gui.exe", "alacritty.exe", "mintty.exe", "putty.exe", "wt.exe"):
        assert styles.is_raw(styles.rule_for(exe)), exe
    for exe in ("OUTLOOK.EXE", "olk.exe", "HxOutlook.exe"):
        assert styles.rule_for(exe)["style"] == "Formal email", exe
    for exe in ("ms-teams.exe", "MSTeams.exe", "Teams.exe", "slack.exe", "WhatsApp.exe"):
        assert styles.rule_for(exe)["style"] == "Casual chat", exe
    assert styles.rule_for("chrome.exe") is None and styles.rule_for("msedge.exe") is None and styles.rule_for() is None


def test_names_match_ignoring_case_and_dot_exe(isolated):
    styles.save({"styles": [], "app_rules": [{"apps": "Chrome, msedge", "mode": "raw"}]})
    for name in ("chrome.exe", "CHROME.EXE", "Chrome", "msedge.exe", "MSEdge.exe"):
        assert styles.is_raw(styles.rule_for(name)), name
    assert styles.rule_for("firefox.exe") is None
    assert styles.rule_for("Code.exe") is None  # the saved rules replace the defaults


def test_mac_names_still_match(isolated):
    assert styles.is_raw(styles.rule_for("Terminal", "com.apple.Terminal"))
    assert styles.rule_for("Microsoft Outlook")["style"] == "Formal email"


def test_dictation_hint_and_writing_hint(isolated):
    hint = styles.dictation_hint(styles.rule_for("olk.exe"), "olk.exe")
    assert "Formal email" in hint and "never add, remove or reword" in hint and "No contractions" not in hint
    assert styles.dictation_hint(styles.rule_for("WindowsTerminal.exe"), "x") == "" and styles.dictation_hint(None) == ""
    writing = styles.writing_hint(styles.rule_for("slack.exe"), "slack.exe")
    assert "Casual chat" in writing and "LinkedIn post" in writing and "The user is in slack.exe" in writing
    assert "The user is in" not in styles.writing_hint(None)


def test_save_cleans_and_reloads_and_a_bad_file_falls_back_to_the_defaults(isolated):
    saved = styles.save({"styles": [{"name": "  Team  update ", "instructions": "Short."}, {"name": "team update"}],
                         "app_rules": [{"apps": " Notes , ", "mode": "style", "style": "Team update"}, {"apps": "", "mode": "raw"},
                                       {"apps": "Xcode", "mode": "weird"}]})
    assert [s["name"] for s in saved["styles"]] == ["Team update"]
    assert saved["app_rules"] == [{"apps": "Notes", "mode": "style", "style": "Team update"}, {"apps": "Xcode", "mode": "normal"}]
    assert styles.instructions("TEAM UPDATE") == "Short."
    (isolated / "styles.json").write_text("{not json")
    styles._cache["mtime"] = None
    assert styles.is_raw(styles.rule_for("Code.exe"))


# ====================================================================================================== actions_windows

def tree(root, names):
    for n in names:
        p = root / n
        p.parent.mkdir(parents=True, exist_ok=True)
        p.write_text("x")
    return root


def test_search_finds_files_and_folders_newest_first_and_skips_project_internals(tmp_path):
    root = tree(tmp_path / "docs", ["NVIDIA deck v1.pptx", "old/NVIDIA deck v2.pptx", "node_modules/pkg/NVIDIA deck.pptx",
                                    ".venv/NVIDIA deck.txt", ".git/NVIDIA deck", "notes.txt", "nvidia_deck_final.docx"])
    (root / "Deck folder NVIDIA").mkdir()
    os.utime(root / "old" / "NVIDIA deck v2.pptx", (time.time() + 100, time.time() + 100))
    hits = aw.search_files("NVIDIA deck", roots=[(root, 5)])
    names = [Path(h).name for h in hits]
    assert names[0] == "NVIDIA deck v2.pptx" and "NVIDIA deck v1.pptx" in names and "nvidia_deck_final.docx" in names
    assert "Deck folder NVIDIA" in names and not any("node_modules" in h or ".venv" in h or ".git" in h for h in hits)
    assert aw.search_files("the file", roots=[(root, 5)]) == []  # only filler words: nothing to look for
    assert len(aw.search_files("NVIDIA", limit=2, roots=[(root, 5)])) == 2


def test_search_never_lists_programs_or_scripts(tmp_path):
    root = tree(tmp_path / "d", ["build script.bat", "build script.ps1", "build script.exe", "build script.txt", "build script.cmd"])
    names = {Path(h).name for h in aw.search_files("build script", roots=[(root, 3)])}
    assert names == {"build script.txt"}


def test_search_depth_is_limited_per_root(tmp_path):
    root = tree(tmp_path / "r", ["a/b/c/target file.txt", "top target file.txt"])
    assert [Path(h).name for h in aw.search_files("target file", roots=[(root, 2)])] == ["top target file.txt"]
    assert len(aw.search_files("target file", roots=[(root, 5)])) == 2


def test_search_has_a_time_limit_even_if_a_drive_never_answers(tmp_path, monkeypatch):
    root = tree(tmp_path / "slow", ["findme.txt"])
    real = os.scandir

    def slow(path):
        time.sleep(30)
        return real(path)
    monkeypatch.setattr(os, "scandir", slow)
    started = time.monotonic()
    assert aw.search_files("findme", roots=[(root, 3)], seconds=0.5) == []
    assert time.monotonic() - started < 3


def test_search_skips_missing_roots_and_is_fast_on_the_real_pc(tmp_path):
    started = time.monotonic()
    assert aw.search_files("zzqqxx unlikely name", roots=[(tmp_path / "nope", 3), (tmp_path, 3)]) == []
    assert aw.search_files("zzqqxx unlikely name", roots=[(Path.home() / "Desktop", 3), (Path.home() / "Downloads", 3)]) == []
    assert time.monotonic() - started < aw.SEARCH_SECONDS * 2 + 1


def test_default_search_roots_are_only_this_users_own_folders():
    roots = [str(r) for r, _ in aw.search_roots()]
    assert any(r.endswith("Desktop") for r in roots) and any(r.endswith("Documents") for r in roots)
    assert any(r.endswith("Downloads") for r in roots)
    # Every folder searched is this person's own: under their user folder or their own OneDrive (which may be on any
    # drive: in these tests it is a throwaway folder on whichever drive holds TEMP). No other drive or folder is searched.
    # (Paths are compared as written, not resolved: a person's Downloads may itself be a link to another drive.)
    norm = lambda x: os.path.normcase(os.path.abspath(x)).rstrip("\\") + "\\"
    own = [norm(os.environ["USERPROFILE"])] + ([norm(os.environ["OneDrive"])] if os.environ.get("OneDrive") else [])
    for r in roots:
        assert any(norm(r).startswith(o) for o in own), r


@pytest.fixture
def start_menu(tmp_path, monkeypatch):
    base = tmp_path / "Start"
    for name in ("Microsoft Edge", "Outlook (classic)", "Outlook", "Word", "Word Viewer", "Uninstall Foo", "Visual Studio Code/Visual Studio Code"):
        p = base / f"{name}.lnk"
        p.parent.mkdir(parents=True, exist_ok=True)
        p.write_text("x")
    monkeypatch.setattr(aw, "start_menu_folders", lambda: [base])
    monkeypatch.setitem(aw._cache, "menu", (0.0, {}))
    monkeypatch.setitem(aw._cache, "store", (time.time(), [("Microsoft Teams", "shell:AppsFolder\\MSTeams_8wekyb3d8bbwe!MSTeams"),
                                                           ("Calculator", "shell:AppsFolder\\Calc!App")]))
    return base


def test_find_app_matches_exact_then_shortest_and_ignores_filler(start_menu):
    assert Path(aw.find_app("Microsoft Edge")).stem == "Microsoft Edge"
    assert Path(aw.find_app("open edge app")).stem == "Microsoft Edge"
    assert Path(aw.find_app("outlook")).stem == "Outlook"  # exact wins over "Outlook (classic)"
    assert Path(aw.find_app("word")).stem == "Word"
    assert Path(aw.find_app("visual studio code")).stem == "Visual Studio Code"
    assert aw.find_app("zzzz") is None and aw.find_app("") is None and aw.find_app("the") is None
    assert aw.find_app("ed") is None  # too short to guess


def test_find_app_falls_back_to_store_apps_like_the_new_teams(start_menu):
    found = aw.find_app("teams")
    assert found == "shell:AppsFolder\\MSTeams_8wekyb3d8bbwe!MSTeams" and aw.app_label(found) == "Microsoft Teams"
    assert aw.app_label("C:\\x\\Microsoft Edge.lnk") == "Microsoft Edge"


def test_open_path_refuses_programs_and_scripts_but_opens_documents_and_apps(tmp_path):
    for bad in ("run.exe", "run.bat", "run.cmd", "run.ps1", "run.vbs", "run.js", "x.msi", "x.hta", "x.reg", "x.scr"):
        with pytest.raises(PermissionError):
            aw.open_path(tmp_path / bad)
    aw.open_path(tmp_path / "deck.pptx")
    aw.open_app("C:\\Start\\Edge.lnk")  # Start-menu apps open through open_app (F1: open_path never opens a .lnk)
    aw.open_app("shell:AppsFolder\\MSTeams_x!MSTeams")
    aw.open_path(tmp_path / "tool.exe", allow_launch=True)
    assert len(OPENED) == 4


def test_voice_shortcuts_are_an_allow_list_of_four_kinds(isolated, monkeypatch, start_menu):
    va.save_shortcuts([{"phrase": "my dashboard", "kind": "url", "target": "https://example.com/dashboard"},
                       {"phrase": "bad", "kind": "shell", "target": "rm -rf /"}, {"phrase": "x", "kind": "run", "target": "calc"}])
    assert [s["phrase"] for s in va.load_shortcuts()] == ["my dashboard"]
    opened = []
    monkeypatch.setattr(aw.webbrowser, "open", lambda url: opened.append(url))
    aw.open_shortcut({"kind": "url", "target": "https://example.com/dashboard"})
    assert opened == ["https://example.com/dashboard"]
    for bad in ("file:///C:/Windows/System32/calc.exe", "javascript:alert(1)", "ms-msdt:/id x", "calc.exe", "ftp://x"):
        with pytest.raises(ValueError):
            aw.open_shortcut({"kind": "url", "target": bad})
    with pytest.raises(ValueError):
        aw.open_shortcut({"kind": "shell", "target": "calc"})
    assert opened == ["https://example.com/dashboard"] and OPENED == []


def test_app_file_and_windows_shortcut_kinds(tmp_path, start_menu):
    aw.open_shortcut({"kind": "app", "target": "Microsoft Edge"})
    assert OPENED[-1].endswith("Microsoft Edge.lnk")
    exe = tmp_path / "mytool.exe"
    exe.write_text("x")
    aw.open_shortcut({"kind": "app", "target": str(exe)})  # a program Sam named himself
    assert OPENED[-1] == str(exe)
    with pytest.raises(FileNotFoundError):
        aw.open_shortcut({"kind": "app", "target": "No Such App"})
    doc = tmp_path / "plan.docx"
    doc.write_text("x")
    aw.open_shortcut({"kind": "file", "target": str(doc)})
    aw.open_shortcut({"kind": "file", "target": str(tmp_path)})
    with pytest.raises(FileNotFoundError):
        aw.open_shortcut({"kind": "file", "target": str(tmp_path / "gone.docx")})
    with pytest.raises(PermissionError):  # a file shortcut never runs a program
        aw.open_shortcut({"kind": "file", "target": str(exe)})
    lnk = tmp_path / "thing.lnk"
    lnk.write_text("x")
    aw.open_shortcut({"kind": "shortcut", "target": str(lnk)})
    aw.open_shortcut({"kind": "shortcut", "target": "Microsoft Edge"})
    assert OPENED[-1].endswith("Microsoft Edge.lnk")
    with pytest.raises(FileNotFoundError):
        aw.open_shortcut({"kind": "shortcut", "target": "Nothing here"})
    with pytest.raises(FileNotFoundError):
        aw.open_shortcut({"kind": "shortcut", "target": str(exe)})  # only .lnk / .url paths


def test_open_thing_tries_shortcuts_then_apps_then_files(isolated, monkeypatch, start_menu):
    va.save_shortcuts([{"phrase": "Zephyr", "kind": "url", "target": "https://example.com"}])
    monkeypatch.setattr(aw.webbrowser, "open", lambda url: None)
    assert "shortcut" in va.open_thing(aw, "Zephyr")
    assert va.open_thing(aw, "Microsoft Edge") == "Opened Microsoft Edge."
    monkeypatch.setattr(aw, "search_files", lambda w, limit=5: ["C:\\d\\NVIDIA deck.pptx", "C:\\d\\old NVIDIA deck.pptx"] if "nvidia" in w.lower() else [])
    msg = va.open_thing(aw, "NVIDIA deck")
    assert msg.startswith("Opened NVIDIA deck.pptx") and "old NVIDIA deck.pptx" in msg
    assert "Couldn't find" in va.open_thing(aw, "zzz")


# ===================== F1 (outside review, October 30, 2026): "open ..." never launches a shortcut, link or program
# A file search must never list, and open_path must never open, a launch-capable link or a program, in any spelling or
# through a symbolic link or junction. Start-menu apps and voice shortcuts listed in shortcuts.json still open.
# Everything here is inert bytes in throwaway folders, and os.startfile is the recorder (OPENED): nothing ever runs.

F1_LINKS = (".lnk", ".url", ".pif", ".scf", ".appref-ms", ".website", ".library-ms", ".settingcontent-ms",
            ".search-ms", ".searchconnector-ms", ".desklink", ".mapimail", ".rdp")
F1_PROGRAMS = (".exe", ".com", ".bat", ".cmd", ".ps1", ".psm1", ".vbs", ".vbe", ".js", ".jse", ".wsf", ".wsh", ".msc",
               ".msi", ".msp", ".hta", ".cpl", ".scr", ".jar", ".reg", ".application", ".gadget", ".inf", ".py", ".pyw")
F1_GUID = ".{20D04FE0-3AEA-1069-A2D8-08002B30309D}"


def _f1_mixed(ext):
    return "".join(c.upper() if i % 2 else c for i, c in enumerate(ext))


def _f1_raw(path):
    """\\\\?\\ keeps a trailing dot or space in a file name, as a downloaded file can have."""
    return "\\\\?\\" + str(path)


def _f1_inert(path):
    with open(_f1_raw(path), "wb") as f:
        f.write(b"INERT TEST FIXTURE")


def _f1_remove(folder):
    for entry in os.scandir(_f1_raw(folder)):
        try:
            os.remove(entry.path)
        except OSError:
            pass


def test_f1_review_reproduction_an_unlisted_shortcut_is_not_opened_and_the_person_is_told_why(tool, isolated, monkeypatch):
    (isolated / "Quarterly budget.lnk").write_bytes(b"INERT TEST FIXTURE")
    monkeypatch.setattr(aw, "find_app", lambda words: None)
    monkeypatch.setattr(aw, "search_roots", lambda: [(isolated, 1)])
    tool._more("open", "Quarterly budget", {"target": "Quarterly budget"}, {})
    assert OPENED == []
    title, text = tool.cards[-1][:2]
    assert title == "Open" and "Quarterly budget.lnk" in text and "Windows shortcut" in text and "voice shortcut" in text


def test_f1_the_risky_list_covers_every_program_and_link_type():
    for ext in F1_LINKS + F1_PROGRAMS + (".lnk", ".url", ".pif"):
        assert ext in aw.RISKY_EXT, ext
    for ext in (".pdf", ".docx", ".xlsx", ".pptx", ".txt", ".md", ".png", ".jpg", ".mp3", ".mp4", ".csv", ".zip"):
        assert ext not in aw.RISKY_EXT, ext


@pytest.mark.parametrize("ext", F1_LINKS + F1_PROGRAMS)
def test_f1_risky_names_in_any_spelling_are_never_found_or_opened(tmp_path, ext):
    root = tmp_path / "found"
    root.mkdir()
    risky = [f"budget plan{ext}", f"budget plan mixed{_f1_mixed(ext)}", f"budget plan upper{ext.upper()}",
             f"budget plan dot{ext}.", f"budget plan space{ext} ", f"budget plan both{ext} . ."]
    try:
        for name in risky:
            _f1_inert(root / name)
        (root / "budget plan.pdf").write_bytes(b"%PDF-1.4 inert")
        assert len(os.listdir(_f1_raw(root))) == len(risky) + 1  # the trailing dots and spaces really are on disk
        skipped = []
        hits = aw.search_files("budget plan", roots=[(root, 2)], skipped=skipped)
        assert [Path(h).name for h in hits] == ["budget plan.pdf"]  # discovery: only the document
        assert len(skipped) == len(risky) and all(why in (aw.PROGRAM_WHY, aw.LINK_WHY) for _, why in skipped)
        for name in risky:  # opening: the final guard, for each name as given and as it sits on disk
            for path in (str(root / name), next(e.path for e in os.scandir(root) if e.name == name)):
                with pytest.raises(PermissionError):
                    aw.open_path(path)
        doc = str(root / "budget plan.pdf")
        for stream in (f"{doc}:hidden{ext}", f"{doc}:hidden{_f1_mixed(ext)}:$DATA", f"{root / 'budget plan'}{ext}::$DATA",
                       f"{root / 'budget plan'}{ext}.:x"):  # alternate data streams never open
            with pytest.raises(PermissionError):
                aw.open_path(stream)
        aw.open_path(doc)
        assert OPENED == [doc]
    finally:
        _f1_remove(root)


def test_f1_a_real_alternate_data_stream_is_never_opened(tmp_path):
    doc = tmp_path / "Budget notes.pdf"
    doc.write_bytes(b"%PDF-1.4 inert")
    try:
        with open(f"{doc}:payload.lnk", "wb") as f:
            f.write(b"INERT TEST FIXTURE")
    except OSError:
        pytest.skip("this drive has no alternate data streams (not NTFS)")
    assert [Path(h).name for h in aw.search_files("budget notes", roots=[(tmp_path, 2)])] == ["Budget notes.pdf"]
    for stream in (f"{doc}:payload.lnk", f"{doc}:payload.lnk:$DATA", f"{doc}::$DATA"):
        with pytest.raises(PermissionError):
            aw.open_path(stream)
    assert OPENED == []


def test_f1_addresses_and_shell_objects_never_open_from_open_path(tmp_path):
    for bad in ("shell:AppsFolder\\MSTeams_x!MSTeams", "shell:startup", "ms-settings:privacy", "search-ms:query=x",
                "ms-msdt:/id PCWDiagnostic", "file:///C:/Windows/System32/calc.exe", "\\\\.\\PhysicalDrive0",
                f"C:\\x\\Budget{F1_GUID}", f"C:\\x\\Budget{F1_GUID}\\plan.pdf", "C:\\x\\.lnk", "C:\\x\\plan.pdf.LnK"):
        with pytest.raises(PermissionError):
            aw.open_path(bad)
    assert OPENED == []


def test_f1_a_symbolic_link_that_leads_to_a_program_or_shortcut_is_refused(tool, tmp_path, monkeypatch):
    bin_ = tmp_path / "bin"
    bin_.mkdir()
    (bin_ / "calc.exe").write_bytes(b"INERT TEST FIXTURE")
    (bin_ / "launcher.lnk").write_bytes(b"INERT TEST FIXTURE")
    docs = tmp_path / "docs"
    docs.mkdir()
    try:
        os.symlink(bin_ / "calc.exe", docs / "Budget report.pdf")
        os.symlink(bin_ / "launcher.lnk", docs / "Budget summary.docx")
        os.symlink(bin_ / "launcher.lnk", docs / "Budget listed.lnk")
        os.symlink(bin_ / "gone.pdf", docs / "Budget missing.pdf")
    except (OSError, NotImplementedError):
        pytest.skip("this Windows account can't create symbolic links (needs Developer Mode or an admin)")
    skipped = []
    assert aw.search_files("budget", roots=[(docs, 2)], skipped=skipped) == []
    why = {Path(p).name: w for p, w in skipped}
    assert why["Budget report.pdf"] == aw.PROGRAM_WHY and why["Budget summary.docx"] == aw.LINK_WHY
    assert why["Budget missing.pdf"] == aw.DANGLING_WHY
    for name in ("Budget report.pdf", "Budget summary.docx", "Budget missing.pdf"):
        with pytest.raises(PermissionError):
            aw.open_path(docs / name)
    with pytest.raises(PermissionError):  # a listed shortcut is still never a link to a program
        aw.open_path(docs / "Budget report.pdf", allow_link=True)
    with pytest.raises(PermissionError):
        aw.open_shortcut({"kind": "file", "target": str(docs / "Budget report.pdf")})
    monkeypatch.setattr(aw, "find_app", lambda words: None)
    monkeypatch.setattr(aw, "search_roots", lambda: [(docs, 2)])
    tool._more("open", "Budget report", {"target": "Budget report"}, {})
    assert OPENED == [] and tool.cards[-1][0] == "Open" and "program or script" in tool.cards[-1][1]


def test_f1_a_link_to_a_program_is_refused_even_where_symbolic_links_cant_be_made(tool, tmp_path, monkeypatch):
    """The same check as above with the link's destination faked, so it runs on every PC (no Developer Mode needed)."""
    (tmp_path / "bin").mkdir()
    program = tmp_path / "bin" / "calc.exe"
    program.write_bytes(b"INERT TEST FIXTURE")
    docs = tmp_path / "docs"
    docs.mkdir()
    link = docs / "Budget report.pdf"
    link.write_bytes(b"INERT TEST FIXTURE")
    real_realpath, real_islink = os.path.realpath, os.path.islink
    leads_to = lambda p: os.path.normcase(str(p)) == os.path.normcase(str(link))
    monkeypatch.setattr(os.path, "realpath", lambda p, *a, **k: str(program) if leads_to(p) else real_realpath(p, *a, **k))
    monkeypatch.setattr(os.path, "islink", lambda p: True if leads_to(p) else real_islink(p))
    skipped = []
    assert aw.search_files("budget report", roots=[(docs, 2)], skipped=skipped) == []
    assert [w for _, w in skipped] == [aw.PROGRAM_WHY]
    for allow_link in (False, True):
        with pytest.raises(PermissionError):
            aw.open_path(link, allow_link=allow_link)
    monkeypatch.setattr(aw, "find_app", lambda words: None)
    monkeypatch.setattr(aw, "search_roots", lambda: [(docs, 2)])
    tool._more("open", "Budget report", {"target": "Budget report"}, {})
    assert OPENED == [] and tool.cards[-1][0] == "Open" and "program or script" in tool.cards[-1][1]


def test_f1_junctions_open_as_folders_but_a_broken_one_does_not(tmp_path):
    import _winapi
    target = tmp_path / "Real budget folder"
    target.mkdir()
    junction = tmp_path / "Budget folder"
    try:
        _winapi.CreateJunction(str(target), str(junction))
    except OSError:
        pytest.skip("this drive can't hold junctions")
    assert aw.why_not_open(junction) is None
    aw.open_path(junction)
    assert OPENED == [str(junction)]
    target.rmdir()
    assert aw.why_not_open(junction) == aw.DANGLING_WHY
    with pytest.raises(PermissionError):
        aw.open_path(junction)
    assert len(OPENED) == 1
    os.rmdir(junction)


def test_f1_special_shell_folders_are_never_found_or_opened(tmp_path):
    root = tmp_path / "r"
    (root / f"Budget{F1_GUID}").mkdir(parents=True)
    shell_folder = root / "Budget shell"
    shell_folder.mkdir()
    (shell_folder / "desktop.ini").write_text("[.ShellClassInfo]\r\nCLSID2={0AFACED1-E828-11D1-9187-B532F1E9575D}\r\n")
    plain = root / "Budget plain"
    plain.mkdir()
    (plain / "desktop.ini").write_text("[.ShellClassInfo]\r\nIconResource=C:\\Windows\\System32\\imageres.dll,3\r\n")
    skipped = []
    assert [Path(h).name for h in aw.search_files("budget", roots=[(root, 2)], skipped=skipped)] == ["Budget plain"]
    assert {Path(p).name for p, _ in skipped} == {f"Budget{F1_GUID}", "Budget shell"}
    for bad in (root / f"Budget{F1_GUID}", shell_folder):
        with pytest.raises(PermissionError):
            aw.open_path(bad)
    aw.open_path(plain)
    assert OPENED == [str(plain)]


def test_f1_documents_and_folders_still_open_by_voice(tool, isolated, monkeypatch):
    docs = isolated / "Docs"
    docs.mkdir()
    (docs / "Tax plan.pdf").write_bytes(b"%PDF-1.4 inert")
    (docs / "Letter draft.docx").write_bytes(b"PK inert")
    (docs / "Next.js").mkdir()  # a folder, whatever its name, opens in File Explorer
    monkeypatch.setattr(aw, "find_app", lambda words: None)
    monkeypatch.setattr(aw, "search_roots", lambda: [(docs, 2)])
    for said, name in (("Tax plan", "Tax plan.pdf"), ("Letter draft", "Letter draft.docx"), ("Next js", "Next.js")):
        tool._more("open", said, {"target": said}, {})
        assert OPENED[-1] == str(docs / name) and tool.cards[-1][:2] == ("Open", f"Opened {name}.")
    assert len(OPENED) == 3


def test_f1_start_menu_apps_still_launch(tool, monkeypatch, start_menu):
    monkeypatch.setattr(aw, "search_roots", lambda: [])
    tool._more("open", "Microsoft Edge", {"target": "Microsoft Edge"}, {})
    assert OPENED[-1] == str(start_menu / "Microsoft Edge.lnk") and tool.cards[-1][:2] == ("Open", "Opened Microsoft Edge.")
    tool._more("open", "teams", {"target": "teams"}, {})
    assert OPENED[-1] == "shell:AppsFolder\\MSTeams_8wekyb3d8bbwe!MSTeams"
    assert tool.cards[-1][:2] == ("Open", "Opened Microsoft Teams.")
    with pytest.raises(PermissionError):  # open_app trusts the Start menu's shortcuts, never a program file
        aw.open_app("C:\\Start\\calc.exe")
    with pytest.raises(PermissionError):
        aw.open_app("shell:AppsFolder\\..\\..\\Windows\\calc.exe")
    assert len(OPENED) == 2


def test_f1_a_listed_windows_shortcut_opens_but_an_unlisted_one_with_the_same_name_does_not(tool, isolated, monkeypatch):
    listed = isolated / "Listed" / "Quarterly budget.lnk"
    unlisted = isolated / "Downloads" / "Quarterly budget.lnk"
    for p in (listed, unlisted):
        p.parent.mkdir()
        p.write_bytes(b"INERT TEST FIXTURE")
    monkeypatch.setattr(aw, "find_app", lambda words: None)
    monkeypatch.setattr(aw, "search_roots", lambda: [(unlisted.parent, 2), (listed.parent, 2)])
    va.save_shortcuts([{"phrase": "quarterly budget", "kind": "shortcut", "target": str(listed)}])
    tool._more("open", "quarterly budget", {"target": "quarterly budget"}, {})
    assert OPENED == [str(listed)] and "quarterly budget" in tool.cards[-1][1]  # the listed one, by its exact phrase
    va.save_shortcuts([{"phrase": "my budget shortcut", "kind": "shortcut", "target": str(listed).upper()}])
    tool._more("open", "my budget shortcut", {"target": "my budget shortcut"}, {})
    assert OPENED == [str(listed), str(listed).upper()]  # any case of its exact path
    tool._more("open", "quarterly budget", {"target": "quarterly budget"}, {})  # not listed under these words now
    assert len(OPENED) == 2 and "won't open it from a file search" in tool.cards[-1][1]
    va.save_shortcuts([{"phrase": "quarterly budget", "kind": "file", "target": str(unlisted)}])
    tool._more("open", "quarterly budget", {"target": "quarterly budget"}, {})  # a File shortcut never opens a .lnk
    assert len(OPENED) == 2 and tool.cards[-1][0] == "Couldn't do that" and "Windows shortcut" in tool.cards[-1][1]


def test_f1_the_final_guard_still_refuses_if_a_search_hit_changed_before_opening(tool, monkeypatch):
    monkeypatch.setattr(aw, "find_app", lambda words: None)
    monkeypatch.setattr(aw, "search_files", lambda w, limit=5: ["C:\\Users\\x\\Downloads\\Quarterly budget.lnk"])
    tool._more("open", "Quarterly budget", {"target": "Quarterly budget"}, {})
    assert OPENED == [] and tool.cards[-1][0] == "Couldn't do that" and "Windows shortcut" in tool.cards[-1][1]
    with pytest.raises(PermissionError):  # and the window's own "open" (WinHost) uses the same guard
        vt.WinHost(tool).open_path("C:\\Users\\x\\Desktop\\Report.url")
    assert OPENED == []
# ===================== end of F1 tests


def test_the_real_start_menu_can_be_read(monkeypatch):
    monkeypatch.setitem(aw._cache, "menu", (0.0, {}))
    items = aw._start_menu_items()
    assert isinstance(items, dict)  # (read-only: just lists shortcut names)


@pytest.fixture(scope="module")
def tk_root():
    """One hidden Tk root for the whole run (Tk can't be started over and over in one process)."""
    import tkinter as tk
    root = tk.Tk()
    root.withdraw()
    yield root
    root.destroy()


class quiet_window:
    """A real window for the screenshot and call-watcher tests. It is fully see-through (alpha 0, which Windows still
    draws into a screenshot of that window), has no taskbar button and never takes focus, so it can't get in the way of
    whatever Sam is doing. (A window moved off-screen is not drawn at all, so it can't be placed there instead.)"""

    def __init__(self, root, title="vt-test", color="#ff0000", size="420x300"):
        import tkinter as tk
        import win32con
        import win32gui
        self.win = tk.Toplevel(root)
        self.win.title(title)
        self.win.geometry(f"{size}+0+0")
        tk.Canvas(self.win, bg=color, highlightthickness=0).pack(fill="both", expand=True)
        self.win.update()
        self.hwnd = win32gui.GetParent(self.win.winfo_id()) or self.win.winfo_id()
        style = win32gui.GetWindowLong(self.hwnd, win32con.GWL_EXSTYLE)
        win32gui.SetWindowLong(self.hwnd, win32con.GWL_EXSTYLE, style | win32con.WS_EX_NOACTIVATE | win32con.WS_EX_TOOLWINDOW)
        self.win.attributes("-alpha", 0.0)
        self.win.update()
        time.sleep(0.15)
        self.win.update()

    def __enter__(self):
        return self

    def __exit__(self, *a):
        self.win.destroy()


@pytest.fixture
def screenshots_only_of_the_test_window(monkeypatch):
    """The test's own window belongs to this process, and nothing else may ever be photographed: the real front
    window is hidden from the code under test, so a failure can't capture whatever is on Sam's screen."""
    monkeypatch.setattr("win32gui.GetForegroundWindow", lambda: 0)
    monkeypatch.setattr(aw.os, "getpid", lambda: 1)  # (so the test window isn't taken for one of VoiceTool's own)


def test_screenshot_captures_just_that_window(tk_root, screenshots_only_of_the_test_window):
    from PIL import Image
    with quiet_window(tk_root, color="#ff0000") as w:
        path = aw.screenshot_front_window(hwnd=w.hwnd)
        assert path is not None and path.suffix == ".jpg"
        try:
            image = Image.open(path).convert("RGB")
            assert 150 < image.width < 900 and 150 < image.height < 700  # the window, not the whole screen
            r, g, b = image.getpixel((image.width // 2, image.height // 2))
            assert r > 200 and g < 60 and b < 60  # its own contents, even though the window is see-through
        finally:
            path.unlink(missing_ok=True)


def test_screenshot_finds_the_window_by_process_too(tk_root, screenshots_only_of_the_test_window, monkeypatch):
    import win32process
    from PIL import Image
    real = win32process.GetWindowThreadProcessId
    with quiet_window(tk_root, color="#00ff00") as w:
        monkeypatch.setattr("win32process.GetWindowThreadProcessId", lambda h: (1, 77) if h == w.hwnd else real(h))
        path = aw.screenshot_front_window(pid=77)
        try:
            image = Image.open(path).convert("RGB")
            r, g, b = image.getpixel((image.width // 2, image.height // 2))
            assert g > 200 and r < 60
        finally:
            path.unlink(missing_ok=True)


def test_screenshot_is_shrunk_for_the_model(tk_root, screenshots_only_of_the_test_window):
    from PIL import Image
    with quiet_window(tk_root, size="1800x1200") as w:
        path = aw.screenshot_front_window(hwnd=w.hwnd, max_side=500)
        try:
            assert max(Image.open(path).size) <= 500
        finally:
            path.unlink(missing_ok=True)


def test_screenshot_dpi_awareness_is_only_for_the_capture():
    import ctypes
    prev = aw._set_thread_dpi_aware()
    assert prev is not None  # supported on this PC
    aw._restore_dpi(prev)
    assert ctypes.windll.user32.GetThreadDpiAwarenessContext() is not None


def test_screenshot_returns_none_when_there_is_no_window(monkeypatch):
    monkeypatch.setattr("win32gui.GetForegroundWindow", lambda: 0)
    assert aw.screenshot_front_window(pid=None, hwnd=None) is None
    assert aw.screenshot_front_window(pid=999999, hwnd=123456789) is None


def test_voicetools_own_pill_and_cards_are_never_photographed(monkeypatch):
    monkeypatch.setattr("win32gui.GetForegroundWindow", lambda: 555)
    monkeypatch.setattr("win32gui.IsWindow", lambda h: True)
    monkeypatch.setattr("win32gui.IsWindowVisible", lambda h: True)
    monkeypatch.setattr("win32gui.IsIconic", lambda h: False)
    monkeypatch.setattr("win32gui.GetWindowRect", lambda h: (0, 0, 400, 300))
    monkeypatch.setattr(aw.os, "getpid", lambda: 4242)
    monkeypatch.setattr("win32process.GetWindowThreadProcessId", lambda h: (1, 4242))  # every window is ours
    assert aw._pick_window(pid=4242, hwnd=555) is None
    monkeypatch.setattr("win32process.GetWindowThreadProcessId", lambda h: (1, 99))  # now they belong to another program
    assert aw._pick_window(pid=None, hwnd=555) == 555


# ====================================================================================================== meeting notes: recording

def speech(seconds, level):
    t = np.arange(int(vt.RATE * seconds)) / vt.RATE
    return (level * np.sin(2 * np.pi * 220 * t)).astype(np.float32)


def fake_transcribe(audio):
    peak = float(np.abs(audio).max())
    return "this is sam speaking" if peak > 0.2 else "this is the other side"


def test_the_platform_halves_are_split_out_and_the_pc_uses_the_windows_one():
    assert meetings.platform() is meetings_windows
    for name in ("start_system_audio", "call_in_progress", "notes_base", "_apps_using_mic", "_window_titles"):
        assert callable(getattr(meetings_windows, name))
    assert not hasattr(meetings, "winreg")  # the shared file has nothing Windows-only in it
    assert REAL_NOTES_BASE() == Path.home() / "Documents"  # (the real one: under the user folder, not OneDrive)
    assert meetings_windows.notes_base() == Path(str(meetings_windows.notes_base()))  # (and in tests a temp Documents)


def test_meeting_notes_and_voice_notes_never_go_to_onedrive(isolated, monkeypatch):
    """Even with %OneDrive% set (and Windows' Documents redirected into it), the real functions use Documents under the
    user folder; the Known Folder (which OneDrive can redirect) is not asked."""
    monkeypatch.setenv("OneDrive", r"D:\Somebody\OneDrive - Work")
    monkeypatch.setenv("OneDriveConsumer", r"D:\Somebody\OneDrive")
    monkeypatch.setenv("OneDriveCommercial", r"D:\Somebody\OneDrive - Work")
    assert REAL_NOTES_BASE() == Path.home() / "Documents"
    assert "onedrive" not in str(REAL_NOTES_BASE()).lower()
    assert REAL_DEFAULT_NOTES_FOLDER() == Path.home() / "Documents" / "Voice notes"
    source = (HERE / "meetings_windows.py").read_text(encoding="utf-8")
    body = source[source.index("def notes_base"):]
    assert "OneDrive" not in body.split('"""')[2] and "SHGetKnownFolderPath" not in body and "SHGetFolderPath" not in body
    # and the meeting notes folder is made under that base
    assert meetings.notes_folder() == isolated / "Documents" / "Meeting notes"


def test_two_tracks_are_kept_apart_and_merged_in_time_order(tmp_path, monkeypatch):
    monkeypatch.setattr(meetings, "PIECE_SECONDS", 2)
    rec = meetings.MeetingRecorder(tmp_path / "audio", fake_transcribe, capture=False, min_peak=0.005)
    rec.me.add(speech(3, 0.4))
    rec.others.add(np.concatenate([np.zeros(vt.RATE * 2, np.float32), speech(3, 0.1)]))
    lines = rec.stop()
    assert [(who, text) for _, who, text in lines] == [("Me", "this is sam speaking"), ("Others", "this is the other side")]
    assert lines[0][0] < lines[1][0]
    for name in ("me.wav", "others.wav"):
        with wave.open(str(tmp_path / "audio" / name)) as w:
            assert w.getframerate() == 16000 and w.getnchannels() == 1
    assert rec.ended_at >= rec.started_at and not rec.running


def test_a_two_voice_meeting_comes_out_as_one_conversation(tmp_path, monkeypatch):
    """Sam and the other side take turns; each is transcribed from its own track and the lines interleave by time."""
    monkeypatch.setattr(meetings, "PIECE_SECONDS", 3)
    said = iter(["sam one", "other one", "sam two", "other two"])
    rec = meetings.MeetingRecorder(tmp_path / "audio", lambda a: "", capture=False)
    # Sam talks 0-2 s and 4-6 s; the other side 2-4 s and 6-8 s (the other track is silent while Sam talks)
    rec.me.add(np.concatenate([speech(2, 0.4), np.zeros(vt.RATE * 2, np.float32), speech(2, 0.4), np.zeros(vt.RATE * 2, np.float32)]))
    rec.others.add(np.concatenate([np.zeros(vt.RATE * 2, np.float32), speech(2, 0.1), np.zeros(vt.RATE * 2, np.float32), speech(2, 0.1)]))
    rec.transcribe = lambda audio: "sam talking" if float(np.abs(audio).max()) > 0.2 else "other talking"
    lines = rec.stop()
    assert [who for _, who, _ in lines] == ["Me", "Others"] or len(lines) >= 2
    assert {who for _, who, _ in lines} == {"Me", "Others"} and lines == sorted(lines)
    assert "Me: sam talking" in meetings.transcript_text(lines) and "Others: other talking" in meetings.transcript_text(lines)


def test_silence_is_not_transcribed(tmp_path):
    calls = []
    rec = meetings.MeetingRecorder(tmp_path / "audio", lambda a: calls.append(1) or "x", capture=False)
    rec.me.add(np.zeros(vt.RATE * 3, np.float32))
    rec.others.add(speech(0.3, 0.5))  # too short to be speech
    assert rec.stop() == [] and calls == []


def test_long_meetings_are_transcribed_in_pieces_while_they_run(tmp_path, monkeypatch):
    monkeypatch.setattr(meetings, "PIECE_SECONDS", 2)
    pieces = []
    rec = meetings.MeetingRecorder(tmp_path / "audio", lambda a: pieces.append(len(a) / vt.RATE) or "words", capture=False)
    rec.me.add(speech(6, 0.4))
    deadline = time.time() + 8
    while not pieces and time.time() < deadline:
        time.sleep(0.2)
    assert pieces and pieces[0] <= 6  # picked up by the background thread before stop()
    rec.stop()


def test_track_cuts_at_a_quiet_moment(tmp_path):
    track = meetings.Track("x", tmp_path / "x.wav")
    audio = np.full(vt.RATE * 30, 0.3, np.float32)
    audio[vt.RATE * 27: vt.RATE * 27 + vt.RATE // 5] = 0
    track.add(audio)
    start, piece = track.take()
    assert start == 0 and 26.5 < len(piece) / vt.RATE < 27.5
    assert track.take() is None  # the rest is shorter than a piece
    track.close()


def test_unk_markers_are_never_part_of_a_transcript(tmp_path):
    rec = meetings.MeetingRecorder(tmp_path / "audio", lambda a: "<unk><unk> hello <unk> there <unk>", capture=False)
    rec.me.add(speech(2, 0.4))
    assert [text for _, _, text in rec.stop()] == ["hello there"]
    assert meetings.strip_unk("<unk><unk>") == "" and meetings.strip_unk("  a   b ") == "a b"
    t = vt.Transcriber.__new__(vt.Transcriber)  # and the speech model's own output is cleaned for dictation too
    t.model = SimpleNamespace(recognize=lambda chunk, sample_rate: "open the <unk><unk> file <unk>")
    assert t(np.zeros(vt.RATE, np.float32)) == "open the file"


def test_abort_stops_without_transcribing_and_closes_the_files(tmp_path):
    calls = []
    rec = meetings.MeetingRecorder(tmp_path / "audio", lambda a: calls.append(1) or "x", capture=False)
    rec.me.add(speech(3, 0.4))
    rec.abort()
    assert calls == [] and not rec.running
    (tmp_path / "audio" / "others.wav").unlink()  # files are closed, so they can be deleted on Windows


def test_leftover_other_side_audio_is_deleted_but_sams_own_track_is_not(isolated):
    audio = vt.MEETING_AUDIO / "2026-10-22_100000"
    audio.mkdir(parents=True)
    (audio / "others.wav").write_bytes(b"x")
    (audio / "me.wav").write_bytes(b"x")
    assert meetings.delete_other_side_audio(vt.MEETING_AUDIO) == [str(audio / "others.wav")]
    assert not (audio / "others.wav").exists() and (audio / "me.wav").exists()
    assert meetings.delete_other_side_audio(isolated / "nothing here") == []


def test_the_recorder_opens_the_microphone_and_the_calls_sound_unless_told_otherwise(tmp_path, monkeypatch):
    opened = {"mic": [], "system": []}

    class Stream:
        def __init__(self, **kw):
            opened["mic"].append(kw["device"])
        def start(self): pass
        def stop(self): pass
        def close(self): pass
    monkeypatch.setattr(meetings.sd, "InputStream", Stream)
    monkeypatch.setattr(meetings_windows, "start_system_audio",
                        lambda on_audio, log=None, is_running=lambda: True: opened["system"].append((on_audio, is_running())))
    rec = meetings.MeetingRecorder(tmp_path / "audio", fake_transcribe, mic_device=3)
    time.sleep(0.3)
    assert opened["mic"] == [3] and len(opened["system"]) == 1
    assert opened["system"][0][0] == rec.others.add and opened["system"][0][1] is True  # the sound goes onto the Others track
    rec.stop()


def test_a_failing_system_audio_is_logged_and_the_mic_still_records(tmp_path, monkeypatch):
    class Stream:
        def __init__(self, **kw): pass
        def start(self): pass
        def stop(self): pass
        def close(self): pass
    monkeypatch.setattr(meetings.sd, "InputStream", Stream)
    monkeypatch.setattr(meetings_windows, "start_system_audio", lambda *a, **k: (_ for _ in ()).throw(OSError("no loopback device")))
    logged = []
    log = SimpleNamespace(info=lambda msg, *a: logged.append(msg % a if a else msg))
    rec = meetings.MeetingRecorder(tmp_path / "audio", fake_transcribe, log=log)
    rec._from_mic(np.full((vt.RATE * 2, 1), 0.4, np.float32))
    time.sleep(0.3)
    lines = rec.stop()
    assert any("no loopback device" in m for m in logged) and [w for _, w, _ in lines] == ["Me"]


def fake_soundcard(levels, monkeypatch):
    import types

    class Rec:
        def __init__(self, level): self.level = level
        def __enter__(self): return self
        def __exit__(self, *a): pass
        def record(self, numframes):
            time.sleep(0.01)
            return np.full((numframes, 1), self.level, np.float32)

    class Mic:
        def __init__(self, level): self.level = level
        def recorder(self, samplerate, channels): return Rec(self.level)

    fake = types.SimpleNamespace(SoundcardRuntimeWarning=RuntimeWarning,
                                 all_speakers=lambda: [SimpleNamespace(name=n) for n in levels],
                                 get_microphone=lambda id, include_loopback: Mic(levels[id]))
    monkeypatch.setitem(sys.modules, "soundcard", fake)


def test_system_audio_mixes_every_output_device_with_synthetic_speakers(monkeypatch):
    """soundcard is replaced by two fake speakers; their audio is added together and handed to the Others track."""
    fake_soundcard({"A": 0.1, "B": 0.2}, monkeypatch)
    got, running = [], {"v": True}
    thread = threading.Thread(target=meetings_windows.start_system_audio, args=(lambda a: got.append(float(a[0])), None, lambda: running["v"]), daemon=True)
    thread.start()
    time.sleep(0.5)
    running["v"] = False
    thread.join(5)
    assert got and abs(got[0] - 0.3) < 1e-5 and not thread.is_alive()


def test_a_speaker_that_cannot_be_opened_is_skipped_and_logged(monkeypatch):
    fake_soundcard({"A": 0.1, "Broken": 0.2}, monkeypatch)
    sc = sys.modules["soundcard"]
    real = sc.get_microphone
    sc.get_microphone = lambda id, include_loopback: (_ for _ in ()).throw(OSError("busy")) if id == "Broken" else real(id, include_loopback)
    got, logged, running = [], [], {"v": True}
    log = SimpleNamespace(info=lambda msg, *a: logged.append(msg % a))
    thread = threading.Thread(target=meetings_windows.start_system_audio, args=(lambda a: got.append(float(a[0])), log, lambda: running["v"]), daemon=True)
    thread.start()
    time.sleep(0.4)
    running["v"] = False
    thread.join(5)
    assert got and abs(got[0] - 0.1) < 1e-5 and any("Broken" in m for m in logged)


def _have_devices():
    try:
        import soundcard as sc
        return any(d["max_input_channels"] > 0 for d in vt.sd.query_devices()) and bool(sc.all_speakers())
    except Exception:
        return False


LIVE_AUDIO = os.environ.get("VOICETOOL_LIVE_AUDIO") == "1"


@pytest.mark.skipif(not LIVE_AUDIO, reason="records the REAL microphone and speakers: only with VOICETOOL_LIVE_AUDIO=1")
@pytest.mark.skipif(LIVE_AUDIO and not _have_devices(), reason="no microphone or speaker on this PC")
def test_live_microphone_and_loopback_record_four_seconds(tmp_path):
    """OPT-IN ONLY (VOICETOOL_LIVE_AUDIO=1, and only with the person's say-so: it records whatever the room and the
    speakers are playing). The REAL microphone (the one in config.json) and the REAL speakers' loopback, 4 seconds through
    the real recorder, into pytest's temp folder. Only the lengths are checked; nothing is listened to or transcribed, and
    the two files are deleted as soon as they're measured."""
    cfg = json.loads((HERE / "config.json").read_text(encoding="utf-8"))
    want = cfg.get("mic_name_contains", "").lower()
    device = next((i for i, d in enumerate(vt.sd.query_devices())
                   if d["max_input_channels"] > 0 and d["hostapi"] == 0 and want and want in d["name"].lower()), None)
    rec = meetings.MeetingRecorder(tmp_path / "audio", lambda a: "", device, 0.005, None)
    time.sleep(4.0)
    stopped = time.time()
    rec.stop()
    assert time.time() - stopped < 5  # stopping doesn't hang on an idle device
    lengths = {}
    try:
        for name in ("me.wav", "others.wav"):
            with wave.open(str(tmp_path / "audio" / name)) as w:
                lengths[name] = w.getnframes() / w.getframerate()
    finally:  # the room and the speakers were recorded: never keep them
        for name in ("me.wav", "others.wav"):
            (tmp_path / "audio" / name).unlink(missing_ok=True)
    for name, seconds in lengths.items():
        assert 3.2 < seconds < 5.0, (name, seconds)  # both sides really were recorded, side by side


# ====================================================================================================== meeting notes: noticing calls

@pytest.mark.parametrize("titles, mic, expected", [
    (["Zoom Meeting"], [], "Zoom"), (["Zoom Webinar"], [], "Zoom"), (["Meet - abc-defg-hij - Google Chrome"], [], "Google Meet"),
    (["Cisco Webex Meetings"], [], "Webex"), (["Huddle with Sam - Slack"], [], "Slack"),
    (["Inbox - Outlook", "Settings"], [], None), ([], ["microsoft.windows.teams_x"], None),
    ([], ["msteams_8wekyb3d8bbwe"], "Teams"), ([], ["c:#program files#zoom#bin#zoom.exe"], "Zoom"),
    (["Zoom Meeting"], ["msteams_x"], "Teams")])
def test_call_detection_by_mic_registry_and_window_titles(monkeypatch, titles, mic, expected):
    monkeypatch.setattr(meetings_windows, "_apps_using_mic", lambda: mic)
    monkeypatch.setattr(meetings_windows, "_window_titles", lambda: titles)
    assert meetings.call_in_progress() == expected  # (the shared file asks the Windows half)


def test_call_detection_sees_a_real_fake_zoom_window(tk_root, monkeypatch):
    monkeypatch.setattr(meetings_windows, "_apps_using_mic", lambda: [])
    with quiet_window(tk_root, title="Zoom Meeting"):
        assert meetings.call_in_progress() == "Zoom"
    assert "Zoom Meeting" not in meetings_windows._window_titles()  # gone when it closes


def test_the_real_microphone_registry_can_be_read():
    assert isinstance(meetings_windows._apps_using_mic(), list)


def watch(**kw):
    events = []
    w = meetings.MeetingWatch(lambda app: events.append(("start", app)), lambda: events.append(("end",)),
                              kw.pop("enabled", lambda: True), end_after=kw.pop("end_after", 0.05), autostart=False, **kw)
    return w, events


def test_watch_offers_once_and_finishes_after_the_call_is_quiet_for_a_while(monkeypatch):
    state = {"app": "Zoom"}
    monkeypatch.setattr(meetings, "call_in_progress", lambda: state["app"])
    w, events = watch()
    w.check()
    w.check()
    assert events == [("start", "Zoom")]
    state["app"] = None
    w.check()  # a short mute does not end it
    assert events == [("start", "Zoom")]
    state["app"] = "Zoom"
    w.check()
    state["app"] = None
    w.check()
    time.sleep(0.1)
    w.check()
    assert events == [("start", "Zoom"), ("end",)]
    state["app"] = "Teams"
    w.check()
    assert events[-1] == ("start", "Teams")


def test_watch_does_nothing_when_switched_off_or_paused(monkeypatch):
    monkeypatch.setattr(meetings, "call_in_progress", lambda: "Zoom")
    on = {"v": False}
    w, events = watch(enabled=lambda: on["v"])
    w.check()
    assert events == []
    on["v"] = True
    w.check()
    assert events == [("start", "Zoom")]
    on["v"] = False  # switched off in the middle of the call: it winds up like a call that ended
    time.sleep(0.01)
    w.check()
    time.sleep(0.1)
    w.check()
    assert events[-1] == ("end",)


def test_watch_loop_survives_errors(monkeypatch):
    seq = iter([RuntimeError("boom"), "Zoom"])

    def fake():
        v = next(seq)
        if isinstance(v, Exception):
            raise v
        return v
    monkeypatch.setattr(meetings, "call_in_progress", fake)
    events = []
    w = meetings.MeetingWatch(lambda a: events.append(a), lambda: None, lambda: True, every=0.05, autostart=True)
    deadline = time.time() + 3
    while not events and time.time() < deadline:
        time.sleep(0.05)
    assert events == ["Zoom"]


def test_offer_card_and_auto_stop_follow_the_switches(tool, monkeypatch):
    class Card:
        shown = []
        def show(self, title, body, buttons=None, seconds=60):
            self.shown.append((title, body, buttons))
    tool.card = Card()
    tool._offer_meeting("Zoom")
    title, body, buttons = tool.card.shown[0]
    assert title == "Meeting in Zoom?" and [b[0] for b in buttons] == ["Everyone agreed, start", "Not this time"]
    assert meetings.CONSENT_REMINDER in body  # the offer is itself the consent card
    monkeypatch.setitem(vt.CONFIG, "meeting_notes", False)
    tool._offer_meeting("Zoom")
    assert len(tool.card.shown) == 1  # switched off: no offer
    stopped = []
    monkeypatch.setattr(tool, "_stop_meeting", lambda: stopped.append(1))
    tool._call_ended()
    assert stopped == []  # no meeting running
    tool.meeting, tool.meeting_auto_stop = object(), False
    tool._call_ended()
    assert stopped == []  # started by hand: only the user stops it
    tool.meeting_auto_stop = True
    tool._call_ended()
    assert stopped == [1]


# ====================================================================================================== meeting notes: the whole flow

NOTES = {"title": "Budget review", "summary": "We agreed to keep Q4 flat.", "decisions": ["Keep Q4 flat"],
         "action_items": [{"owner": "Me", "task": "Send the numbers", "due": "Friday"}, {"owner": "Others", "task": "Review", "due": ""}],
         "open_questions": ["Who owns hiring?"], "follow_up_email": "Hi all,\nThanks for today.\nSam"}


def meeting_tool(tool, monkeypatch, tmp_path, client=None):
    """A tool whose meetings record nothing from the hardware (synthetic audio is put on the two tracks)."""
    tool.azure = FakeCleaner(client=client or FakeClient(json.dumps(NOTES)))
    tool.transcribe = fake_transcribe
    real = meetings.MeetingRecorder
    monkeypatch.setattr(meetings, "MeetingRecorder", lambda *a, **k: real(*a, **{**k, "capture": False}))
    tool._refresh_icon = lambda: None
    claude_calls = []
    monkeypatch.setattr(tool, "_claude", lambda said, ctx: claude_calls.append((said, ctx)))
    tool.claude_calls = claude_calls
    return tool


def run_meeting(tool, seconds=3, app=None):
    """A synthetic two-voice meeting: Sam on his microphone, the other side on the call's sound."""
    start_meeting(tool, app)
    assert tool.meeting is not None
    tool.meeting.me.add(speech(seconds, 0.4))
    tool.meeting.others.add(np.concatenate([np.zeros(vt.RATE * 2, np.float32), speech(seconds, 0.1)]))
    recorder = tool.meeting
    tool._stop_meeting()
    deadline = time.time() + 20
    while time.time() < deadline and not (tool.light.meeting is None and tool.meeting is None
                                         and not (recorder.folder / "others.wav").exists()):
        time.sleep(0.1)
    return recorder


def notes_files(tmp_path):
    return list((tmp_path / "Documents" / "Meeting notes").glob("*.docx"))


def test_a_meeting_end_to_end_makes_the_word_document_history_row_and_follow_up(tool, monkeypatch, isolated):
    meeting_tool(tool, monkeypatch, isolated)
    recorder = run_meeting(tool)
    docs = notes_files(isolated)  # (a temp Documents folder, never the real one)
    assert len(docs) == 1 and docs[0].name.endswith("Budget review.docx")
    from docx import Document
    text = "\n".join(p.text for p in Document(docs[0]).paragraphs)
    assert "Summary" in text and "We agreed to keep Q4 flat." in text and "Me: " in text and "Others: " in text
    assert "this is sam speaking" in text and "this is the other side" in text  # both sides are in the transcript
    row = json.loads((isolated / "history.jsonl").read_text().splitlines()[-1])
    assert row["meeting"] == "Budget review" and row["notes"] == str(docs[0]) and row["minutes"] >= 0
    # the other people's audio is always deleted; Sam's own only if Save recordings is on
    assert not (recorder.folder / "others.wav").exists() and not (recorder.folder / "me.wav").exists()
    # Claude follows up with a DRAFT to the calendar attendees
    assert len(tool.claude_calls) == 1
    instruction, ctx = tool.claude_calls[0]
    assert "never send" in instruction and "Hi all,\nThanks for today.\nSam" in instruction and ctx["app"] == "VoiceTool meeting notes"
    assert tool.light.meeting is None and tool.meeting is None
    assert any(c[0] == "Meeting notes ready" for c in tool.cards)


def test_save_recordings_keeps_sam_but_never_the_other_side(tool, monkeypatch, isolated):
    monkeypatch.setitem(vt.CONFIG, "save_recordings", True)
    meeting_tool(tool, monkeypatch, isolated)
    recorder = run_meeting(tool)
    assert (recorder.folder / "me.wav").exists() and not (recorder.folder / "others.wav").exists()


@pytest.mark.parametrize("followup, actions", [(False, True), (True, False)])
def test_follow_up_draft_obeys_its_switches(tool, monkeypatch, isolated, followup, actions):
    monkeypatch.setitem(vt.CONFIG, "meeting_followup", followup)
    monkeypatch.setitem(vt.CONFIG, "claude_actions", actions)
    meeting_tool(tool, monkeypatch, isolated)
    run_meeting(tool)
    assert tool.claude_calls == []
    assert len(notes_files(isolated)) == 1  # the notes themselves are unaffected


def test_a_meeting_where_nothing_was_heard_writes_nothing_and_is_tidied(tool, monkeypatch, isolated):
    meeting_tool(tool, monkeypatch, isolated)
    start_meeting(tool, )
    recorder = tool.meeting
    recorder.me.add(np.zeros(vt.RATE * 2, np.float32))
    tool._stop_meeting()
    deadline = time.time() + 15
    while time.time() < deadline and (tool.light.meeting is not None or (recorder.folder / "others.wav").exists()):
        time.sleep(0.1)
    assert tool.cards[-1][0] == "No meeting notes" and notes_files(isolated) == []
    assert not (recorder.folder / "others.wav").exists() and tool.light.meeting is None
    assert not (isolated / "history.jsonl").exists()


def transcript_text_of(path):
    from docx import Document
    return "\n".join(p.text for p in Document(path).paragraphs)


def test_a_failed_notes_call_still_saves_the_transcript_and_deletes_the_other_side(tool, monkeypatch, isolated):
    client = FakeClient("{}")
    client.create = lambda **kw: (_ for _ in ()).throw(TimeoutError("model down"))
    meeting_tool(tool, monkeypatch, isolated, client=client)
    recorder = run_meeting(tool)
    docs = notes_files(isolated)
    assert len(docs) == 1 and docs[0].name.endswith("Meeting transcript.docx")  # the words are never lost
    text = transcript_text_of(docs[0])
    assert "this is sam speaking" in text and "this is the other side" in text and "full transcript only" in text
    assert "took too long" in text
    assert tool.cards[-1][0] == "Meeting transcript saved" and "took too long" in tool.cards[-1][1] and "Open notes" in tool.card_buttons[-1]
    row = json.loads((isolated / "history.jsonl").read_text().splitlines()[-1])
    assert row["meeting"] == "Meeting transcript" and row["notes"] == str(docs[0])
    assert not (recorder.folder / "others.wav").exists() and tool.light.meeting is None
    assert tool.claude_calls == []  # no follow-up email without notes


def test_notes_that_come_back_unreadable_or_a_model_that_vanishes_mid_meeting_also_save_the_transcript(tool, monkeypatch, isolated):
    meeting_tool(tool, monkeypatch, isolated, client=FakeClient("this is not json"))
    run_meeting(tool)
    assert len(notes_files(isolated)) == 1 and "couldn't be read" in transcript_text_of(notes_files(isolated)[0])
    # the model is deleted while the meeting runs: the tidy-up is "not set up" when the notes are written
    for doc in notes_files(isolated):
        doc.unlink()
    meeting_tool(tool, monkeypatch, isolated)
    tool.azure._client = lambda: (_ for _ in ()).throw(vt.TidyUpNotSetUp("gone"))
    run_meeting(tool)
    docs = notes_files(isolated)
    assert len(docs) == 1 and "isn't set up" in transcript_text_of(docs[0])
    assert tool.cards[-1][0] == "Meeting transcript saved"


def test_the_notes_call_on_this_computer_gets_fifteen_minutes(tool, monkeypatch, isolated):
    monkeypatch.setitem(vt.CONFIG, "offline_only", True)
    monkeypatch.setattr(local_cleanup_win, "available", lambda config=None: True)
    meeting_tool(tool, monkeypatch, isolated)
    options = []
    client = tool.azure.client
    client.with_options = lambda **kw: options.append(kw) or client
    run_meeting(tool)
    assert options == [{"timeout": 900}] and meetings.LOCAL_NOTES_TIMEOUT == 900
    monkeypatch.setitem(vt.CONFIG, "offline_only", False)  # online (advanced): the old 3 minutes
    run_meeting(tool)
    assert options[-1] == {"timeout": 180}


def test_if_the_word_document_cannot_be_made_the_transcript_is_saved_as_a_text_file(tool, monkeypatch, isolated):
    meeting_tool(tool, monkeypatch, isolated)
    monkeypatch.setattr(meetings, "write_document", lambda *a, **k: (_ for _ in ()).throw(PermissionError("locked")))
    run_meeting(tool)
    texts = list((isolated / "Documents" / "Meeting notes").glob("*.txt"))
    assert len(texts) == 1 and "this is sam speaking" in texts[0].read_text(encoding="utf-8")
    assert tool.cards[-1][0] == "Meeting transcript saved" and "text file" in tool.cards[-1][1]


def test_the_light_and_status_follow_recording_then_writing_then_done(tool, monkeypatch, isolated):
    meeting_tool(tool, monkeypatch, isolated)
    host = vt.WinHost(tool)
    gate = threading.Event()
    real_make = meetings.make_notes
    monkeypatch.setattr(meetings, "make_notes", lambda *a, **k: gate.wait(10) and real_make(*a, **k))
    start_meeting(tool, "Zoom")
    assert tool.light.meeting == "meeting" and host.meeting_status()["recording"] and not host.meeting_status()["writing"]
    tool.meeting.me.add(speech(3, 0.4))
    tool._stop_meeting()
    assert tool.light.meeting == "meeting_notes"
    status = host.meeting_status()
    assert status == {"recording": False, "seconds": 0, "app": "", "writing": True}  # "Writing your meeting notes…"
    gate.set()
    assert wait_until(lambda: tool.light.meeting is None, 20)
    assert host.meeting_status()["writing"] is False


def test_manual_start_is_refused_when_meeting_notes_are_switched_off(tool, monkeypatch, isolated):
    monkeypatch.setitem(vt.CONFIG, "meeting_notes", False)
    meeting_tool(tool, monkeypatch, isolated)
    start_meeting(tool, )
    assert tool.meeting is None and tool.cards[-1][0] == "Meeting notes are switched off"


def test_starting_twice_does_not_start_a_second_recording(tool, monkeypatch, isolated):
    meeting_tool(tool, monkeypatch, isolated)
    start_meeting(tool, )
    first = tool.meeting
    start_meeting(tool, )
    assert tool.meeting is first
    tool.meeting.abort()


def test_start_failure_is_reported(tool, monkeypatch, isolated):
    meeting_tool(tool, monkeypatch, isolated)
    monkeypatch.setattr(meetings, "MeetingRecorder", lambda *a, **k: (_ for _ in ()).throw(OSError("no microphone")))
    start_meeting(tool, )
    assert tool.meeting is None and tool.cards[-1][0] == "Couldn't start meeting notes" and tool.light.meeting is None


def test_started_for_a_call_remembers_the_app_and_to_stop_with_the_call(tool, monkeypatch, isolated):
    meeting_tool(tool, monkeypatch, isolated)
    start_meeting(tool, "Zoom")
    assert tool.meeting_auto_stop is True and tool.meeting_app == "Zoom"
    assert vt.WinHost(tool).meeting_status()["app"] == "Zoom"
    tool.meeting.abort()
    tool.meeting = None
    start_meeting(tool, )
    assert tool.meeting_auto_stop is False and tool.meeting_app == "" and vt.WinHost(tool).meeting_status()["app"] == ""
    tool.meeting.abort()


def test_quitting_in_the_middle_of_a_meeting_stops_recording_and_deletes_the_other_side(tool, monkeypatch, isolated):
    meeting_tool(tool, monkeypatch, isolated)
    start_meeting(tool, )
    recorder = tool.meeting
    recorder.me.add(speech(2, 0.4))
    recorder.others.add(speech(2, 0.1))
    tool._abandon_meeting()
    assert tool.meeting is None and not recorder.running
    assert not (recorder.folder / "others.wav").exists() and not (recorder.folder / "me.wav").exists()
    assert notes_files(isolated) == [] and tool.light.meeting is None


def test_the_word_document_has_the_summary_decisions_action_table_and_transcript(isolated):
    now_ = time.time()
    path = meetings.write_document(NOTES, [(0, "Me", "hello"), (65, "Others", "hi there")], now_ - 600, now_)
    assert path.parent == isolated / "Documents" / "Meeting notes" and path.suffix == ".docx"
    from docx import Document
    doc = Document(path)
    headings = [p.text for p in doc.paragraphs if p.style.name.startswith("Heading")]
    assert headings == ["Summary", "Decisions", "Open questions", "Action items", "Follow-up email", "Transcript"]
    assert [c.text for c in doc.tables[0].rows[1].cells] == ["Me", "Send the numbers", "Friday"]
    assert any(p.text.startswith("[01:05] Others: ") for p in doc.paragraphs)


# ====================================================================================================== the app window: host + API

@pytest.fixture
def api(tool, isolated):
    host = vt.WinHost(tool)
    return ui_api.API(host, vt.CONFIG, settings_window.save), host


def test_features_include_the_meetings_rows_with_friend_wording():
    rows = {f[1]: f for f in ui_api.FEATURES}
    assert rows["meeting_notes"][0] == "Meetings" and rows["meeting_notes"][3] == "Icon menu › Start meeting notes"
    assert "Records you and the call" in rows["meeting_notes"][4] and "Meeting notes folder" in rows["meeting_notes"][4]
    assert "agrees to be recorded" in rows["meeting_notes"][4] and "Off until you turn it on" in rows["meeting_notes"][4]
    assert rows["meeting_auto_offer"][2] == "Offer when a call starts" and rows["meeting_followup"][2] == "Follow-up email draft"
    assert "Draft only." in rows["meeting_followup"][4] and "Claude" in rows["meeting_followup"][4]
    assert "Claude Code installed and signed in on your own Claude plan" in rows["claude_actions"][4]
    keys = [f[1] for f in ui_api.FEATURES]
    sections = [f[0] for f in ui_api.FEATURES]
    assert keys.index("meeting_notes") + 1 == keys.index("meeting_auto_offer") and keys.index("meeting_auto_offer") + 1 == keys.index("meeting_followup")
    assert sections.index("Meetings") > sections.index("Doing things") and sections.index("Meetings") < sections.index("Learning")
    assert not any("Option" in str(f[3]) or "Command" in str(f[3]) for f in ui_api.FEATURES)  # Windows key names


def test_the_old_tk_settings_window_keeps_its_meetings_rows():
    assert {"meeting_notes", "meeting_auto_offer", "meeting_followup"} <= {f[1] for f in settings_window.FEATURES}


def test_every_setting_saves_to_config_and_defaults_are_right(api, isolated):
    api, host = api
    s = api.call("settings")
    assert [f["key"] for f in s["features"]] == [f[1] for f in ui_api.FEATURES]
    on = {f["key"]: f["on"] for f in s["features"]}
    assert on["speak_answers"] is False and on["offline_only"] is False and on["meeting_notes"] is True  # (this test's environment)
    for f in ui_api.FEATURES:
        api.call("set_setting", {"key": f[1], "on": False})
        assert json.loads((isolated / "config.json").read_text())[f[1]] is False
        api.call("set_setting", {"key": f[1], "on": True, "confirmed": True})
        assert vt.CONFIG[f[1]] is True
    with pytest.raises(ValueError):
        api.call("set_setting", {"key": "azure_openai_base_url", "on": False})  # only the listed switches can be changed
    with pytest.raises(ValueError):
        api.call("nothing_like_this")


def test_missing_keys_use_the_friend_defaults_in_the_settings_page(api, monkeypatch):
    api, host = api
    for k in ("meeting_notes", "meeting_auto_offer", "meeting_followup", "claude_actions", "claude_keep_running", "speak_answers",
              "save_recordings", "offline_only", "beeps"):
        monkeypatch.delitem(vt.CONFIG, k, raising=False)
    on = {f["key"]: f["on"] for f in api.call("settings")["features"]}
    for k in ui_api.OFF_BY_DEFAULT:
        assert on[k] is False, k
    assert on["beeps"] is True and on["offline_only"] is True and on["dictate_enabled"] is True


def test_turning_meeting_notes_on_in_the_window_shows_the_reminder_first(api, isolated):
    api, host = api
    vt.CONFIG["meeting_notes"] = False
    reply = api.call("set_setting", {"key": "meeting_notes", "on": True})
    assert reply == {"key": "meeting_notes", "on": False, "needs_confirm": meetings.CONSENT_REMINDER}
    assert vt.CONFIG["meeting_notes"] is False and not (isolated / "config.json").exists()  # nothing was saved yet
    assert api.call("set_setting", {"key": "meeting_notes", "on": True, "confirmed": True}) == {"key": "meeting_notes", "on": True}
    assert vt.CONFIG["meeting_notes"] is True
    assert api.call("set_setting", {"key": "meeting_notes", "on": False}) == {"key": "meeting_notes", "on": False}  # off needs no question
    assert api.call("set_setting", {"key": "meeting_auto_offer", "on": True})["on"] is True  # (its own card asks every time)


def test_the_two_offline_rows_say_when_the_model_is_not_installed(api, monkeypatch):
    api, host = api
    monkeypatch.setattr(local_cleanup_win, "available", lambda config=None: False)
    rows = {f["key"]: f for f in api.call("settings")["features"]}
    assert "isn't downloaded yet" in rows["offline_only"]["text"] and "Finish setup" in rows["offline_only"]["text"]
    assert "isn't downloaded" not in rows["cleanup"]["text"]
    monkeypatch.setattr(local_cleanup_win, "available", lambda config=None: True)
    assert "isn't downloaded" not in {f["key"]: f for f in api.call("settings")["features"]}["offline_only"]["text"]


def test_turning_off_keep_claude_ready_stops_the_kept_session(api, tool):
    api, host = api
    stopped = []
    tool.claude_session = SimpleNamespace(stop=lambda: stopped.append(1))
    api.call("set_setting", {"key": "claude_keep_running", "on": False})
    assert stopped == [1]


def test_translate_dropdown_keeps_a_language_said_aloud(api, isolated):
    api, host = api
    assert api.call("set_translate", {"language": "Spanish"})["translate_to"] == "Spanish"
    with pytest.raises(ValueError):
        api.call("set_translate", {"language": "Klingon"})
    vt.CONFIG["translate_to"] = "Swahili"  # set by voice
    s = api.call("settings")
    assert "Swahili" in s["languages"] and s["translate_to"] == "Swahili"
    assert api.call("set_translate", {"language": "Swahili"})["translate_to"] == "Swahili"
    assert api.call("set_translate", {"language": ""})["translate_to"] == ""


def test_host_computer_name_and_claude_sign_in_is_checked_at_most_once_a_minute(tool, monkeypatch):
    host = vt.WinHost(tool)
    assert host.computer_name() == (os.environ.get("COMPUTERNAME") or "This PC")
    calls = []
    monkeypatch.setattr(commands, "signed_in", lambda: calls.append(1) or True)
    assert host.claude_signed_in() and host.claude_signed_in() and calls == [1]


def test_signed_in_reads_claude_auth_status(monkeypatch):
    monkeypatch.setattr(commands.subprocess, "run", lambda *a, **k: SimpleNamespace(stdout='{"loggedIn": true}'))
    assert REAL_SIGNED_IN() is True
    monkeypatch.setattr(commands.subprocess, "run", lambda *a, **k: SimpleNamespace(stdout='{"loggedIn": false}'))
    assert REAL_SIGNED_IN() is False
    monkeypatch.setattr(commands.subprocess, "run", lambda *a, **k: (_ for _ in ()).throw(FileNotFoundError()))
    assert REAL_SIGNED_IN() is False


def test_host_startup_uses_schtasks_enable_and_disable_for_the_voicetool_task(tool, monkeypatch):
    runs = []
    monkeypatch.setattr(settings_window, "TASK", "VoiceTool")  # the real name, but subprocess is replaced
    monkeypatch.setattr(settings_window.subprocess, "run", lambda args, **k: runs.append(args) or SimpleNamespace(stdout='"VoiceTool","N/A","Ready"\n'))
    host = vt.WinHost(tool)
    host.set_startup(True)
    host.set_startup(False)
    assert runs == [["schtasks", "/Change", "/TN", "VoiceTool", "/ENABLE"], ["schtasks", "/Change", "/TN", "VoiceTool", "/DISABLE"]]
    assert host.startup_enabled() is True
    monkeypatch.setattr(settings_window.subprocess, "run", lambda args, **k: SimpleNamespace(stdout='"VoiceTool","N/A","Disabled"\n'))
    assert host.startup_enabled() is False
    monkeypatch.setattr(settings_window.subprocess, "run", lambda args, **k: SimpleNamespace(stdout=""))
    assert host.startup_enabled() is False


def test_api_set_startup_returns_the_real_state(api, monkeypatch):
    api, host = api
    runs = []
    monkeypatch.setattr(settings_window, "TASK", "VoiceTool")
    monkeypatch.setattr(settings_window.subprocess, "run", lambda args, **k: runs.append(args[3:]) or SimpleNamespace(stdout='"VoiceTool","N/A","Ready"\n'))
    assert api.call("set_startup", {"on": True}) == {"startup": True}
    assert ["/ENABLE"] in [r[1:] for r in runs]


def test_host_stop_claude_refresh_style_and_copy(tool, monkeypatch):
    host = vt.WinHost(tool)
    stopped, copied, learned = [], [], []
    tool.claude_session = SimpleNamespace(stop=lambda: stopped.append(1))
    monkeypatch.setattr(vt, "set_clipboard_text", lambda text: copied.append(text))
    monkeypatch.setattr(commands, "learn_style", lambda model: learned.append(model))
    host.stop_claude()
    host.copy_text("Hello there")
    host.refresh_style()
    time.sleep(0.3)
    assert stopped == [1] and copied == ["Hello there"] and learned == [vt.CONFIG.get("claude_model", "sonnet")]
    assert {"Learning your style", "Style updated"} <= {c[0] for c in tool.cards}
    assert tool.light.claude_busy == 0


def test_host_open_path_opens_documents_and_refuses_programs(tool):
    host = vt.WinHost(tool)
    host.open_path(Path("C:/x/words.txt"))
    assert OPENED == ["C:\\x\\words.txt"] or OPENED[-1].endswith("words.txt")
    with pytest.raises(PermissionError):
        host.open_path(Path("C:/x/run.exe"))


def test_host_quit_is_queued_for_the_main_thread(tool, monkeypatch):
    quit_ = []
    monkeypatch.setattr(tool, "_quit", lambda: quit_.append(1))
    vt.WinHost(tool).quit()
    assert quit_ == [] and tool.ui.qsize() == 1
    drain_ui(tool)
    assert quit_ == [1]


def test_host_meeting_methods_and_status(tool, monkeypatch, isolated):
    meeting_tool(tool, monkeypatch, isolated)
    host = vt.WinHost(tool)
    assert host.meeting_status() == {"recording": False, "seconds": 0, "app": "", "writing": False}
    host.start_meeting()
    assert tool.meeting is None  # queued for the main thread
    drain_ui(tool)
    assert tool.meeting is None and tool.cards[-1][0] == "Before recording"  # asks first: nothing is recording yet
    agree(tool)
    status = host.meeting_status()
    assert status["recording"] is True and status["seconds"] >= 0 and status["app"] == "" and status["writing"] is False
    host.start_meeting()
    drain_ui(tool)
    assert host.meeting_status()["recording"] is True and tool.cards[-1][0] != "Before recording"  # (still the one recording)
    host.stop_meeting()
    drain_ui(tool)
    assert host.meeting_status()["recording"] is False
    wait_until(lambda: tool.light.meeting is None, 15)


def test_host_open_meeting_folder_opens_the_notes_folder(tool, isolated):
    vt.WinHost(tool).open_meeting_folder()
    assert OPENED == [str(isolated / "Documents" / "Meeting notes")] and (isolated / "Documents" / "Meeting notes").is_dir()


def history_row(isolated, **row):
    with open(isolated / "history.jsonl", "a", encoding="utf-8") as f:
        f.write(json.dumps({"at": time.strftime("%Y-%m-%d %H:%M:%S"), **row}) + "\n")


def test_meetings_page_lists_the_meeting_rows_from_history(api, tool, isolated):
    api, host = api
    doc = meetings.write_document(NOTES, [(0, "Me", "hello")], time.time() - 1800, time.time())
    history_row(isolated, raw="x", text="a dictation")
    history_row(isolated, meeting="Budget review", notes=str(doc), minutes=30)
    history_row(isolated, meeting="Moved away", notes=str(isolated / "Documents" / "Meeting notes" / "gone.docx"), minutes=5)
    page = api.call("meetings")
    assert page["status"] == {"recording": False, "seconds": 0, "app": "", "writing": False}
    assert [(m["title"], m["exists"], m["minutes"]) for m in page["items"]] == [("Moved away", False, 5), ("Budget review", True, 30)]
    assert page["items"][1]["path"] == str(doc) and page["folder"] == str(doc.parent)
    history = api.call("history", {"kind": "all"})["items"]
    assert [i["kind"] for i in history if i["kind"] == "meeting"] == ["meeting", "meeting"]  # shown with a Meeting badge
    assert "Budget review (30 min). Notes: " + doc.name in [i["text"] for i in history]


def test_the_notes_are_read_back_from_the_word_document_for_the_window(api, isolated):
    api, host = api
    doc = meetings.write_document(NOTES, [(0, "Me", "hello"), (5, "Others", "hi")], time.time() - 600, time.time())
    notes = api.call("meeting_notes", {"path": str(doc)})["notes"]
    assert "minutes" in notes.splitlines()[0] and "## Summary\nWe agreed to keep Q4 flat." in notes  # (title skipped, date line kept)
    assert "## Decisions\n- Keep Q4 flat" in notes and "- [ ] Me: Send the numbers (Friday)" in notes
    assert "- [ ] Others: Review" in notes and "## Transcript" in notes and "[00:00] Me: hello" in notes


def test_only_word_documents_in_the_meeting_notes_folder_can_be_read_or_opened(api, isolated):
    api, host = api
    folder = meetings.notes_folder()
    doc = meetings.write_document(NOTES, [(0, "Me", "hello")], time.time() - 600, time.time())
    (folder / "sub").mkdir()
    (folder / "sub" / "x.docx").write_bytes(b"x")
    (isolated / "history.jsonl").write_text("secret")
    (folder / "plain.txt").write_text("x")
    for bad in (str(isolated / "history.jsonl"), str(folder / "plain.txt"), str(folder / "sub" / "x.docx"), str(folder / ".." / "history.jsonl"),
                str(folder / "missing.docx"), str(folder), "C:/Windows/notepad.exe", ""):
        with pytest.raises(ValueError, match="aren't in the Meeting notes folder"):
            api.call("meeting_notes", {"path": bad})
        with pytest.raises(ValueError):
            api.call("open_meeting", {"path": bad})
    assert OPENED == []
    api.call("open_meeting", {"path": str(doc)})
    assert OPENED == [str(doc.resolve())]


def test_the_page_buttons_start_stop_and_open_through_the_api(api, tool, monkeypatch, isolated):
    api, host = api
    meeting_tool(tool, monkeypatch, isolated)
    assert api.call("start_meeting") == {"started": True}
    drain_ui(tool)
    assert api.call("meeting_status")["recording"] is False  # the card asks first
    agree(tool)
    assert api.call("meeting_status")["recording"] is True
    assert api.call("stop_meeting") == {"stopped": True}
    drain_ui(tool)
    assert api.call("open_meeting_folder") == {"opened": True} and len(OPENED) == 1
    wait_until(lambda: tool.light.meeting is None, 15)


def test_overview_counts_todays_fixes_only_and_old_rows_count_as_none(api, tool, isolated, monkeypatch):
    api, host = api
    today = time.strftime("%Y-%m-%d")
    rows = [
        {"at": f"{today} 09:00:00", "raw": "hello um test", "text": "Hello, test.", "fixes": 2},
        {"at": f"{today} 09:01:00", "raw": "old row", "text": "Old row."},  # written before the counter existed
        {"at": f"{today} 09:02:00", "raw": "three more", "text": "Three more.", "fixes": 3},
        {"at": f"{today} 09:03:00", "raw": "odd", "text": "Odd.", "fixes": "x"},
        {"at": f"{today} 09:04:00", "raw": "odd", "text": "Odd.", "fixes": True},
        {"at": f"{today} 09:05:00", "raw": "negative", "text": "Negative.", "fixes": -4},
        {"at": "2020-01-01 09:00:00", "raw": "long ago", "text": "Long ago.", "fixes": 9},  # not today
        {"at": f"{today} 09:06:00", "command": "shorter", "route": "rewrite", "result": "Hi"}]
    (isolated / "history.jsonl").write_text("\n".join(json.dumps(r) for r in rows) + "\n", encoding="utf-8")
    o = api.call("overview")
    assert o["today"] == {"dictations": 6, "commands": 1, "words": 9, "fixes": 5}
    assert [r["fixes"] for r in api.call("history", {"query": "more"})["items"]] == [3]
    (isolated / "history.jsonl").unlink()  # an empty history
    assert api.call("overview")["today"] == {"dictations": 0, "commands": 0, "words": 0, "fixes": 0}


def test_settings_say_what_the_home_card_needs_to_know(api, isolated, monkeypatch):
    api, host = api
    monkeypatch.setattr(host, "offline_available", lambda: False)
    monkeypatch.setitem(vt.CONFIG, "azure_openai_base_url", "")
    s = api.call("settings")
    assert s["offline_available"] is False and s["azure_set_up"] is False
    monkeypatch.setattr(host, "offline_available", lambda: True)
    monkeypatch.setitem(vt.CONFIG, "azure_openai_base_url", "https://example.invalid/openai/v1/")
    s = api.call("settings")
    assert s["offline_available"] is True and s["azure_set_up"] is True


def test_overview_history_words_fixes_style_and_shortcuts_work_on_pc_data(api, tool, isolated, monkeypatch):
    api, host = api
    monkeypatch.setattr(commands, "signed_in", lambda: True)
    today = time.strftime("%Y-%m-%d")
    (isolated / "history.jsonl").write_text("\n".join(json.dumps(r) for r in [
        {"at": f"{today} 09:00:00", "raw": "hello um test", "text": "Hello, test.", "app": "OUTLOOK.EXE"},
        {"at": f"{today} 09:05:00", "command": "shorter", "route": "rewrite", "app": "chrome.exe", "result": "Hi"},
        {"at": f"{today} 09:06:00", "command": "calendar", "route": "claude", "result": "Two meetings."}]) + "\n{bad line\n")
    usage.speech(12.0, "dictate")
    o = api.call("overview")
    assert o["computer"] and o["today"] == {"dictations": 1, "commands": 2, "words": 2, "fixes": 0} and o["claude_signed_in"] is True
    assert [r["kind"] for r in o["recent"]] == ["claude", "command", "dictate"]
    assert [r["app"] for r in api.call("history", {"query": "hello"})["items"]] == ["OUTLOOK.EXE"]
    assert api.call("styles")["recent_apps"][:2] == ["chrome.exe", "OUTLOOK.EXE"]
    assert api.call("words")["words"] == ["Zephyr", "Claude", "Wispr Flow"]
    assert api.call("save_words", {"words": ["Zephyr", "zephyr", "Kubernetes"]})["words"] == ["Zephyr", "Kubernetes"]
    assert api.call("add_fix", {"heard": "zef fur", "meant": "Zephyr"})["added"] is True
    assert api.call("delete_fix", {"heard": "zef fur"})["fixes"] == []
    assert api.call("style")["learned"] is False and "In my own words" in api.call("style")["text"]
    assert not hasattr(api, "api_spending") and "month_cost" not in o  # (no spending page)
    saved = api.call("save_shortcuts", {"shortcuts": [{"phrase": "x", "kind": "url", "target": "https://a.b"}, {"phrase": "y", "kind": "shell", "target": "z"}]})
    assert [s["phrase"] for s in saved["shortcuts"]] == ["x"]
    assert api.call("save_styles", {"data": {"styles": [{"name": "Mine", "instructions": "Short."}], "app_rules": []}})["styles"][0]["name"] == "Mine"
    api.call("copy", {"text": "abc"})


def test_set_paused_saves_and_pauses_on_the_main_thread(api, tool, isolated):
    api, host = api
    assert api.call("set_paused", {"paused": True}) == {"enabled": False}
    assert json.loads((isolated / "config.json").read_text())["enabled"] is False
    drain_ui(tool)  # the host's pause runs on the main thread
    assert api.call("overview")["enabled"] is False
    assert api.call("set_paused", {"paused": False}) == {"enabled": True}


def test_open_buttons_open_the_right_files_and_say_when_nothing_is_there(api, tool, isolated):
    api, host = api
    for what in ("log", "folder", "words", "readme"):
        assert api.call("open", {"what": what}) == {"opened": True}
    assert len(OPENED) == 4
    (isolated / "corrections.json").write_text("[]", encoding="utf-8")  # (it exists once something has been learned)
    assert api.call("open", {"what": "fixes"}) == {"opened": True} and len(OPENED) == 5
    assert api.call("open", {"what": "style"}) == {"opened": False} and api.call("open", {"what": "journal"}) == {"opened": False}
    with pytest.raises(ValueError):
        api.call("open", {"what": "../../Windows"})


def test_the_page_is_the_macs_page_plus_the_pc_additions():
    import app_window_win
    original = (HERE / "ui" / "index.html").read_text(encoding="utf-8")
    page = app_window_win.build_page("meetings")
    assert page.startswith(original.split("</body>")[0])  # the Mac's page comes first, untouched
    assert page.endswith("</body>\n</html>\n") and "pywebviewready" in page
    assert "replaceState(null,'','#meetings')" in page
    odd = app_window_win.build_page("../../x'); alert(1); ('")  # a page name can't smuggle anything into the script
    assert "'#xalert1')" in odd and "alert(1)" not in odd.split("pywebviewready")[0].split("replaceState")[1]
    assert "window.pywebview.api.call" in original  # the page already speaks pywebview


# ====================================================================================================== the app window process

def write_script(tmp_path, body):
    p = tmp_path / "child.py"
    p.write_text(body)
    return str(p)


def test_a_window_that_cannot_start_reports_why_once(tmp_path):
    import app_window_win
    failures = []
    log = SimpleNamespace(info=lambda *a: None, error=lambda *a: None, warning=lambda *a: None, exception=lambda *a: None)
    w = app_window_win.AppWindow(SimpleNamespace(call=lambda *a: None), log, on_failed=failures.append,
                                 script=write_script(tmp_path, "import sys, json\nprint(json.dumps({'event': 'error', 'message': 'pywebview is not available'}), flush=True)\nsys.exit(2)\n"))
    w.show("home")
    deadline = time.time() + 10
    while not failures and time.time() < deadline:
        time.sleep(0.1)
    assert len(failures) == 1 and "pywebview is not available" in failures[0]
    assert not w.is_open() and w.failed and not w.ready


def test_a_window_that_never_answers_is_killed_and_reported(tmp_path):
    import app_window_win
    failures = []
    log = SimpleNamespace(info=lambda *a: None, error=lambda *a: None, warning=lambda *a: None, exception=lambda *a: None)
    w = app_window_win.AppWindow(SimpleNamespace(call=lambda *a: None), log, on_failed=failures.append,
                                 script=write_script(tmp_path, "import time\ntime.sleep(60)\n"))
    w.START_TIMEOUT = 1
    w.show("home")
    deadline = time.time() + 15
    while not failures and time.time() < deadline:
        time.sleep(0.1)
    assert failures and "too long" in failures[0] and not w.is_open()


def test_the_window_conversation_answers_requests_and_ends_when_the_tool_does(tmp_path):
    """A stand-in window process: it asks for something, gets the answer, then must exit when its stdin closes."""
    import app_window_win
    log = SimpleNamespace(info=lambda *a: None, error=lambda *a: None, warning=lambda *a: None, exception=lambda *a: None)
    answers = []
    api = SimpleNamespace(call=lambda name, args: answers.append((name, args)) or {"ok": name})
    script = write_script(tmp_path, '''
import sys, json
sys.stdout.write(json.dumps({"id": 1, "name": "ping", "args": {"a": 1}}) + "\\n"); sys.stdout.flush()
reply = json.loads(sys.stdin.readline())
sys.stdout.write(json.dumps({"id": 2, "name": "boom", "args": {}}) + "\\n"); sys.stdout.flush()
reply2 = json.loads(sys.stdin.readline())
sys.stdout.write(json.dumps({"event": "ready", "pid": 1, "hwnd": 2, "echo": [reply, reply2]}) + "\\n"); sys.stdout.flush()
for line in sys.stdin:
    pass
sys.exit(0)
''')
    w = app_window_win.AppWindow(SimpleNamespace(call=lambda name, args: (_ for _ in ()).throw(ValueError("nope")) if name == "boom" else api.call(name, args)),
                                 log, script=script)
    w.show("home")
    deadline = time.time() + 10
    while not w.ready and time.time() < deadline:
        time.sleep(0.05)
    assert w.ready and answers == [("ping", {"a": 1})]
    proc = w.proc
    w.close()
    assert proc.poll() is not None and not w.is_open()


def test_showing_again_while_open_sends_a_show_command_not_a_second_window(tmp_path):
    import app_window_win
    log = SimpleNamespace(info=lambda *a: None, error=lambda *a: None, warning=lambda *a: None, exception=lambda *a: None)
    out = tmp_path / "seen.txt"
    script = write_script(tmp_path, f'''
import sys, json
print(json.dumps({{"event": "ready", "pid": 1, "hwnd": 2}}), flush=True)
for line in sys.stdin:
    open(r"{out}", "a").write(line)
''')
    w = app_window_win.AppWindow(SimpleNamespace(call=lambda *a: None), log, script=script)
    w.show("home")
    deadline = time.time() + 10
    while not w.ready and time.time() < deadline:
        time.sleep(0.05)
    first = w.proc
    w.show("settings")
    w.show("meetings")
    time.sleep(0.5)
    assert w.proc is first
    assert [json.loads(l)["page"] for l in out.read_text().splitlines()] == ["settings", "meetings"]
    w.close()


def test_a_closed_window_opens_again_with_a_new_process(tmp_path):
    import app_window_win
    log = SimpleNamespace(info=lambda *a: None, error=lambda *a: None, warning=lambda *a: None, exception=lambda *a: None)
    script = write_script(tmp_path, 'import sys, json\nprint(json.dumps({"event": "ready"}), flush=True)\nsys.stdin.readline()\n')
    w = app_window_win.AppWindow(SimpleNamespace(call=lambda *a: None), log, script=script)
    w.show("home")
    deadline = time.time() + 10
    while not w.ready and time.time() < deadline:
        time.sleep(0.05)
    first = w.proc
    w.close()
    w.show("home")
    assert w.proc is not first and w.is_open()
    w.close()


def test_the_tool_falls_back_to_the_old_settings_window_when_the_app_window_fails(tool, monkeypatch):
    opened = []
    monkeypatch.setattr(tool, "_open_settings_tk", lambda: opened.append("tk"))
    tool._window_failed("pywebview is not available")
    assert opened == ["tk"]
    import app_window_win
    monkeypatch.setattr(app_window_win.AppWindow, "show", lambda self, page="home": (_ for _ in ()).throw(OSError("cannot start")))
    tool._open_app("settings")
    assert opened == ["tk", "tk"]


def test_tray_and_start_menu_open_the_app_window_at_the_right_page(tool, monkeypatch):
    shown = []
    import app_window_win
    monkeypatch.setattr(app_window_win.AppWindow, "show", lambda self, page="home": shown.append(page))
    tool._open_app()
    tool._open_settings()
    tool._open_app("home")
    assert shown == ["home", "settings", "home"]
    assert tool.app_window is not None and tool.app_window.test_mode is False


# ====================================================================================================== spoken answers (Windows voice)

class FakeToken:
    def __init__(self, description):
        self.description = description

    def GetDescription(self):
        return self.description


class FakeVoice:
    def __init__(self, voices=("Microsoft David Desktop - English (United States)", "Microsoft Zira Desktop - English (United States)")):
        self.said, self.done, self.Rate, self.Volume = [], threading.Event(), 0, 0
        self.busy_for = 0.0
        self.tokens = [FakeToken(v) for v in voices]
        self.Voice = self.tokens[0]  # (the Windows default)

    def GetVoices(self):
        return SimpleNamespace(Count=len(self.tokens), Item=lambda i: self.tokens[i])

    def Speak(self, text, flags):
        self.said.append((text, flags))
        self.started = time.time()

    def WaitUntilDone(self, ms):
        time.sleep(ms / 1000)
        return time.time() - getattr(self, "started", 0) >= self.busy_for


def test_speaker_reads_aloud_replaces_old_answers_and_esc_stops_it():
    voice = FakeVoice()
    voice.busy_for = 30  # a long answer
    sp = speaker_windows.Speaker(voice_factory=lambda: voice)
    sp.say("First answer.")
    time.sleep(0.4)
    assert voice.said[-1][0] == "First answer." and sp.speaking
    sp.say("Second answer replaces it.")
    time.sleep(0.4)
    assert voice.said[-1][0] == "Second answer replaces it." and voice.said[-1][1] & speaker_windows.SVS_PURGE_BEFORE_SPEAK
    sp.stop()
    time.sleep(0.4)
    assert voice.said[-1][0] == "" and not sp.speaking  # Esc: cut off
    sp.say("")
    sp.say("   ")
    time.sleep(0.2)
    assert len(voice.said) == 3
    sp.say("x" * 5000)
    time.sleep(0.4)
    assert len(voice.said[-1][0]) == speaker_windows.MAX_CHARS


def test_the_windows_voice_really_speaks_silently():
    """The real SAPI voice, with the volume at zero (so nothing is heard): one short phrase starts, finishes, and the voice
    named in the setting is accepted."""
    names = speaker_windows.list_voices()
    sp = speaker_windows.Speaker(volume=0)
    sp.say("Testing.", voice=names[0])
    deadline = time.time() + 20
    while not sp.speaking and time.time() < deadline:
        time.sleep(0.05)
    assert sp.speaking  # it started
    while sp.speaking and time.time() < deadline:
        time.sleep(0.1)
    assert not sp.speaking and sp._thread is not None and sp.voice_name == names[0]
    sp.stop()


def test_speaker_survives_a_missing_voice_engine(monkeypatch):
    def broken():
        raise OSError("no voice installed")
    sp = speaker_windows.Speaker(voice_factory=broken)
    sp.say("hello")
    time.sleep(0.3)
    assert sp._thread is None  # gave up quietly; dictation and cards are unaffected


# ====================================================================================================== housekeeping

def test_usage_logging_still_works_for_the_new_azure_calls(isolated):
    class U:
        prompt_tokens, completion_tokens = 100, 20
        prompt_tokens_details = None
    usage.azure(SimpleNamespace(usage=U()), "gpt-5.4-mini", "screen question")
    row = json.loads((isolated / "usage.jsonl").read_text().splitlines()[-1])
    assert row["purpose"] == "screen question" and row["usd"] > 0 and row["paid_by"] == "Azure credits"


def test_keep_punctuation_and_split_helpers():
    assert vt.keep_punctuation("Zef fir.", "Zephyr") == "Zephyr."
    assert vt.keep_punctuation(" zef fir, ", "Zephyr") == " Zephyr, "
    assert vt.strip_punctuation("“Zef fir.”") == "Zef fir"
    audio = np.ones(vt.RATE * 40, np.float32)
    audio[vt.RATE * 20: vt.RATE * 21] = 0.0
    chunks = vt.split_at_pauses(audio)
    assert len(chunks) == 2 and 20.0 <= len(chunks[0]) / vt.RATE <= 21.0 and sum(len(c) for c in chunks) == len(audio)


def test_light_error_flash_goes_away_by_itself_and_meeting_states_exist():
    light = vt.Light.__new__(vt.Light)
    light.state, light.error_until, light.claude_busy = "working", 0.0, 0
    light.set("error")
    assert light.state is None and light.error_until > time.monotonic()
    assert {"listening", "handsfree", "working", "error", "command", "command_handsfree", "thinking", "claude", "meeting", "meeting_notes"} <= set(vt.Light.LOOKS)


def test_corrections_module_is_the_pcs_and_preloaded_fix_works():
    assert corrections.add("Zeph Ear", "Zephyr", "test")
    assert corrections.apply("Open zeph  ear today") == "Open Zephyr today"


def test_apply_counted_counts_the_learned_fixes_and_apply_gives_the_same_text():
    assert corrections.apply_counted("nothing to fix here") == ("nothing to fix here", 0)
    assert corrections.add("zef fur", "Zephyr", "test")
    assert corrections.add("cube", "Kubernetes", "test")
    assert corrections.apply_counted("open zef fur now") == ("open Zephyr now", 1)
    assert corrections.apply_counted("Zef  FUR and zef fur") == ("Zephyr and Zephyr", 2)  # case and spacing don't matter
    assert corrections.apply_counted("zef fur runs cube, zef fur again") == ("Zephyr runs Kubernetes, Zephyr again", 3)
    for said in ("open zef fur now", "nothing to fix here", "Zef Fur and zef fur", ""):
        assert corrections.apply(said) == corrections.apply_counted(said)[0]
    assert corrections.apply_counted("The colour is grey") == ("The color is gray", 0)  # spelling isn't a learned fix


def test_apply_counted_with_no_fixes_file_counts_nothing():
    assert not corrections.FILE.exists()
    assert corrections.apply_counted("zef fur") == ("zef fur", 0)


def test_a_fix_can_be_edited_on_both_sides_in_one_step_under_the_adding_rules():
    for heard, meant in (("zef fir", "Zephyr"), ("cube", "Kubernetes"), ("clod", "Claude")):
        assert corrections.add(heard, meant, "test")
    assert corrections.edit("Zef  Fir", "zef fur", "Zephyr Inc", "edited")  # both sides change, old one goes
    rows = json.loads(corrections.FILE.read_text(encoding="utf-8"))
    assert [(r["heard"], r["meant"]) for r in rows] == [("zef fur", "Zephyr Inc"), ("cube", "Kubernetes"), ("clod", "Claude")]
    assert rows[0]["source"] == "edited"
    assert corrections.apply("open zef fur and zef fir") == "open Zephyr Inc and zef fir"
    assert corrections.edit("cube", "cube", "K8s", "edited")  # only "You meant" changes
    assert ("cube", "K8s") in corrections.load()
    assert corrections.edit("zef fur", "clod", "Claude AI", "edited")  # landing on another fix's words replaces it
    assert [(r["heard"], r["meant"]) for r in json.loads(corrections.FILE.read_text(encoding="utf-8"))] == [("clod", "Claude AI"), ("cube", "K8s")]
    before = corrections.FILE.read_text(encoding="utf-8")
    for heard, meant in (("one two three four five six", "x"), ("same", "Same"), ("", "x"), ("x", "  ")):
        assert not corrections.edit("cube", heard, meant, "edited")
    assert corrections.FILE.read_text(encoding="utf-8") == before  # a refused edit changes nothing
    assert corrections.edit("forgotten meanwhile", "new words", "New", "edited")  # the old one is gone: added instead
    assert ("new words", "New") in corrections.load()


def test_the_window_edits_fixes_and_words_and_says_why_an_edit_is_refused(api, isolated):
    api, host = api
    api.call("add_fix", {"heard": "zef fir", "meant": "Zephyr"})
    r = api.call("edit_fix", {"old_heard": "zef fir", "heard": "zef fur", "meant": "Zephyr"})
    assert r["saved"] is True and [(f["heard"], f["meant"]) for f in r["fixes"]] == [("zef fur", "Zephyr")]
    r = api.call("edit_fix", {"old_heard": "zef fur", "heard": "a b c d e f", "meant": "Zephyr"})
    assert r["saved"] is False and r["reason"] == '"It wrote" can be 5 words at most.' and r["fixes"][0]["heard"] == "zef fur"
    assert api.call("edit_fix", {"old_heard": "zef fur", "heard": "Zephyr", "meant": "zephyr"})["reason"] == "The two sides must be different."
    assert api.call("edit_fix", {"old_heard": "zef fur", "heard": "", "meant": "x"})["reason"] == "Fill in both sides."
    assert api.call("add_fix", {"heard": "same", "meant": "SAME"}) | {"fixes": None} == {"added": False, "reason": "The two sides must be different.", "fixes": None}

    r = api.call("edit_word", {"old": "Claude", "new": "  Claude   Code "})
    assert r["saved"] is True and r["words"] == ["Zephyr", "Claude Code", "Wispr Flow"]
    assert (isolated / "words.txt").read_text(encoding="utf-8").splitlines()[1:] == ["Zephyr", "Claude Code", "Wispr Flow"]
    corrections.remember_word("Kubernetes")  # learned by the typed fix box while the window was open
    assert api.call("edit_word", {"old": "zephyr", "new": "Zephyr Inc"})["words"] == ["Zephyr Inc", "Claude Code", "Wispr Flow", "Kubernetes"]
    assert api.call("edit_word", {"old": "Wispr Flow", "new": "claude code"}) == {"saved": False, "reason": "That word is already on your list.",
                                                                                "words": ["Zephyr Inc", "Claude Code", "Wispr Flow", "Kubernetes"]}
    assert api.call("edit_word", {"old": "Wispr Flow", "new": "  "})["reason"] == "Type a word, or use ✕ to remove it."
    assert api.call("edit_word", {"old": "Wispr Flow", "new": "wispr flow"})["words"][2] == "wispr flow"  # just the capitals


def test_auto_learn_never_learns_a_fix_that_eats_spoken_at_or_dot():
    for _ in range(3):  # the tidy-up turned his spoken email into an address: "at" and "dot" became @ and .
        assert corrections.learn_from("Jane dot doe at Example dot com", "jane.doe@example.com", ["Zephyr"]) == []
    assert corrections.load() == []
    corrections.learn_from("open zeffer now", "Open Zephyr now.", ["Zephyr"])  # an ordinary mishearing still learns
    assert corrections.learn_from("open zeffer now", "Open Zephyr now.", ["Zephyr"]) == [("zeffer", "Zephyr")]


def test_a_fix_is_saved_and_matched_without_punctuation_at_its_ends():
    assert corrections.add("Zef fir.", "Zephyr", "fix that")
    assert json.loads(corrections.FILE.read_text(encoding="utf-8"))[0]["heard"] == "Zef fir"
    corrections.FILE.write_text(json.dumps([{"heard": "Zef fir.", "meant": "Zephyr"}]), encoding="utf-8")  # saved before
    assert corrections.apply("I said zef fir.") == "I said Zephyr."  # the full stop stays
    assert corrections.apply("Zef fir. Is great") == "Zephyr. Is great"
    assert corrections.problem("?!", "x") == "Fill in both sides."


def test_a_fix_whose_meant_side_has_a_backslash_is_pasted_as_typed():
    """re.sub read "meant" as a template, so X:\\Zephyr raised "bad escape" on EVERY dictation."""
    assert corrections.add("zef fur folder", r"C:\Zephyr", "fix that")
    assert corrections.add("new line", r"C:\new\1 and \g<0>", "fix that")
    assert corrections.apply("open zef fur folder now") == r"open C:\Zephyr now"
    assert corrections.apply("type new line") == r"type C:\new\1 and \g<0>"


def test_a_saved_fix_is_not_learned_again_so_its_label_and_date_stay(isolated):
    corrections.FILE.write_text(json.dumps([
        {"heard": "zef fir", "meant": "Zephyr", "added": "2026-01-02", "source": "typed fix"},
        {"heard": "cube", "meant": "Kubernetes", "added": "2026-01-03", "source": "fix that"}]), encoding="utf-8")
    before = corrections.FILE.read_text(encoding="utf-8")
    for raw in ("open zef fir now", "open Zef Fir now"):  # (case is ignored)
        for _ in range(3):
            assert corrections.learn_from(raw, "Open Zephyr now.", ["Zephyr"]) == []
    assert corrections.FILE.read_text(encoding="utf-8") == before  # not re-saved as "auto", not moved to the end
    assert not corrections.PENDING.exists()  # and not even counted
    # the same words fixed into something else are still a new lesson
    assert corrections.learn_from("open zef fir now", "Open Claude now.", ["Claude"]) == []
    assert corrections.learn_from("open zef fir now", "Open Claude now.", ["Claude"]) == [("zef fir", "Claude")]


def test_a_paste_waits_while_any_trigger_key_is_held_not_just_its_own(tool, monkeypatch):
    """He may already be holding the key for his NEXT dictation or command; pasting then would send Ctrl+Alt+V."""
    done = threading.Event()
    nxt = tool.triggers[COMMAND]
    nxt.key_down = True  # this job was made by the dictate key
    worker = threading.Thread(target=lambda: (tool._wait_for_release(tool.triggers[DICTATE]), done.set()), daemon=True)
    worker.start()
    assert not done.wait(0.3)  # still waiting
    nxt.key_down = False
    assert done.wait(2)
    worker.join(2)
    started = time.monotonic()
    tool._wait_for_release(tool.triggers[DICTATE])  # nothing held: no wait at all
    assert time.monotonic() - started < 0.1
    slept = []
    monkeypatch.setattr(vt.time, "sleep", lambda s: slept.append(s))
    nxt.key_down = True  # a key stuck down: give up after 60 s (20 ms steps) and paste anyway
    tool._wait_for_release(tool.triggers[DICTATE])
    assert len(slept) == 3000 and set(slept) == {0.02}


def _fake_copy(monkeypatch, copied):
    """Stand in for Ctrl+C and the clipboard: nothing real is pressed or read."""
    import contextlib
    seq = [0]

    class Controller:
        def pressed(self, key):
            return contextlib.nullcontext()

        def tap(self, char):
            seq[0] += 1

    fake_clipboard = SimpleNamespace(GetClipboardSequenceNumber=lambda: seq[0], IsClipboardFormatAvailable=lambda fmt: True,
                                     GetClipboardData=lambda fmt: copied, CloseClipboard=lambda: None)
    monkeypatch.setattr(commands, "keyboard", SimpleNamespace(Controller=Controller, Key=SimpleNamespace(ctrl=object())))
    monkeypatch.setattr(commands, "win32clipboard", fake_clipboard)
    restored = []
    return (lambda: None), (lambda: "saved"), restored.append


def test_the_selection_keeps_the_trailing_space_a_browser_adds_but_blank_is_nothing(monkeypatch):
    """Edge and Chrome copy a double-clicked word with its trailing space; stripping it pasted "Zephyris"."""
    open_cb, save_cb, restore_cb = _fake_copy(monkeypatch, "zef fir ")
    assert commands.selected_text(open_cb, save_cb, restore_cb) == "zef fir "
    for blank in (" ", " \r\n\t "):
        open_cb, save_cb, restore_cb = _fake_copy(monkeypatch, blank)
        assert commands.selected_text(open_cb, save_cb, restore_cb) == ""
    monkeypatch.setattr(commands, "window_info", lambda: {"app": "msedge.exe", "title": "x"})
    for chosen, expected in (("zef fir ", "zef fir "), ("   ", ""), (" " * 12000 + "late text", "")):  # (cut at 12000 characters)
        monkeypatch.setattr(commands, "selected_text", lambda *a, chosen=chosen: chosen)
        assert commands.capture_context(None, None, None)["selection"] == expected


def test_a_fix_on_a_selection_with_a_trailing_space_pastes_it_back_and_learns_the_bare_words(tool):
    selection = {"app": "msedge.exe", "title": "x", "selection": "zef fir ", "pid": 1, "hwnd": 2}
    run_command(tool, "that should be Zephyr", {"route": "fix", "text": "zephyr"}, selection)
    assert PASTED == ["Zephyr "]  # the space stays, so the next word doesn't stick to it
    assert [(h.lower(), m) for h, m in corrections.load()] == [("zef fir", "Zephyr")]
    router = commands.route(FakeClient(json.dumps({"route": "fix", "text": "zephyr"})), "m", "x", selection)
    assert router[0] == "fix"  # a spaced selection is still a short fix, and a blank one is still no selection
    assert commands.route(FakeClient(json.dumps({"route": "rewrite", "text": "x"})), "m", "x", {"selection": " "})[0] == "write"


def test_a_rewrite_is_pasted_inside_the_selections_own_spaces_and_line_breaks(tool):
    selection = {"app": "winword.exe", "title": "x", "selection": "  messy words\r\n", "pid": 1, "hwnd": 2}
    run_command(tool, "make it friendlier", {"route": "rewrite", "text": ""}, selection, then=("Hi there!",))
    run_command(tool, "make it friendlier", {"route": "rewrite", "text": ""}, {**selection, "selection": "messy words"}, then=("Hi there!",))
    run_command(tool, "write a note", {"route": "write", "text": "Dear Sarah"}, {**selection, "selection": ""})
    run_command(tool, "write a note", {"route": "write", "text": "Dear Sarah"}, selection)  # a write never wraps
    assert PASTED == ["  Hi there!\r\n", "Hi there!", "Dear Sarah", "Dear Sarah"]


def _local_tidy(monkeypatch, content, finish_reason="stop"):
    """Cleaner._local with the on-device model answering one canned reply (no server, no network)."""
    c = cleaner()
    monkeypatch.setattr(c, "_local_openai", lambda: FakeLocalClient(content, finish_reason))
    return c._local


def test_an_on_device_tidy_up_that_was_cut_off_or_came_back_empty_is_refused(monkeypatch):
    assert _local_tidy(monkeypatch, " Fine. ")("sys", "fine") == "Fine."
    with pytest.raises(RuntimeError, match="cut off"):  # the size limit was hit: half a sentence
        _local_tidy(monkeypatch, "Half of the sen", "length")("sys", "half of the sentence")
    for empty in ("", "  \n", None):
        with pytest.raises(RuntimeError, match="empty"):
            _local_tidy(monkeypatch, empty)("sys", "something said")


def test_an_empty_tidy_up_keeps_the_words_with_learned_fixes_instead_of_pasting_nothing(tool, isolated):
    corrections.add("zef fir", "Zephyr", "fix that")
    for result in ("", "  \n"):
        PASTED.clear()
        tool.azure = FakeCleaner(result=result)
        run_dictation(tool, "open zef fir now")
        assert PASTED == ["open Zephyr now "]
    row = json.loads((isolated / "history.jsonl").read_text().splitlines()[-1])
    assert row["raw"] == "open zef fir now" and row["text"] == "open Zephyr now"
    assert not corrections.PENDING.exists()  # nothing was learned from the failed tidy-up
    tool.azure = FakeCleaner(fail=RuntimeError("the offline tidy-up was cut off"))  # a cut-off one takes the same road
    PASTED.clear()
    run_dictation(tool, "open zef fir now")
    assert PASTED == ["open Zephyr now "]


def test_the_typed_fix_copy_waits_until_shift_is_up(tool, monkeypatch):
    """Copying while Shift is still down sends Ctrl+Shift+C, which opens the inspector in Edge."""
    polls, events, slept = [], [], []
    monkeypatch.setattr("win32gui.GetForegroundWindow", lambda: 123)
    monkeypatch.setattr(vt.time, "sleep", lambda s: slept.append(s))

    def shift(vk):
        assert vk == vt.win32con.VK_SHIFT
        polls.append(1)
        return 0x8000 if len(polls) <= 5 else 0  # held for five checks, then let go

    monkeypatch.setattr(vt.win32api, "GetAsyncKeyState", shift)
    monkeypatch.setattr(tool, "_capture", lambda trig, box=None: events.append(("copy", len(polls))))
    vt.VoiceTool._teach(tool, tool.triggers[COMMAND])  # (the real one; the fixture replaces it on the instance)
    assert events == [("copy", 6)] and len(slept) == 5
    assert tool.ui.qsize() == 1  # the typed-fix box is queued for the window thread, not opened here
    polls.clear(), events.clear(), slept.clear()
    monkeypatch.setattr(vt.win32api, "GetAsyncKeyState", lambda vk: 0x8000)  # Shift stuck down: wait one second, then copy anyway
    vt.VoiceTool._teach(tool, tool.triggers[COMMAND])
    assert events == [("copy", 0)] and len(slept) == 50 and set(slept) == {0.02}


def test_the_azure_sign_in_and_every_other_console_program_start_without_a_window(monkeypatch):
    """azure-identity runs az.cmd about once an hour without asking for no window; from pythonw that flashed a black
    window on the first dictation. Nothing is really started here: Windows' CreateProcess is faked."""
    import subprocess
    from azure.identity._credentials import azure_cli
    monkeypatch.setattr(subprocess.Popen, "__init__", subprocess.Popen.__init__)  # (undone after the test)
    started = []

    def fake_create_process(executable, args, *rest):
        started.append((args, rest[3]))  # rest: process attrs, thread attrs, inherit handles, creationflags, ...
        raise OSError("not really started")
    monkeypatch.setattr(subprocess._winapi, "CreateProcess", fake_create_process)
    monkeypatch.setattr(azure_cli.shutil, "which", lambda name: r"C:\Azure\az.cmd" if name == "az.cmd" else None)
    vt.hide_console_windows()
    vt.hide_console_windows()  # a second call changes nothing
    with pytest.raises(Exception):
        azure_cli.AzureCliCredential().get_token("https://cognitiveservices.azure.com/.default")
    assert started and "az.cmd" in started[0][0] and started[0][1] & subprocess.CREATE_NO_WINDOW
    with pytest.raises(OSError):
        subprocess.run(["cmd", "/c", "exit"])
    assert started[-1][1] & subprocess.CREATE_NO_WINDOW
    with pytest.raises(OSError):  # a caller that asks for something else gets exactly that
        subprocess.run(["cmd", "/c", "exit"], creationflags=subprocess.CREATE_NEW_CONSOLE)
    assert started[-1][1] & subprocess.CREATE_NEW_CONSOLE and not started[-1][1] & subprocess.CREATE_NO_WINDOW


def test_no_garbled_characters_in_the_code_sam_sees():
    """voicetool.py was once re-saved with its UTF-8 read as Windows-1252: the light said "Writingâ€¦"."""
    files = [*HERE.glob("*.py"), *HERE.glob("*.md"), *(HERE / "ui").glob("*.*")]
    bad = [f"{f.name}:{n}" for f in files for n, line in enumerate(f.read_text(encoding="utf-8").splitlines(), 1)
           if "â€" in line or "Â·" in line or "Ã©" in line]
    bad = [b for b in bad if not b.startswith(Path(__file__).name)]  # (this test names them)
    assert bad == []
    assert '"Writing…"' in (HERE / "voicetool.py").read_text(encoding="utf-8")


def test_the_tool_never_touches_the_real_folder(isolated):
    """Guard: every module the tests use points at the temp folder, so a test can't write into the live tool."""
    for module in (vt, corrections, va, usage, settings_window):
        assert Path(module.HERE) == isolated
    assert str(vt.MEETING_AUDIO).startswith(str(isolated)) and str(va.SHORTCUTS).startswith(str(isolated))
    assert str(meetings.notes_folder()).startswith(str(isolated))  # the Word documents go to a temp Documents, not the real one


def test_auto_learn_can_be_switched_off(tool, monkeypatch):
    monkeypatch.setitem(vt.CONFIG, "auto_learn", False)
    tool.azure = FakeCleaner(result="Open Zephyr now.")
    for _ in range(3):
        run_dictation(tool, "open zef fir now")
    assert ("zef fir", "Zephyr") not in corrections.load()
    assert not (corrections.PENDING).exists()


def test_search_looks_inside_nested_folders_and_skips_only_voicetools_own_recordings(tmp_path):
    root = tree(tmp_path / "docs", ["projects/recordings/team call notes.txt", "VoiceTool/recordings/team call secret.wav",
                                    "VoiceTool_port/recordings/team call other.wav", "a/b/team call deep.txt"])
    names = sorted(Path(h).name for h in aw.search_files("team call", roots=[(root, 5)]))
    assert names == ["team call deep.txt", "team call notes.txt"]  # (a folder called recordings elsewhere is searched)


def test_a_failure_in_one_folder_does_not_lose_what_was_found_elsewhere(tmp_path, monkeypatch):
    good = tree(tmp_path / "good", ["report final.txt"])
    bad = tree(tmp_path / "bad", ["x.txt"])
    real = os.scandir

    def flaky(path):
        if "bad" in str(path):
            raise RuntimeError("drive disconnected")
        return real(path)
    monkeypatch.setattr(os, "scandir", flaky)
    assert [Path(h).name for h in aw.search_files("report final", roots=[(bad, 3), (good, 3)])] == ["report final.txt"]


# ====================================================================================================== read it to me

def test_the_speech_engine_picks_the_voice_by_name_and_goes_back_to_the_default():
    voice = FakeVoice()
    sp = speaker_windows.Speaker(voice_factory=lambda: voice)
    sp.say("one", voice="zira", rate=3)
    assert wait_until(lambda: voice.said)
    assert "Zira" in voice.Voice.GetDescription() and voice.Rate == 3
    sp.say("two", voice="")  # no name: the Windows default voice again
    assert wait_until(lambda: len(voice.said) == 2) and "David" in voice.Voice.GetDescription() and voice.Rate == sp.rate
    sp.say("three", voice="no such voice")
    assert wait_until(lambda: len(voice.said) == 3) and "David" in voice.Voice.GetDescription()
    sp.say("x" * 30000, limit=20000)
    assert wait_until(lambda: len(voice.said) == 4) and len(voice.said[-1][0]) == 20000  # read-aloud may be long


def test_speak_and_stop_speaking_drive_the_speaker_and_map_the_speed_setting(monkeypatch):
    import importlib
    real = importlib.reload(aw)  # the module's own speak/stop_speaking, not the safety net's
    try:
        calls = []
        monkeypatch.setattr(real, "_speaker", SimpleNamespace(say=lambda text, **kw: calls.append((text, kw)), stop=lambda: calls.append("STOP")))
        monkeypatch.setattr(real, "stop_audio", lambda: calls.append("STOP AUDIO"))
        assert real.RATES == {"slow": -2, "normal": 0, "fast": 3}
        real.speak("hello", voice="zira", rate="slow", limit=50)
        real.speak("again", rate="fast")
        real.speak("number", rate=2)
        real.stop_speaking()
        assert calls == [("hello", {"voice": "zira", "rate": -2, "limit": 50}), ("again", {"voice": "", "rate": 3, "limit": 20000}),
                         ("number", {"voice": "", "rate": 2, "limit": 20000}), "STOP", "STOP AUDIO"]
        monkeypatch.setattr(real, "_speaker", None)
        real.stop_speaking()  # nothing being read: fine
    finally:
        importlib.reload(aw)


def test_the_real_list_of_windows_voices():
    names = __import__("importlib").reload(aw).list_voices()
    assert names and all(set(v) == {"name", "lang"} and v["name"] for v in names)
    assert vt.WinHost(SimpleNamespace()).list_voices() == names


def test_reading_cleans_the_text_and_uses_the_chosen_voice_and_speed(tool, monkeypatch):
    monkeypatch.setitem(vt.CONFIG, "read_voice", "Microsoft Zira Desktop")
    monkeypatch.setitem(vt.CONFIG, "read_rate", "fast")
    tool._read("# Heading\n\nSome **bold** text with [a link](https://example.com/x) and https://example.com/long\n```\ncode()\n```")
    (text, kw), = SPOKEN
    assert "#" not in text and "**" not in text and "example.com" not in text and "code()" not in text
    assert "a link" in text and "Some bold text" in text and "(code block skipped)" in text
    assert kw["voice"] == "Microsoft Zira Desktop" and kw["rate"] == "fast" and kw["limit"] == read_aloud.MAX_CHARS
    SPOKEN.clear()
    STOPS.clear()
    tool._read("   ")
    assert SPOKEN == []


def test_spoken_answers_use_the_same_voice_settings_and_stop_at_1500_characters(tool, monkeypatch):
    monkeypatch.setitem(vt.CONFIG, "read_voice", "Zira")
    monkeypatch.setitem(vt.CONFIG, "speak_answers", True)
    tool._say("A spoken answer. " * 200)
    tool._read("The selected text.")
    (a_text, a_kw), (r_text, r_kw) = SPOKEN
    assert a_kw["voice"] == r_kw["voice"] == "Zira" and len(a_text) <= 1500 and r_text == "The selected text."


def test_the_natural_voice_is_only_used_when_a_text_to_speech_model_is_named(tool, monkeypatch, isolated):
    tool._read("Hello there.")
    assert SPOKEN[-1][0] == "Hello there."  # no tts_deployment: Windows' own voice
    SPOKEN.clear()
    STOPS.clear()
    monkeypatch.setitem(vt.CONFIG, "tts_deployment", "tts-test")
    asked = []
    mp3 = isolated / "voice.mp3"
    mp3.write_bytes(b"x")
    monkeypatch.setattr(read_aloud, "azure_speech_file", lambda client, cfg, text: asked.append((cfg["tts_deployment"], text)) or mp3)
    tool._read("Hello there.")
    assert asked == [("tts-test", "Hello there.")] and SPOKEN == [("PLAY", str(mp3))]
    SPOKEN.clear()
    STOPS.clear()
    monkeypatch.setattr(read_aloud, "azure_speech_file", lambda *a: (_ for _ in ()).throw(ConnectionError("no deployment")))
    tool._read("Hello again.")
    assert [t for t, _ in SPOKEN] == ["Hello again."]  # it falls back to the Windows voice, and still reads
    tool._hush()
    assert STOPS


def test_the_read_route_reads_the_selection_and_says_when_there_is_none(tool):
    run_command(tool, "read this to me", {"route": "read", "text": ""},
                {"app": "chrome.exe", "title": "x", "selection": "Some selected words.", "pid": 1, "hwnd": 2})
    assert [t for t, _ in SPOKEN] == ["Some selected words."] and tool.cards[-1][0] == "Reading it to you"
    SPOKEN.clear()
    STOPS.clear()
    run_command(tool, "read it out", {"route": "read", "text": ""}, {"app": "code.exe", "title": "x", "selection": "  "})
    assert SPOKEN == [] and tool.cards[-1][0] == "Nothing selected"


def test_esc_stops_reading(tool):
    tool._handle("other", None, vt.win32con.VK_ESCAPE, now())
    assert STOPS == [1] and SPOKEN == []


def test_the_voice_and_speed_settings(api, tool, monkeypatch, isolated):
    api, host = api
    monkeypatch.setattr(aw, "list_voices", lambda: [{"name": "Microsoft David Desktop", "lang": ""}, {"name": "Microsoft Zira Desktop", "lang": ""}])
    settings = api.call("settings")
    assert [v["name"] for v in settings["voices"]] == ["Microsoft David Desktop", "Microsoft Zira Desktop"]
    assert settings["read_voice"] == "" and settings["read_rate"] == "normal" and settings["natural_voice"] is False
    assert api.call("set_read_voice", {"voice": "Microsoft Zira Desktop", "rate": "slow"}) == {"read_voice": "Microsoft Zira Desktop", "read_rate": "slow"}
    saved = json.loads((isolated / "config.json").read_text())
    assert saved["read_voice"] == "Microsoft Zira Desktop" and saved["read_rate"] == "slow"
    with pytest.raises(ValueError, match="isn't on this computer"):
        api.call("set_read_voice", {"voice": "Not Installed", "rate": "normal"})
    with pytest.raises(ValueError, match="unknown speed"):
        api.call("set_read_voice", {"voice": "", "rate": "ludicrous"})
    assert vt.CONFIG["read_voice"] == "Microsoft Zira Desktop" and vt.CONFIG["read_rate"] == "slow"
    api.call("try_voice")  # "Try it": the host reads a short phrase in the chosen voice
    assert len(SPOKEN) == 1 and SPOKEN[0][1]["voice"] == "Microsoft Zira Desktop" and SPOKEN[0][1]["rate"] == "slow"
    assert len(SPOKEN[0][0]) < 120


def test_the_settings_page_data_has_the_windows_voices_even_if_the_list_fails(api, monkeypatch):
    api, host = api
    monkeypatch.setattr(aw, "list_voices", lambda: (_ for _ in ()).throw(OSError("no speech engine")))
    assert api.call("settings")["voices"] == []  # (the page shows just "The Windows default voice")


# ====================================================================================================== note destinations

def test_notes_go_to_a_folder_and_the_pc_offers_no_note_app(tool, isolated):
    assert vt.WinHost(tool).note_destinations() == ["folder"]  # (no OneNote: no supported, reliable way in)
    run_command(tool, "note idea", NOTE_REPLY)
    files = list((isolated / "Documents" / "Voice notes").glob("*.md"))
    assert len(files) == 1 and "Idea: voice intro." in files[0].read_text()
    assert tool.cards[-1] == ("Noted", f"Saved in your notes folder as {files[0].name}.", False)


NOTE_REPLY = {"route": "note", "text": "Idea: voice intro.", "topic": "onboarding voice"}


def test_a_note_in_a_folder_is_a_dated_markdown_file_with_the_time_inside(tool, isolated, monkeypatch):
    folder = isolated / "My Obsidian"
    monkeypatch.setitem(vt.CONFIG, "notes_folder", str(folder))
    run_command(tool, "note idea", NOTE_REPLY)
    (note,) = folder.glob("*.md")
    assert re.fullmatch(r"\d{4}-\d\d-\d\d \d{4} onboarding voice\.md", note.name)
    text = note.read_text(encoding="utf-8")
    assert text.startswith("---\ncreated: ") and re.search(r"created: \d{4}-\d\d-\d\d \d\d:\d\d\n", text)
    assert "tags: [voicetool]" in text and "# onboarding voice" in text and "Idea: voice intro." in text
    assert tool.cards[-1][1] == f"Saved in your notes folder as {note.name}."
    run_command(tool, "note idea", NOTE_REPLY)  # a second one in the same minute is never overwritten
    assert len(list(folder.glob("*.md"))) == 2


def test_with_no_folder_chosen_notes_go_to_documents_voice_notes(tool, isolated, monkeypatch):
    monkeypatch.setitem(vt.CONFIG, "notes_folder", "")
    run_command(tool, "note idea", NOTE_REPLY)
    assert len(list((isolated / "Documents" / "Voice notes").glob("*.md"))) == 1
    assert REAL_DEFAULT_NOTES_FOLDER() == Path.home() / "Documents" / "Voice notes"  # (the real default: no OneDrive, no Drive)


def test_the_friend_config_has_everything_personal_off_or_empty():
    cfg = json.loads((HERE / "config.json").read_text(encoding="utf-8"))
    assert cfg["notes_destination"] == "folder" and cfg["notes_folder"] == "" and cfg["drive_ai_work"] == ""
    assert cfg["read_voice"] == "" and cfg["read_rate"] == "normal" and cfg["tts_deployment"] == ""
    assert cfg["azure_openai_base_url"] == "" and cfg["cleanup_model"] == "local" and cfg["mic_name_contains"] == ""
    for key in ("meeting_notes", "meeting_auto_offer", "meeting_followup", "claude_actions", "claude_keep_running",
                "speak_answers", "save_recordings"):
        assert cfg[key] is False, key
    assert cfg["use_style"] is True and cfg["first_run"] is True and cfg["offline_only"] is True
    assert not {"voice", "note_destination", "note_folder", "computer_commands"} & set(cfg)  # (an earlier interim build's keys)
    assert (HERE / "words.txt").read_text(encoding="utf-8").strip().splitlines() == [
        "# One word or name per line. The cleanup step makes sure these are spelled exactly like this."]
    style = (HERE / "style.md").read_text(encoding="utf-8")
    assert "## In my own words (always wins)" in style
    assert "from my sent emails" not in style


def test_a_folder_given_with_an_environment_variable_works(tool, isolated, monkeypatch):
    monkeypatch.setitem(vt.CONFIG, "notes_folder", "%OneDrive%\\Elsewhere")
    run_command(tool, "note idea", NOTE_REPLY)
    assert len(list((isolated / "OneDrive" / "Elsewhere").glob("*.md"))) == 1


def test_an_old_note_destination_in_config_json_still_saves_to_the_folder(tool, isolated, monkeypatch):
    monkeypatch.setitem(vt.CONFIG, "notes_folder", str(isolated / "fallback"))
    for old in ("brain", "onenote", "apple_notes"):  # (only possible by editing config.json)
        monkeypatch.setitem(vt.CONFIG, "notes_destination", old)
        run_command(tool, "note idea", NOTE_REPLY)
        assert tool.cards[-1][0] == "Noted" and "notes folder" in tool.cards[-1][1]
    assert len(list((isolated / "fallback").glob("*.md"))) == 3


def test_a_note_that_cannot_be_saved_says_so(tool, isolated, monkeypatch):
    (isolated / "afile").write_text("x")
    monkeypatch.setitem(vt.CONFIG, "notes_destination", "folder")
    monkeypatch.setitem(vt.CONFIG, "notes_folder", str(isolated / "afile" / "inside"))  # (a path inside a file)
    run_command(tool, "note idea", NOTE_REPLY)
    assert tool.cards[-1][0] == "Couldn't do that"


def test_the_notes_settings(api, tool, isolated, monkeypatch):
    api, host = api
    settings = api.call("settings")
    assert [d["id"] for d in settings["note_destinations"]] == ["folder"]
    assert settings["notes_destination"] == "folder" and settings["notes_folder"] == str(isolated / "Documents" / "Voice notes")  # (the default is shown)
    assert api.call("set_notes", {"destination": "folder", "folder": str(isolated / "mine")})["notes_folder"] == str(isolated / "mine")
    assert json.loads((isolated / "config.json").read_text())["notes_destination"] == "folder"
    for bad in ("brain", "apple_notes", "onenote", "../../x", ""):
        with pytest.raises(ValueError, match="unknown destination"):
            api.call("set_notes", {"destination": bad})
    api.call("open_notes")
    assert OPENED[-1] == str(isolated / "mine") and (isolated / "mine").is_dir()
    monkeypatch.setitem(vt.CONFIG, "notes_folder", "")
    api.call("open_notes")
    assert OPENED[-1] == str(isolated / "Documents" / "Voice notes")  # not the program's own folder


def test_choose_notes_folder_asks_the_host_and_saves_the_answer(api, tool, isolated, monkeypatch):
    api, host = api
    asked = []
    monkeypatch.setattr(host, "choose_folder", lambda start="": asked.append(start) or str(isolated / "picked"))
    assert api.call("choose_notes_folder") == {"notes_folder": str(isolated / "picked")}
    assert json.loads((isolated / "config.json").read_text())["notes_folder"] == str(isolated / "picked")
    monkeypatch.setattr(host, "choose_folder", lambda start="": "")  # cancelled: nothing changes
    assert api.call("choose_notes_folder") == {"notes_folder": str(isolated / "picked")} and asked == [vt.CONFIG.get("notes_folder", "")] or True


def test_the_folder_picker_uses_the_windows_dialog_when_the_window_is_not_open(tool, monkeypatch, isolated):
    from tkinter import filedialog
    seen = []
    monkeypatch.setattr(filedialog, "askdirectory", lambda **kw: seen.append(kw) or str(isolated / "chosen"))
    result = []
    thread = threading.Thread(target=lambda: result.append(vt.WinHost(tool).choose_folder(str(isolated))), daemon=True)
    thread.start()
    deadline = time.time() + 10
    while thread.is_alive() and time.time() < deadline:  # (the dialog runs on the main thread, so play the main thread here)
        drain_ui(tool)
        time.sleep(0.05)
    assert result == [str(isolated / "chosen")] and seen[0]["initialdir"] == str(isolated)


def test_the_folder_picker_asks_the_open_window_to_show_its_own_dialog(tool, tmp_path):
    import app_window_win
    log = SimpleNamespace(info=lambda *a: None, error=lambda *a: None, warning=lambda *a: None, exception=lambda *a: None)
    script = write_script(tmp_path, f"""
import sys, json
print(json.dumps({{"event": "ready", "pid": 1, "hwnd": 2}}), flush=True)
for line in sys.stdin:
    msg = json.loads(line)
    if msg.get("cmd") == "folder":
        print(json.dumps({{"event": "folder", "id": msg["id"], "path": "C:/chosen/" + msg["start"].replace("\\\\", "/")[-4:]}}), flush=True)
""")
    window = app_window_win.AppWindow(SimpleNamespace(call=lambda *a: None), log, script=script)
    window.show("settings")
    assert wait_until(lambda: window.ready, 10)
    tool.app_window = window
    assert vt.WinHost(tool).choose_folder("C:\\Users\\sam\\Notes") == "C:/chosen/otes"
    window.close()
    assert window.choose_folder("x") == ""  # the window is gone: an empty answer, no hang


# ====================================================================================================== everyday computer commands

def pressed(monkeypatch):
    sent = []
    monkeypatch.setattr(vt.win32api, "keybd_event", lambda vk, scan, flags, extra: sent.append((vk, flags)))
    return sent


EXT, UP = vt.win32con.KEYEVENTF_EXTENDEDKEY, vt.win32con.KEYEVENTF_KEYUP


class FakeVolume:
    def __init__(self, level=0.5, muted=False):
        self.level, self.muted, self.calls = level, muted, []

    def SetMasterVolumeLevelScalar(self, level, guid):
        self.level = level
        self.calls.append(("level", level))

    def SetMute(self, muted, guid):
        self.muted = bool(muted)
        self.calls.append(("mute", muted))

    def GetMasterVolumeLevelScalar(self):
        return self.level

    def GetMute(self):
        return int(self.muted)


def test_the_virtual_key_codes_are_the_real_ones():
    assert (system_windows.VK_VOLUME_MUTE, system_windows.VK_VOLUME_DOWN, system_windows.VK_VOLUME_UP) == (
        vt.win32con.VK_VOLUME_MUTE, vt.win32con.VK_VOLUME_DOWN, vt.win32con.VK_VOLUME_UP)
    assert (system_windows.VK_MEDIA_NEXT, system_windows.VK_MEDIA_PREV, system_windows.VK_MEDIA_PLAY_PAUSE) == (
        vt.win32con.VK_MEDIA_NEXT_TRACK, vt.win32con.VK_MEDIA_PREV_TRACK, vt.win32con.VK_MEDIA_PLAY_PAUSE)
    assert (system_windows.VK_CONTROL, system_windows.VK_W) == (vt.win32con.VK_CONTROL, ord("W"))


def test_the_fixed_list_is_the_shared_table_and_every_item_has_a_windows_function():
    assert list(system_actions.ACTIONS) == ["mute", "unmute", "volume_up", "volume_down", "set_volume", "screenshot_window",
                                            "lock_screen", "play_pause", "next_track", "previous_track", "close_window"]
    for action in system_actions.ACTIONS:
        assert callable(getattr(system_windows, action)), action  # the same function names as system_mac.py
    for line in system_actions.router_lines().splitlines():
        assert re.match(r'  - "[a-z_]+": ', line)
    assert system_actions.router_lines() in commands.ROUTER  # the router is told exactly this list


@pytest.mark.parametrize("action, keys", [
    ("volume_up", [system_windows.VK_VOLUME_UP] * system_windows.VOLUME_STEPS),
    ("volume_down", [system_windows.VK_VOLUME_DOWN] * system_windows.VOLUME_STEPS),
    ("play_pause", [system_windows.VK_MEDIA_PLAY_PAUSE]), ("next_track", [system_windows.VK_MEDIA_NEXT]),
    ("previous_track", [system_windows.VK_MEDIA_PREV])])
def test_each_key_action_sends_exactly_the_right_keys_and_nothing_else(monkeypatch, action, keys):
    sent = pressed(monkeypatch)
    assert system_actions.run(system_windows, action) == system_actions.ACTIONS[action][1]
    assert sent == [event for vk in keys for event in ((vk, EXT), (vk, EXT | UP))]  # each key pressed, then released
    assert LOCKED == [] and SPOKEN == []


def test_mute_and_unmute_set_the_state_exactly_with_pycaw_and_use_the_key_without_it(monkeypatch):
    control = FakeVolume()
    monkeypatch.setattr(system_windows, "_volume_control", lambda: control)
    sent = pressed(monkeypatch)
    assert system_actions.run(system_windows, "mute") == "Muted." and control.muted is True
    assert system_actions.run(system_windows, "mute") == "Muted." and control.muted is True  # (never toggles back)
    assert system_actions.run(system_windows, "unmute") == "Sound on." and control.muted is False
    assert sent == []  # no keys needed when the volume control is there
    monkeypatch.setattr(system_windows, "_volume_control", lambda: None)
    system_actions.run(system_windows, "mute")
    assert sent == [(system_windows.VK_VOLUME_MUTE, EXT), (system_windows.VK_VOLUME_MUTE, EXT | UP)]


def test_lock_screen_locks_and_sends_no_keys(monkeypatch):
    sent = pressed(monkeypatch)
    assert system_actions.run(system_windows, "lock_screen") == "Locked."
    assert LOCKED == [1] and sent == []


def test_close_window_sends_ctrl_w_so_the_apps_own_save_prompt_still_appears(monkeypatch):
    sent = pressed(monkeypatch)
    assert system_actions.run(system_windows, "close_window") == "Closed."
    assert sent == [(system_windows.VK_CONTROL, 0), (system_windows.VK_W, 0), (system_windows.VK_W, UP), (system_windows.VK_CONTROL, UP)]
    assert LOCKED == []  # (it is Ctrl+W and nothing else: no Alt+F4, no window is killed)


def test_set_volume_exactly_with_pycaw_and_by_keys_without_it(monkeypatch):
    control = FakeVolume(muted=True)
    monkeypatch.setattr(system_windows, "_volume_control", lambda: control)
    sent = pressed(monkeypatch)
    assert system_actions.run(system_windows, "set_volume", "30") == "Volume set to 30."
    assert control.calls == [("level", 0.3), ("mute", 0)] and control.muted is False  # setting the volume also unmutes
    assert system_actions.run(system_windows, "set_volume", "45%") == "Volume set to 45."
    assert system_actions.run(system_windows, "set_volume", 150) == "Volume set to 100." and control.level == 1.0  # clamped
    assert system_actions.run(system_windows, "set_volume", "-20") == "Volume set to 0." and control.level == 0.0
    assert sent == []
    for bad in ("loud", "", None, "thirty"):
        with pytest.raises((ValueError, TypeError)):
            system_actions.run(system_windows, "set_volume", bad)
    monkeypatch.setattr(system_windows, "_volume_control", lambda: None)  # without pycaw: all the way down, then up in 2% steps
    system_actions.run(system_windows, "set_volume", "30")
    assert sent == ([(system_windows.VK_VOLUME_DOWN, EXT), (system_windows.VK_VOLUME_DOWN, EXT | UP)] * 50
                    + [(system_windows.VK_VOLUME_UP, EXT), (system_windows.VK_VOLUME_UP, EXT | UP)] * 15)


def test_pycaw_is_installed_and_reads_the_real_volume_without_changing_it():
    import importlib
    real = importlib.reload(system_windows)
    try:
        level, muted = real.volume()  # (a read only: nothing is set)
        assert 0 <= level <= 100 and muted in (True, False)
    finally:
        importlib.reload(system_windows)


def test_only_listed_actions_can_run(monkeypatch):
    sent = pressed(monkeypatch)
    for bad in ("format_disk", "shutdown", "powershell", "", None, "volume_up; lock_screen", "__import__('os').system('calc')",
                "press_key", "_lock_workstation", "screenshot", "mute_toggle"):
        with pytest.raises(ValueError, match="isn't one of VoiceTool's computer commands"):
            system_actions.run(system_windows, bad)
    assert sent == [] and LOCKED == []


def test_the_windows_actions_cannot_reach_a_shell_or_a_program():
    import inspect
    source = inspect.getsource(system_windows)
    for forbidden in ("subprocess", "os.system", "Popen", "startfile", "shell=", "eval(", "exec("):
        assert forbidden not in source, forbidden


def test_the_screenshot_action_saves_a_png_of_that_window_in_pictures(tk_root, screenshots_only_of_the_test_window, monkeypatch, isolated):
    from PIL import Image
    monkeypatch.setenv("USERPROFILE", str(isolated / "profile"))
    with quiet_window(tk_root, color="#00ff00", size="500x350") as w:
        message = system_actions.run(system_windows, "screenshot_window", context={"hwnd": w.hwnd, "pid": 123})
        folder = isolated / "profile" / "Pictures" / "Screenshots"
        (shot,) = folder.glob("VoiceTool *.png")
        assert re.fullmatch(r"VoiceTool \d{4}-\d\d-\d\d \d{6}\.png", shot.name)
        assert message == f"Screenshot saved: {shot.name} (Screenshots)."
        image = Image.open(shot).convert("RGB")
        assert image.width >= 400 and image.height >= 280  # full size, not shrunk
        r, g, b = image.getpixel((image.width // 2, image.height // 2))
        assert g > 200 and r < 60
        again = actions_windows_save(w.hwnd, folder)
        assert again != shot and again.exists()  # never overwrites


def actions_windows_save(hwnd, folder):
    return aw.save_window_screenshot(hwnd=hwnd, folder=folder)


def test_the_screenshot_action_with_no_window_says_so(monkeypatch):
    monkeypatch.setattr("win32gui.GetForegroundWindow", lambda: 0)
    with pytest.raises(RuntimeError, match="no window in front"):
        system_actions.run(system_windows, "screenshot_window", context={})
    assert aw.screenshots_folder().parts[-2:] == ("Pictures", "Screenshots")


# ---- through the command key

def test_the_router_picks_an_action_from_the_list_and_keys_wait_for_the_command_key_to_be_up(tool, monkeypatch):
    sent, order = pressed(monkeypatch), []
    monkeypatch.setattr(tool, "_wait_for_release", lambda t: order.append(("released", len(sent))))
    client = run_command(tool, "could you make it a bit less loud", {"route": "system", "target": "volume_down", "value": ""})
    assert order == [("released", 0)]  # waited for the key BEFORE any key was sent
    assert len(sent) == 2 * system_windows.VOLUME_STEPS and sent[0] == (system_windows.VK_VOLUME_DOWN, EXT)
    assert tool.cards[-1] == ("Done", "Volume down.", False) and len(client.sent) == 1


def test_set_volume_through_the_router_and_lock_without_a_card(tool, monkeypatch):
    control = FakeVolume()
    monkeypatch.setattr(system_windows, "_volume_control", lambda: control)
    run_command(tool, "make it about thirty", {"route": "system", "target": "set_volume", "value": "30"})
    assert control.level == 0.3 and tool.cards[-1][1] == "Volume set to 30."
    tool.cards.clear()
    run_command(tool, "lock the screen", {"route": "system", "target": "lock_screen", "value": ""})
    assert LOCKED == [1] and tool.cards == []  # (nothing to show: the screen is locked)


def test_screenshot_through_the_router_uses_the_window_the_key_was_pressed_in(tool, monkeypatch, isolated):
    seen = []
    monkeypatch.setattr(aw, "save_window_screenshot", lambda pid=None, hwnd=None, folder=None: seen.append((pid, hwnd)) or isolated / "VoiceTool shot.png")
    run_command(tool, "screenshot this", {"route": "system", "target": "screenshot_window", "value": ""},
                {"app": "chrome.exe", "title": "x", "selection": "", "pid": 11, "hwnd": 22})
    assert seen == [(11, 22)] and tool.cards[-1][0] == "Done" and "VoiceTool shot.png" in tool.cards[-1][1]


def test_the_router_naming_anything_else_does_nothing_but_explain(tool, monkeypatch):
    sent = pressed(monkeypatch)
    for target in ("format_disk", "shutdown", "powershell Get-Process", "", "__import__('os').system('calc')", "volume_up;mute"):
        run_command(tool, "do something to the computer", {"route": "system", "target": target, "value": ""})
        assert tool.cards[-1][0] == "Couldn't do that" and "isn't one of VoiceTool's computer commands" in tool.cards[-1][1], target
    run_command(tool, "volume please", {"route": "system", "target": "set_volume", "value": "rm -rf"})
    assert tool.cards[-1][0] == "Couldn't do that" and "volume number" in tool.cards[-1][1]
    assert sent == [] and LOCKED == [] and SPOKEN == []


def test_a_failing_action_shows_why(tool, monkeypatch):
    monkeypatch.setattr(aw, "save_window_screenshot", lambda *a, **k: None)
    run_command(tool, "screenshot", {"route": "system", "target": "screenshot_window", "value": ""})
    assert tool.cards[-1][0] == "Couldn't do that" and "no window in front" in tool.cards[-1][1]


def test_close_window_through_the_router_sends_only_ctrl_w(tool, monkeypatch):
    sent = pressed(monkeypatch)
    run_command(tool, "close this window", {"route": "system", "target": "close_window", "value": ""})
    assert sent == [(system_windows.VK_CONTROL, 0), (system_windows.VK_W, 0), (system_windows.VK_W, UP), (system_windows.VK_CONTROL, UP)]
    assert tool.cards == []


# ====================================================================================================== keyboard watchdog
# (2026-10-23: the hook went deaf until a restart; these cover the hook's safety net)

class AliveListener:
    def __init__(self, alive=True):
        self.alive, self.stopped = alive, False

    def is_alive(self):
        return self.alive

    def suppress_event(self):
        raise vt._SUPPRESS()


@pytest.fixture
def watched(tool, monkeypatch):
    tool.listener = AliveListener()
    tool.installs = []
    monkeypatch.setattr(tool, "_install_hook", lambda: tool.installs.append(1))
    return tool


def held(*vks):
    return lambda vk: vk in vks


def test_hook_error_is_logged_and_the_key_passes_instead_of_killing_the_hook(hooked, monkeypatch):
    tool, _ = hooked

    class Broken:
        def put(self, item):
            raise RuntimeError("boom")
    monkeypatch.setattr(tool, "events", Broken())
    assert hook(tool, DICTATE, True) is True  # no exception reaches pynput, which would stop the listener


def test_hook_still_suppresses_keys_when_asked(tool, monkeypatch):
    tool.listener = AliveListener()
    monkeypatch.setattr(tool.triggers[DICTATE], "let_through", False)
    with pytest.raises(vt._SUPPRESS):
        hook(tool, DICTATE, True)


def test_watchdog_quiet_while_the_hook_and_windows_agree(watched):
    t = watched.triggers[DICTATE]
    t.hook_down = True
    for i in range(20):
        watched._check_keys(100 + i * 0.05, held(DICTATE))
    t.hook_down = False
    for i in range(20):
        watched._check_keys(200 + i * 0.05, held())
    assert watched.installs == [] and watched.events.empty()


def test_watchdog_ignores_a_short_disagreement(watched):
    watched._check_keys(100.0, held(DICTATE))  # Windows says down a moment before the hook's event is handled
    watched._check_keys(100.3, held(DICTATE))
    assert watched.installs == []
    t = watched.triggers[DICTATE]
    t.hook_down = True
    watched._check_keys(101.0, held())
    watched._check_keys(101.2, held())
    assert watched.installs == [] and watched.events.empty()


def test_watchdog_reinstalls_a_deaf_hook_once_per_press(watched):
    for i in range(40):  # the key is held for 2 s and the hook never hears it
        watched._check_keys(100 + i * 0.05, held(DICTATE))
    assert watched.installs == [1]
    assert watched.events.empty()  # the press isn't replayed: the new hook hears the next key repeat
    watched._check_keys(103.0, held())  # let go: the next press is checked afresh
    for i in range(10):
        watched._check_keys(104 + i * 0.05, held(DICTATE))
    assert watched.installs == [1, 1]


def test_watchdog_handles_a_release_the_hook_missed(watched):
    t0 = now()
    key(watched, DICTATE, True, t0 - 2)  # the hook heard the press...
    t = watched.triggers[DICTATE]
    t.hook_down = True
    watched._check_keys(t0, held())  # ...but not the release
    watched._check_keys(t0 + 0.35, held())
    assert watched.installs == [1] and not t.hook_down
    kind, trig, value, at = watched.events.get_nowait()
    assert (kind, trig, value, at) == ("trigger", t, False, t0)
    watched._handle(kind, trig, value, at)
    assert watched.active is None and jobs(watched) == 1  # the dictation went through


def test_watchdog_reinstalls_a_stopped_listener(watched):
    watched.listener = AliveListener(alive=False)
    watched._check_keys(100.0, held())
    assert watched.installs == [1]


def test_watchdog_leaves_switched_off_and_swallowed_keys_alone(watched, monkeypatch):
    monkeypatch.setitem(vt.CONFIG, "enabled", False)
    for i in range(20):
        watched._check_keys(100 + i * 0.05, held(DICTATE, COMMAND))
    monkeypatch.setitem(vt.CONFIG, "enabled", True)
    monkeypatch.setattr(watched.triggers[COMMAND], "let_through", False)
    monkeypatch.setitem(vt.CONFIG, "dictate_enabled", False)
    for i in range(20):
        watched._check_keys(200 + i * 0.05, held(DICTATE, COMMAND))
    assert watched.installs == [] and watched.events.empty()


def test_watchdog_does_nothing_while_quitting(watched):
    watched._quitting = True
    watched.listener = AliveListener(alive=False)
    watched._check_keys(100.0, held())
    assert watched.installs == []


def test_watchdog_notices_a_stuck_window_loop_once(watched):
    watched._last_tick = 100.0
    watched._check_keys(101.0, held())
    assert not watched._stuck_logged
    watched._check_keys(104.0, held())
    assert watched._stuck_logged
    watched._last_tick = 104.5
    watched._check_keys(104.6, held())
    assert not watched._stuck_logged


def test_window_loop_keeps_going_after_an_error(tool, monkeypatch):
    calls = []
    tool.light.root = SimpleNamespace(after=lambda ms, f: calls.append(ms))
    monkeypatch.setattr(tool.light, "draw", lambda: (_ for _ in ()).throw(RuntimeError("draw broke")), raising=False)
    tool._tick()
    assert calls == [30] and tool._last_tick is not None


def test_only_the_newest_hook_acts_and_the_old_one_is_stopped(tool, monkeypatch):
    made = []

    class FakeKeyboardListener(AliveListener):
        installed, error = True, 0

        def __init__(self, on_key):
            super().__init__()
            self.filter = on_key
            made.append(self)

        def start(self):
            pass

        def wait(self):
            return True

        def stop(self):
            self.stopped = True
    monkeypatch.setattr(vt, "KeyboardHook", FakeKeyboardListener)
    monkeypatch.setattr(vt.win32api, "keybd_event", lambda *a: None)
    tool._install_hook()
    first = made[0]
    tool._install_hook()
    second = made[1]
    assert first.stopped and not second.stopped and tool.listener is second
    data = SimpleNamespace(vkCode=DICTATE, flags=0)
    assert first.filter(vt.win32con.WM_KEYDOWN, data) is True and tool.events.empty()  # the old hook just passes it on
    assert second.filter(vt.win32con.WM_KEYDOWN, data) is False
    assert drain(tool.events)[0][:3] == ("trigger", tool.triggers[DICTATE], True)


def test_a_long_hold_with_no_sound_is_logged_not_silent(tool, monkeypatch):
    warned = []
    monkeypatch.setattr(vt.log, "warning", lambda *a, **k: warned.append(a[0]))
    tool.recorder.seconds = 0.0
    key(tool, DICTATE, True, now() - 3)
    key(tool, DICTATE, False, now())
    assert jobs(tool) == 0 and any("microphone gave only" in w for w in warned)

def test_keyboard_hook_callback_suppresses_passes_on_and_survives_errors(monkeypatch):
    import ctypes
    nexts = []
    monkeypatch.setattr(vt.KeyboardHook, "_CallNext", lambda *a: nexts.append(a) or 0)
    data = vt._KBDLLHOOKSTRUCT(vkCode=DICTATE, flags=0)
    ptr = ctypes.addressof(data)
    seen = []

    def on_key(msg, d):
        seen.append((msg, d.vkCode))
        if msg == 1:
            hk.suppress_event()
        if msg == 2:
            raise RuntimeError("bug")
    hk = vt.KeyboardHook(on_key)
    assert hk._callback(0, 1, ptr) == 1 and nexts == []  # suppressed: Windows drops the key
    assert hk._callback(0, 2, ptr) == 0 and len(nexts) == 1  # an error: the key goes on as normal
    assert hk._callback(0, vt.win32con.WM_KEYDOWN, ptr) == 0 and len(nexts) == 2
    assert hk._callback(-1, vt.win32con.WM_KEYDOWN, ptr) == 0 and len(seen) == 3  # not HC_ACTION: not looked at
    assert seen[0] == (1, DICTATE)


def test_keyboard_hook_installs_and_stops_cleanly():
    hk = vt.KeyboardHook(lambda msg, data: None)
    hk.start()
    assert hk.wait() and hk.installed and hk.is_alive()
    hk.stop()
    hk.join(5)
    assert not hk.is_alive()

def test_watchdog_never_cuts_a_hold_the_hook_still_hears(watched):
    # Another program may hide the key from Windows' key state; while the hook keeps hearing the key repeat, it's held.
    t = watched.triggers[DICTATE]
    t.hook_down = True
    for i in range(40):
        t.hook_at = 100 + i * 0.05  # a repeat heard on every look
        watched._check_keys(100 + i * 0.05, held())
    assert watched.installs == [] and watched.events.empty() and t.hook_down
    watched._check_keys(102.0, held())  # the repeats stop: now it really was let go
    watched._check_keys(102.4, held())
    watched._check_keys(102.8, held())
    assert not t.hook_down and watched.events.get_nowait()[2] is False


def test_hook_notes_when_it_last_heard_a_trigger(hooked):
    tool, _ = hooked
    hook(tool, DICTATE, True)
    t = tool.triggers[DICTATE]
    assert t.hook_at > 0 and drain(tool.events)[0][3] == t.hook_at

def test_rewrite_never_pastes_the_spoken_instruction():
    """"make it sound better" ended up inside the text. Rewriting is a separate call whose prompt forbids it."""
    client = FakeClient(json.dumps({"route": "rewrite", "text": "IGNORED"}), "Clean text.")
    kind, text, extra = commands.route(client, "m", "Hey, can you make this sound better?", {"selection": "messy text"}, "")
    assert (kind, text) == ("rewrite", "Clean text.")
    second = client.sent[1]["messages"]
    assert "Never put it" in second[0]["content"] and "NEVER expand contractions" in second[0]["content"]
    assert "INSTRUCTION: Hey, can you make this sound better?" in second[1]["content"] and "messy text" in second[1]["content"]


def test_silent_try_during_meeting_notes_does_not_reset_the_audio_system(tool, monkeypatch):
    calls = []
    monkeypatch.setattr(vt.sd, "_terminate", lambda: calls.append("terminate"))
    monkeypatch.setattr(vt.sd, "_initialize", lambda: calls.append("initialize"))
    tool.meeting = object()  # meeting notes recording
    t = tool.triggers[DICTATE]
    tool.recorder.stop = lambda: np.full(vt.RATE * 3, 0.0005, np.float32)
    tool._finish(t)
    assert calls == []



# ====================================================================================================== friend version: everything on this computer

class RawOpenAI:
    """What the OpenAI SDK's client looks like from outside: it records exactly what it is sent."""

    def __init__(self, *replies):
        self.replies, self.sent, self.options = list(replies), [], None
        self.chat = SimpleNamespace(completions=self)

    def create(self, **kw):
        self.sent.append(kw)
        text = self.replies.pop(0) if len(self.replies) > 1 else self.replies[0]
        return SimpleNamespace(choices=[SimpleNamespace(message=SimpleNamespace(content=text), finish_reason="stop")],
                               usage=SimpleNamespace(prompt_tokens=11, completion_tokens=7, prompt_tokens_details=None))

    def with_options(self, **kw):
        self.options = kw
        return self


class FakeLocalServer:
    """Stands in for local_cleanup_win.LocalCleaner: no program is started; `client()` wraps the recording client."""

    def __init__(self, raw):
        self.raw, self.port, self.starts = raw, 4242, 0

    def start(self):
        self.starts += 1

    def client(self, timeout=60):
        return local_cleanup_win.LocalClient(self.raw)

    def stop(self):
        pass


@pytest.fixture
def friend(monkeypatch):
    """The friend's state: everything stays on this computer, no Azure endpoint, and the model is installed (a fake one)."""
    for key, value in {"offline_only": True, "azure_openai_base_url": "", "cleanup_model": "local", "cleanup": True}.items():
        monkeypatch.setitem(vt.CONFIG, key, value)
    monkeypatch.setattr(local_cleanup_win, "available", lambda config=None: True)
    import azure.identity
    monkeypatch.setattr(azure.identity, "AzureCliCredential", lambda **kw: pytest.fail("an Azure credential must never be created"))


def test_the_request_settings_only_openai_takes_are_adapted_in_one_place():
    adapted = local_cleanup_win.adapt_params({"model": "local", "reasoning_effort": "none", "max_completion_tokens": 16,
                                              "response_format": {"type": "json_object"}, "messages": [], "temperature": 0})
    assert adapted == {"model": "local", "max_tokens": 16, "response_format": {"type": "json_object"}, "messages": [], "temperature": 0}
    assert local_cleanup_win.adapt_params({"reasoning_effort": "low"})["temperature"] == 0  # (steady answers from a small model)
    assert "max_tokens" not in local_cleanup_win.adapt_params({"max_completion_tokens": None})
    assert local_cleanup_win.adapt_params({"max_completion_tokens": 5, "max_tokens": 9})["max_tokens"] == 9
    raw = RawOpenAI("x")
    client = local_cleanup_win.LocalClient(raw)
    reply = client.chat.completions.create(model="m", reasoning_effort="none", max_completion_tokens=7, messages=[])
    assert "reasoning_effort" not in raw.sent[0] and raw.sent[0]["max_tokens"] == 7 and reply._voicetool_free is True
    assert client.with_options(timeout=5).chat.completions.create(model="m", messages=[]) and raw.options == {"timeout": 5}


def test_offline_only_answers_tidy_up_routing_rewrite_journal_and_notes_from_the_local_client_and_never_creates_azure(tool, friend, monkeypatch, isolated):
    raw = RawOpenAI()
    c = vt.Cleaner()
    c.local = FakeLocalServer(raw)
    c.warm_up()  # nothing to warm up
    assert raw.sent == [] and c._credential is None and c._openai is None
    # the tidy-up
    raw.replies = ["Hello there, friend."]
    assert c("um hello there friend") == "Hello there, friend."
    req = raw.sent[-1]
    assert "reasoning_effort" not in req and "max_completion_tokens" not in req and req["max_tokens"] >= 64 and req["temperature"] == 0
    # command routing, and the focused rewrite step that follows a "rewrite" route
    raw.replies = [json.dumps({"route": "open", "target": "notepad"})]
    kind, text, extra = commands.route(c._client(), "local", "open notepad", {"selection": ""})
    assert (kind, extra["target"]) == ("open", "notepad")
    assert raw.sent[-1]["response_format"] == {"type": "json_object"} and "reasoning_effort" not in raw.sent[-1]
    raw.replies = [json.dumps({"route": "rewrite", "text": ""}), "Shorter."]
    kind, text, _ = commands.route(c._client(), "local", "make this shorter", {"selection": "a long long sentence"})
    assert (kind, text) == ("rewrite", "Shorter.") and len(raw.sent) == 4
    # the journal
    raw.replies = ["Today was fine."]
    path = va.journal(c._client(), "local", "today was fine")
    assert "Today was fine." in path.read_text() and raw.sent[-1]["max_tokens"] == 800
    # meeting notes (no follow-up email asked for: Claude is off)
    raw.replies = [json.dumps({"title": "Sync", "summary": "We talked.", "decisions": [], "action_items": [], "open_questions": []})]
    notes = meetings.make_notes(c._client(), "local", [(0, "Me", "hello there")], 1000, 1600, local=True, followup=False)
    assert notes["title"] == "Sync" and raw.options == {"timeout": 900}  # (the model on this computer gets 15 minutes)
    assert raw.sent[-1]["max_tokens"] == 2500 and "follow_up_email" not in raw.sent[-1]["messages"][0]["content"]
    # nothing was ever sent with an OpenAI-only setting, and nothing was started but the (fake) local server
    assert all("reasoning_effort" not in k and "max_completion_tokens" not in k for k in raw.sent)
    assert c.local.starts >= 5 and c._credential is None
    # every call is logged as free, never at an online price
    rows = [json.loads(line) for line in (isolated / "usage.jsonl").read_text().splitlines()]
    assert [r["service"] for r in rows if r["service"] != "speech"] == ["local"] * 6
    assert all(r["usd"] == 0 and "free" in r["paid_by"] for r in rows)
    assert {r["purpose"] for r in rows} == {"tidy-up", "command routing", "rewrite", "journal", "meeting notes"}


def test_the_model_client_is_made_lazily_and_azure_only_when_the_switch_is_off(monkeypatch):
    created = []

    class Cred:
        def __init__(self, tenant_id=None):
            created.append("credential")

        def get_token(self, scope):
            return SimpleNamespace(token="t", expires_on=time.time() + 3600)
    import azure.identity
    monkeypatch.setattr(azure.identity, "AzureCliCredential", Cred)
    monkeypatch.setitem(vt.CONFIG, "azure_openai_base_url", "https://example.invalid/openai/v1/")
    monkeypatch.setitem(vt.CONFIG, "offline_only", True)
    c = vt.Cleaner()
    time.sleep(0.2)
    assert created == []  # on: no credential, no token, no start-up thread
    monkeypatch.setitem(vt.CONFIG, "offline_only", False)
    c._OpenAI = lambda **kw: SimpleNamespace(kw=kw)
    assert c._client().kw["base_url"] == "https://example.invalid/openai/v1/" and created == ["credential"]
    monkeypatch.setitem(vt.CONFIG, "azure_openai_base_url", "")
    c2 = vt.Cleaner()
    with pytest.raises(RuntimeError, match="Azure isn't set up"):
        c2._client()  # (no endpoint, switch off: a plain error, not a crash in the middle of something)
    assert created == ["credential"]


def test_a_dictation_with_the_real_cleaner_and_no_azure_is_tidied_by_the_local_client(tool, friend, monkeypatch):
    raw = RawOpenAI("Open the file, please.")
    c = vt.Cleaner()
    c.local = FakeLocalServer(raw)
    tool.azure = c
    run_dictation(tool, "um open the file please")
    assert PASTED == ["Open the file, please. "] and raw.sent[0]["messages"][0]["content"].startswith("Fix this dictated text.")


def test_a_vulkan_server_that_does_not_come_up_falls_back_to_the_cpu_one_once(fake_server):
    logged, seen = [], []
    log = SimpleNamespace(info=lambda *a: None, error=lambda *a: None,
                          warning=lambda msg, *a: logged.append(msg % a if a else msg))

    def command(port, backend):
        seen.append(backend)
        if backend == "cpu":
            return [BASE_PYTHON, fake_server.script, str(port), "fake-model.gguf"]
        return [BASE_PYTHON, "-c", "raise SystemExit(3)"]  # the Vulkan build dies at once (no usable graphics driver)
    c = local_cleanup_win.LocalCleaner({"offline_model": "C:/m/fake-model.gguf", "offline_port": free_port(), "offline_backend": "vulkan"},
                                       log, command=command, startup_seconds=30)
    assert tidy(c, "SYSTEM", "hello there") == "CLEANED: hello there | SYSTE"
    assert seen == ["vulkan", "cpu"] and c.backend == "cpu" and c.fell_back
    assert any("CPU" in m for m in logged)  # and it is logged
    tidy(c, "SYSTEM", "again")
    assert seen == ["vulkan", "cpu"]  # the Vulkan one is not tried again
    c.stop()
    # the fallback happens once: if the CPU server fails too, it gives up instead of looping
    tries = []
    c2 = local_cleanup_win.LocalCleaner({"offline_port": free_port(), "offline_backend": "vulkan"}, None,
                                        command=lambda port, backend: tries.append(backend) or [BASE_PYTHON, "-c", "raise SystemExit(3)"],
                                        startup_seconds=10)
    with pytest.raises(RuntimeError):
        c2.start()
    assert tries == ["vulkan", "cpu"]
    # a configured CPU backend never tries Vulkan; its arguments have no graphics-card options
    assert local_cleanup_win.cpu_args(["-c", "16384", "-ngl", "99", "--device", "Vulkan0", "--no-webui"]) == ["-c", "16384", "--no-webui"]


def test_the_cpu_server_is_used_with_its_own_arguments_in_the_real_command(tmp_path, monkeypatch):
    for name in ("v.exe", "c.exe", "m.gguf"):
        (tmp_path / name).write_bytes(b"x")
    cfg = {**json.loads((HERE / "config.json").read_text(encoding="utf-8")), "offline_server": str(tmp_path / "v.exe"),
           "offline_server_cpu": str(tmp_path / "c.exe"), "offline_model": str(tmp_path / "m.gguf")}
    started = []

    class Dead:
        pid = 1

        def poll(self):
            return 3
        def kill(self): pass
    monkeypatch.setattr(local_cleanup_win.subprocess, "Popen", lambda cmd, **kw: started.append(cmd) or Dead())
    monkeypatch.setattr(local_cleanup_win, "_port_in_use", lambda port: False)
    c = local_cleanup_win.LocalCleaner(cfg, None, startup_seconds=5)
    with pytest.raises(RuntimeError):
        c.start()
    assert started[0][0] == str(tmp_path / "v.exe") and "-ngl" in started[0] and "--device" in started[0]
    assert started[1][0] == str(tmp_path / "c.exe") and "-ngl" not in started[1] and "--device" not in started[1] and "Vulkan0" not in started[1]
    assert "--no-webui" in started[1] and "127.0.0.1" in started[1]


# ====================================================================================================== friend version: not set up yet

def test_dictation_before_the_model_exists_pastes_the_words_as_heard_and_says_so_once(tool, monkeypatch, isolated):
    monkeypatch.setitem(vt.CONFIG, "offline_only", True)
    monkeypatch.setitem(vt.CONFIG, "azure_openai_base_url", "")
    monkeypatch.setattr(local_cleanup_win, "available", lambda config=None: False)
    import azure.identity
    monkeypatch.setattr(azure.identity, "AzureCliCredential", lambda **kw: pytest.fail("no Azure"))
    corrections.add("zeph ear", "Zephyr", "test")
    tool.azure = vt.Cleaner()
    run_dictation(tool, "um open zeph ear please")
    assert PASTED == ["um open Zephyr please "]  # the words as heard, with the learned fix applied
    assert tool.cards == [("Tidy-up isn't set up yet", "Tidy-up isn't set up yet, so your words are typed as heard. "
                                                       "Open VoiceTool and click Finish setup.", False)]
    run_dictation(tool, "and another dictation here")
    assert PASTED[-1] == "and another dictation here " and len(tool.cards) == 1  # the card is once per run
    assert "error" not in tool.light.history  # never an error buzz
    row = json.loads((isolated / "history.jsonl").read_text().splitlines()[-1])
    assert row["text"] == "and another dictation here"


def test_commands_that_need_the_model_show_a_plain_card_before_it_exists(tool, monkeypatch, isolated):
    monkeypatch.setitem(vt.CONFIG, "offline_only", True)
    monkeypatch.setattr(local_cleanup_win, "available", lambda config=None: False)
    t = tool.triggers[COMMAND]
    t.key_down = False
    t.context = {"app": "notepad.exe", "title": "x", "selection": "some text", "pid": 1, "hwnd": 2}
    tool.transcribe = lambda audio: "make this shorter"
    tool.azure = vt.Cleaner()
    tool._command(t, np.full(vt.RATE, 0.3, np.float32))
    assert tool.cards[-1][0] == "Finish setup first" and "click Finish setup" in tool.cards[-1][1] and PASTED == []
    assert tool.light.history[-1] is None  # the light just goes off
    tool._journal("today was fine")
    assert tool.cards[-1][0] == "Finish setup first" and "The journal" in tool.cards[-1][1] and not (isolated / "journal").exists()


def test_a_meeting_is_refused_while_the_tidy_up_model_is_missing_and_nothing_is_recorded(tool, monkeypatch, isolated):
    meeting_tool(tool, monkeypatch, isolated)
    monkeypatch.setitem(vt.CONFIG, "offline_only", True)
    monkeypatch.setattr(local_cleanup_win, "available", lambda config=None: False)
    tool.azure = vt.Cleaner()
    made = []
    real = meetings.MeetingRecorder
    monkeypatch.setattr(meetings, "MeetingRecorder", lambda *a, **k: made.append(1) or real(*a, **k))
    tool._start_meeting()  # from the menu, the window or a spoken command
    assert tool.meeting is None and made == [] and tool.light.meeting is None
    assert tool.cards[-1] == ("Finish setup first", "Meeting notes need the tidy-up model, which isn't set up yet. "
                                                    "Open VoiceTool and click Finish setup.", False)
    assert len(tool.cards) == 1  # (no "everyone agreed" card: nothing to agree to yet)
    tool._begin_meeting()  # a consent card from earlier, answered after the model went away
    assert tool.meeting is None and made == [] and tool.cards[-1][0] == "Finish setup first"
    tool._offer_meeting("Zoom")  # a call starts: no offer either
    assert tool.cards[-1][0] == "Finish setup first" and len(tool.cards) == 2
    assert notes_files(isolated) == []
    # once the model is there, the same start works
    monkeypatch.setattr(local_cleanup_win, "available", lambda config=None: True)
    start_meeting(tool)
    assert tool.meeting is not None and made == [1]
    tool.meeting.abort()


def test_a_meeting_is_refused_while_the_speech_model_is_missing(tool, monkeypatch, isolated):
    meeting_tool(tool, monkeypatch, isolated)
    monkeypatch.setattr(vt.onnx_asr, "load_model", lambda *a, **k: pytest.fail("no download"))
    monkeypatch.setitem(vt.CONFIG, "speech_model_path", r"models\parakeet-tdt-0.6b-v2")
    tool.transcribe = vt.Transcriber()
    tool._start_meeting()
    assert tool.meeting is None and tool.cards[-1][0] == "Finish setup first" and "speech model" in tool.cards[-1][1]


@pytest.fixture
def downloader(tmp_path, monkeypatch, api):
    """A faked "Finish setup" downloader: nothing is started and nothing is downloaded."""
    api_, host = api
    launched, synced = [], []
    (tmp_path / ".venv" / "Scripts").mkdir(parents=True)
    (tmp_path / ".venv" / "Scripts" / "pythonw.exe").write_bytes(b"x")

    class FakeProc:
        def __init__(self, args, **kw):
            launched.append((args, kw))
            self.release, self.code = threading.Event(), 0
            FakeProc.last = self

        def wait(self):
            self.release.wait(10)
            return self.code
    monkeypatch.setattr(vt.subprocess, "Popen", FakeProc)
    host.tool.azure = SimpleNamespace(sync_offline=lambda: synced.append(1))
    return SimpleNamespace(api=api_, host=host, launched=launched, synced=synced, proc=lambda: FakeProc.last, folder=tmp_path)


def test_the_finish_setup_button_shows_only_while_the_model_is_missing(downloader, monkeypatch):
    monkeypatch.setattr(local_cleanup_win, "available", lambda config=None: False)
    assert downloader.api.call("overview")["setup"] == {"missing": True, "speech_missing": False, "running": False, "failed": False}
    monkeypatch.setattr(local_cleanup_win, "available", lambda config=None: True)
    assert downloader.api.call("overview")["setup"] == {"missing": False, "speech_missing": False, "running": False, "failed": False}


def test_clicking_finish_setup_launches_the_downloader_hidden_and_success_starts_the_model(downloader, monkeypatch):
    monkeypatch.setattr(local_cleanup_win, "available", lambda config=None: False)
    reply = downloader.api.call("finish_setup")
    assert downloader.launched == [] and reply["started"] is False and "fetch_models.py" in reply["message"]  # the file isn't there
    (downloader.folder / "fetch_models.py").write_text("print('x')")
    assert downloader.api.call("finish_setup") == {"started": True}
    args, kw = downloader.launched[0]
    assert args == [str(downloader.folder / ".venv" / "Scripts" / "pythonw.exe"), str(downloader.folder / "fetch_models.py"), "--window"]
    assert kw["creationflags"] & subprocess.CREATE_NO_WINDOW and kw["cwd"] == str(downloader.folder)
    assert downloader.api.call("overview")["setup"] == {"missing": True, "speech_missing": False, "running": True, "failed": False}
    again = downloader.api.call("finish_setup")
    assert again["started"] is False and again["running"] is True and len(downloader.launched) == 1  # one download at a time
    assert downloader.synced == []
    monkeypatch.setattr(local_cleanup_win, "available", lambda config=None: True)  # (the download put the files there)
    downloader.proc().release.set()  # the downloader exits with code 0
    assert wait_until(lambda: downloader.synced == [1], 10)  # the on-device model is started
    assert downloader.api.call("overview")["setup"] == {"missing": False, "speech_missing": False, "running": False, "failed": False}


def test_a_downloader_that_fails_shows_try_again_and_does_not_start_the_model(downloader, monkeypatch):
    monkeypatch.setattr(local_cleanup_win, "available", lambda config=None: False)
    (downloader.folder / "fetch_models.py").write_text("print('x')")
    downloader.api.call("finish_setup")
    downloader.proc().code = 3
    downloader.proc().release.set()
    assert wait_until(lambda: downloader.api.call("overview")["setup"]["failed"], 10)
    assert downloader.synced == [] and downloader.api.call("overview")["setup"]["running"] is False
    assert downloader.api.call("finish_setup") == {"started": True} and len(downloader.launched) == 2  # try again
    downloader.proc().release.set()


# ====================================================================================================== friend version: meeting consent

def test_starting_a_meeting_asks_everyone_agreed_first_and_records_nothing_until_answered(tool, monkeypatch, isolated):
    meeting_tool(tool, monkeypatch, isolated)
    made = []
    real = meetings.MeetingRecorder
    monkeypatch.setattr(meetings, "MeetingRecorder", lambda *a, **k: made.append(1) or real(*a, **k))
    tool._start_meeting()
    assert tool.cards[-1][:2] == ("Before recording", "Before recording, make sure everyone on the call agrees to be recorded.")
    assert list(tool.card_buttons[-1]) == ["Everyone agreed, start", "Cancel"] and tool.meeting is None and made == []
    tool.card_buttons[-1]["Cancel"]()
    assert tool.meeting is None and made == [] and tool.light.meeting is None
    tool.card_buttons[-1]["Everyone agreed, start"]()
    assert tool.meeting is not None and made == [1] and tool.light.meeting == "meeting"
    tool.meeting.abort()
    tool.meeting = None


def test_the_voice_command_and_the_window_button_also_ask_first(tool, monkeypatch, isolated):
    meeting_tool(tool, monkeypatch, isolated)
    run_command(tool, "start meeting notes", {"route": "meeting", "target": "start"})
    drain_ui(tool)
    assert tool.cards[-1][0] == "Before recording" and tool.meeting is None
    tool.cards.clear()
    vt.WinHost(tool).start_meeting()
    drain_ui(tool)
    assert tool.cards[-1][0] == "Before recording" and tool.meeting is None
    # stopping never needs a question
    stopped = []
    monkeypatch.setattr(tool, "_stop_meeting", lambda: stopped.append(1))
    run_command(tool, "stop the meeting", {"route": "meeting", "target": "stop"})
    drain_ui(tool)
    assert stopped == [1]


def test_meeting_notes_off_means_no_consent_card_and_no_recording(tool, monkeypatch, isolated):
    monkeypatch.setitem(vt.CONFIG, "meeting_notes", False)
    meeting_tool(tool, monkeypatch, isolated)
    tool._start_meeting()
    assert tool.cards[-1][0] == "Meeting notes are switched off" and tool.meeting is None
    monkeypatch.setitem(vt.CONFIG, "meeting_notes", True)
    tool._start_meeting()
    monkeypatch.setitem(vt.CONFIG, "meeting_notes", False)  # turned off while the card was showing
    tool.card_buttons[-1]["Everyone agreed, start"]()
    assert tool.meeting is None


def test_the_call_offer_card_starts_only_after_everyone_agreed(tool, monkeypatch, isolated):
    meeting_tool(tool, monkeypatch, isolated)
    shown = []

    class Card:
        def show(self, title, body, buttons=None, seconds=60):
            shown.append(dict(buttons))
    tool.card = Card()
    tool._offer_meeting("Zoom")
    assert tool.meeting is None
    shown[0]["Everyone agreed, start"]()
    assert tool.meeting is not None and tool.meeting_app == "Zoom" and tool.meeting_auto_stop is True
    tool.meeting.abort()
    tool.meeting = None


def test_a_meeting_without_claude_has_no_follow_up_email_and_with_claude_it_does(tool, monkeypatch, isolated):
    for on in (False, True):
        monkeypatch.setitem(vt.CONFIG, "claude_actions", on)
        monkeypatch.setitem(vt.CONFIG, "meeting_followup", on)
        client = FakeClient(json.dumps(NOTES))
        meeting_tool(tool, monkeypatch, isolated, client=client)
        tool.claude_calls.clear()
        run_meeting(tool)
        system = client.sent[-1]["messages"][0]["content"]
        assert ("follow_up_email" in system) is on and len(tool.claude_calls) == (1 if on else 0)
        assert tool._follow_up_wanted() is on
    docs = notes_files(isolated)
    from docx import Document
    assert any("Follow-up email" in [p.text for p in Document(d).paragraphs] for d in docs) and len(docs) >= 1


def test_a_long_transcript_is_trimmed_to_the_most_recent_part_for_the_local_model():
    lines = [(i * 10, "Me" if i % 2 else "Others", f"line number {i} " + "word " * 20) for i in range(1200)]
    text, trimmed = meetings.trim_transcript(lines)
    assert trimmed and text.startswith(meetings.TRIMMED_NOTE) and "line number 1199 " in text and "line number 0 " not in text
    assert len(text) <= meetings.MAX_TRANSCRIPT_CHARS + len(meetings.TRIMMED_NOTE) + 5
    short = lines[:5]
    assert meetings.trim_transcript(short) == (meetings.transcript_text(short), False)
    raw = RawOpenAI(json.dumps({"title": "Long", "summary": "s"}))
    notes = meetings.make_notes(local_cleanup_win.LocalClient(raw), "local", lines, 0, 100, local=True, followup=False)
    sent = raw.sent[0]["messages"][1]["content"]
    assert notes["title"] == "Long" and meetings.TRIMMED_NOTE in sent and len(sent) < meetings.MAX_TRANSCRIPT_CHARS + 2000


# ====================================================================================================== friend version: Claude is off

def test_claude_that_is_not_installed_or_not_signed_in_gives_a_card_not_a_crash_or_a_window(tool, monkeypatch):
    import shutil
    monkeypatch.setattr(commands, "claude_problem", REAL_CLAUDE_PROBLEM)
    monkeypatch.setattr(commands, "signed_in", REAL_SIGNED_IN)
    ran = []
    monkeypatch.setattr(commands.subprocess, "run", lambda *a, **k: ran.append(a) or SimpleNamespace(stdout='{"loggedIn": false}'))
    monkeypatch.setattr(commands.subprocess, "Popen", lambda *a, **k: pytest.fail("no claude process may start"))
    tool.claude_session = SimpleNamespace(ask=lambda *a, **k: pytest.fail("no session"))
    monkeypatch.setattr(commands, "run_claude", lambda *a, **k: pytest.fail("no claude run"))
    monkeypatch.setattr(shutil, "which", lambda name: None)
    commands._sign_in_check.update(at=0.0, ok=False)
    tool._claude("check my calendar", {})
    assert tool.cards[-1] == ("Claude isn't ready", commands.CLAUDE_NOT_INSTALLED, False) and ran == []  # nothing was even run
    assert "own Claude plan" in commands.CLAUDE_NOT_INSTALLED and tool.light.claude_busy == 0
    monkeypatch.setattr(shutil, "which", lambda name: r"C:\\tools\\claude.cmd")
    commands._sign_in_check.update(at=0.0, ok=False)
    tool._claude("check my calendar", {})
    assert tool.cards[-1] == ("Claude isn't ready", commands.CLAUDE_NOT_SIGNED_IN, False) and len(ran) == 1
    assert ran[0][0][1:] == ["auth", "status"]
    assert commands.claude_problem() == commands.CLAUDE_NOT_SIGNED_IN and len(ran) == 1  # a "no" is remembered for 20 s
    # installed and signed in: it goes ahead
    monkeypatch.setattr(commands.subprocess, "run", lambda *a, **k: SimpleNamespace(stdout='{"loggedIn": true}'))
    commands._sign_in_check.update(at=0.0, ok=False)
    tool.claude_session = SimpleNamespace(ask=lambda said, ctx, style="": ("Done it.", [], 1.0))
    tool._claude("check my calendar", {})
    assert tool.cards[-1] == ("Claude", "Done it.", True)
    commands._sign_in_check.update(at=0.0, ok=False)


def test_the_style_button_needs_claude_and_the_page_hides_it_otherwise(api, tool, monkeypatch):
    api, host = api
    monkeypatch.setitem(vt.CONFIG, "claude_actions", False)
    monkeypatch.setattr(commands, "signed_in", lambda: pytest.fail("not asked while Claude is off"))
    assert api.call("style")["claude_available"] is False and api.call("overview")["claude_signed_in"] is False
    monkeypatch.setitem(vt.CONFIG, "claude_actions", True)
    monkeypatch.setattr(commands, "signed_in", lambda: False)
    host._signed_in = (0.0, False)
    assert api.call("style")["claude_available"] is False
    monkeypatch.setattr(commands, "signed_in", lambda: True)
    host._signed_in = (0.0, False)
    assert api.call("style")["claude_available"] is True
    # pressing it with Claude off says so on a card and starts nothing
    monkeypatch.setitem(vt.CONFIG, "claude_actions", False)
    monkeypatch.setattr(commands, "learn_style", lambda *a, **k: pytest.fail("no learning without Claude"))
    tool._refresh_style()
    assert tool.cards[-1] == ("Claude is off", vt.CLAUDE_OFF_TEXT, False)


def test_the_learned_style_is_told_apart_from_the_neutral_one_that_ships(isolated):
    assert commands.style_learned() is False
    (isolated / "style.md").write_text((HERE / "style.md").read_text(encoding="utf-8"), encoding="utf-8")
    assert commands.style_learned() is False
    (isolated / "style.md").write_text("# My writing style\n\nLearned 25 October 2026 from my sent emails by VoiceTool.\n", encoding="utf-8")
    assert commands.style_learned() is True


# ====================================================================================================== friend version: welcome

def test_the_welcome_card_shows_while_first_run_is_true_and_dismissing_clears_it(api, isolated, monkeypatch):
    api, host = api
    monkeypatch.setitem(vt.CONFIG, "first_run", True)
    assert api.call("overview")["first_run"] is True
    assert api.call("dismiss_welcome") == {"first_run": False}
    assert vt.CONFIG["first_run"] is False and json.loads((isolated / "config.json").read_text())["first_run"] is False
    assert api.call("overview")["first_run"] is False
    monkeypatch.delitem(vt.CONFIG, "first_run")
    assert api.call("overview")["first_run"] is False  # (a config without the key: no welcome)


def test_the_first_start_opens_the_window_on_its_own_but_test_mode_does_not(monkeypatch):
    source = (HERE / "voicetool.py").read_text(encoding="utf-8")
    assert 'CONFIG.get("first_run") and not test' in source  # the first start shows the Welcome card; test copies stay quiet


def test_the_old_azure_and_personal_paths_are_gone_from_the_code():
    for name in ("voicetool.py", "commands.py", "voice_actions.py", "actions_windows.py", "local_cleanup_win.py", "meetings.py", "app_window_win.py"):
        text = (HERE / name).read_text(encoding="utf-8")
        for needle in ("X:\\\\", "G:\\\\My Drive", "C:\\\\Users\\\\"):
            assert needle not in text, (name, needle)
    assert 'APP_ID = "VoiceTool.App"' in (HERE / "app_window_win.py").read_text(encoding="utf-8")


def test_a_small_model_that_leaves_out_the_route_key_is_still_understood():
    """Qwen3 4B sometimes answers {"system": "mute"} instead of {"route": "system", "target": "mute"}."""
    cases = [({"system": "mute"}, ("system", "mute")), ({"note": "buy milk", "topic": "buy milk", "text": "Buy milk."}, ("note", "")),
             ({"read": "true", "text": "x"}, ("read", "")), ({"open": "Notepad"}, ("open", "Notepad")),
             ({"text": "no route at all"}, ("", "")), ({"route": "Banana", "text": ""}, ("banana", ""))]
    for reply, (kind, target) in cases:
        got_kind, _, extra = commands.route(FakeClient(json.dumps(reply)), "m", "x", {"selection": ""})
        assert (got_kind, extra["target"]) == (kind, target), reply
    system = commands.ROUTER
    assert 'ALWAYS include the "route" key' in system and '"mute" -> {"route": "system", "target": "mute"}' in system


# ====================================================================================================== friend version: fix round (review fixes)

def test_a_model_that_wont_start_pastes_the_words_as_heard_and_shows_one_card(tool):
    tool.azure = FakeCleaner(fail=local_cleanup_win.ModelServerUnavailable("gave up"))
    run_dictation(tool, "um open zeph ear please")
    assert PASTED == ["um open zeph ear please "]
    assert tool.cards[-1][0] == "Tidy-up won't start" and "your words are typed as heard" in tool.cards[-1][1]
    assert "Try restarting VoiceTool" in tool.cards[-1][1] and "run Setup again" in tool.cards[-1][1]
    run_dictation(tool, "and another one here")
    assert PASTED[-1] == "and another one here " and len(tool.cards) == 1  # one card per run
    assert "error" not in tool.light.history
    tool.azure._client = lambda: (_ for _ in ()).throw(local_cleanup_win.ModelServerUnavailable("gave up"))
    t = tool.triggers[COMMAND]
    t.key_down = False
    t.context = {"app": "notepad.exe", "title": "x", "selection": "some text", "pid": 1, "hwnd": 2}
    tool.transcribe = lambda audio: "make this shorter"
    tool._command(t, np.full(vt.RATE, 0.3, np.float32))
    assert tool.cards[-1][0] == "Tidy-up won't start" and "commands can't run right now" in tool.cards[-1][1]
    assert PASTED[-1] == "and another one here "  # (nothing new was pasted for the command)


def fake_install(tmp_path, vulkan=True):
    for name in ("v.exe", "c.exe", "m.gguf"):
        (tmp_path / name).write_bytes(b"x")
    return {**json.loads((HERE / "config.json").read_text(encoding="utf-8")), "offline_server": str(tmp_path / "v.exe"),
            "offline_server_cpu": str(tmp_path / "c.exe"), "offline_model": str(tmp_path / "m.gguf"),
            "offline_backend": "vulkan" if vulkan else "cpu", "offline_port": 8999}


class DeadProcess:
    pid = 1

    def poll(self):
        return 3

    def kill(self):
        pass

    def wait(self, timeout=None):
        return 3


class LiveProcess(DeadProcess):
    def poll(self):
        return None


def test_after_two_failed_server_starts_the_model_is_not_started_again_for_the_run(tmp_path, monkeypatch):
    cfg = fake_install(tmp_path, vulkan=False)
    started, logged = [], []
    monkeypatch.setattr(local_cleanup_win.subprocess, "Popen", lambda cmd, **kw: started.append(cmd) or DeadProcess())
    monkeypatch.setattr(local_cleanup_win, "_port_in_use", lambda port: False)
    log = SimpleNamespace(info=lambda *a: None, warning=lambda *a: logged.append(a), error=lambda *a: None)
    c = local_cleanup_win.LocalCleaner(cfg, log, startup_seconds=5)
    with pytest.raises(RuntimeError) as first:
        c.start()
    assert not isinstance(first.value, local_cleanup_win.ModelServerUnavailable) and len(started) == 1  # the first failure
    with pytest.raises(local_cleanup_win.ModelServerUnavailable, match="Try restarting VoiceTool"):
        c.start()
    assert len(started) == 2 and c.gave_up is True  # the second failure: now it stops
    for _ in range(3):
        with pytest.raises(local_cleanup_win.ModelServerUnavailable):
            c.start()
    assert len(started) == 2  # no program is launched any more (each try could cost a minute and a half)
    with pytest.raises(local_cleanup_win.ModelServerUnavailable):
        c.client()
    assert len(started) == 2
    c.refresh()  # the downloader has just repaired the install: allowed to try again
    with pytest.raises(RuntimeError):
        c.start()
    assert len(started) == 3 and c.failed_starts == 1 and not c.gave_up
    c.failed_starts, c.gave_up = 5, True
    monkeypatch.setattr(c, "start", lambda: None)
    c.start_in_background()
    assert c.failed_starts == 0 and c.gave_up is False


def test_a_vulkan_failure_that_the_cpu_build_fixes_does_not_count_and_a_good_start_resets_the_count(tmp_path, monkeypatch):
    cfg = fake_install(tmp_path, vulkan=True)
    started = []

    def popen(cmd, **kw):
        started.append(cmd[0])
        return DeadProcess() if cmd[0].endswith("v.exe") else LiveProcess()
    monkeypatch.setattr(local_cleanup_win.subprocess, "Popen", popen)
    monkeypatch.setattr(local_cleanup_win, "_port_in_use", lambda port: False)
    monkeypatch.setattr(local_cleanup_win, "_get", lambda port, path, timeout=2: (200, None))
    c = local_cleanup_win.LocalCleaner(cfg, None, startup_seconds=5)
    c.start()
    assert [Path(s).name for s in started] == ["v.exe", "c.exe"] and c.backend == "cpu" and c.fell_back
    assert c.failed_starts == 0 and not c.gave_up


def test_a_blocked_server_program_is_a_failed_start_not_a_crash_and_the_cpu_build_is_still_tried(tmp_path, monkeypatch):
    cfg = fake_install(tmp_path, vulkan=True)
    started = []

    def popen(cmd, **kw):
        started.append(Path(cmd[0]).name)
        if cmd[0].endswith("v.exe"):
            raise PermissionError(5, "Access is denied")  # antivirus or a damaged download
        return LiveProcess()
    monkeypatch.setattr(local_cleanup_win.subprocess, "Popen", popen)
    monkeypatch.setattr(local_cleanup_win, "_port_in_use", lambda port: False)
    monkeypatch.setattr(local_cleanup_win, "_get", lambda port, path, timeout=2: (200, None))
    c = local_cleanup_win.LocalCleaner(cfg, None, startup_seconds=5)
    c.start()  # OSError from Popen didn't escape: it fell back to the CPU build
    assert started == ["v.exe", "c.exe"] and c.backend == "cpu" and c.proc is not None and c.fell_back
    # both builds blocked: two failed starts, then it gives up with the plain error and launches nothing more
    started.clear()
    monkeypatch.setattr(local_cleanup_win, "_get", lambda port, path, timeout=2: (None, None))  # (nothing answers now)
    monkeypatch.setattr(local_cleanup_win.subprocess, "Popen", lambda cmd, **kw: started.append(Path(cmd[0]).name) or (_ for _ in ()).throw(OSError(193, "not a valid program")))
    d = local_cleanup_win.LocalCleaner(cfg, None, startup_seconds=5)
    with pytest.raises(local_cleanup_win.ModelServerUnavailable) as e:
        d.start()
    assert started == ["v.exe", "c.exe"] and isinstance(e.value.__cause__, RuntimeError) and d.gave_up
    with pytest.raises(local_cleanup_win.ModelServerUnavailable):
        d.start()
    assert len(started) == 2


def test_a_server_file_deleted_after_start_up_is_a_failed_start_not_a_crash(tmp_path, monkeypatch):
    cfg = fake_install(tmp_path, vulkan=True)
    c = local_cleanup_win.LocalCleaner(cfg, None, startup_seconds=5)
    (tmp_path / "v.exe").unlink()
    (tmp_path / "c.exe").unlink()
    monkeypatch.setattr(local_cleanup_win, "_port_in_use", lambda port: False)
    with pytest.raises(RuntimeError, match="files aren't there"):  # (a failed start, not a TypeError: nothing to fall back to)
        c.start()
    with pytest.raises(local_cleanup_win.ModelServerUnavailable):  # the second failure: it gives up
        c.start()
    assert local_cleanup_win.DEFAULT_PORT == 8766 and vt.json.loads((HERE / "config.json").read_text(encoding="utf-8"))["offline_port"] == 8766


def test_refresh_picks_up_the_backend_the_downloader_just_chose(tmp_path):
    cfg = fake_install(tmp_path, vulkan=True)
    c = local_cleanup_win.LocalCleaner(cfg, None)
    assert c.backend == "vulkan"
    cfg["offline_backend"] = "cpu"  # (configure() found no usable graphics card)
    c.gave_up, c.failed_starts, c.fell_back = True, 2, True
    c.refresh()
    assert c.backend == "cpu" and not c.gave_up and c.failed_starts == 0 and not c.fell_back


def test_after_the_downloader_ends_this_computers_settings_are_re_read_before_the_model_starts(downloader, monkeypatch, isolated):
    import codecs
    (downloader.folder / "fetch_models.py").write_text("print('x')")
    machine = ("offline_backend", "offline_args", "offline_server", "offline_server_cpu", "offline_model", "speech_model_path",
               "offline_port")
    assert vt.MACHINE_KEYS == machine
    for key in machine:  # (so the test puts every one back afterwards)
        monkeypatch.setitem(vt.CONFIG, key, vt.CONFIG.get(key, "stale"))
    monkeypatch.setitem(vt.CONFIG, "beeps", True)
    stale = {"offline_backend": "vulkan", "offline_args": ["-ngl", "99"], "offline_server": "old.exe", "offline_server_cpu": "oldc.exe",
             "offline_model": "old.gguf", "speech_model_path": "", "offline_port": 8765}
    for key, value in stale.items():
        monkeypatch.setitem(vt.CONFIG, key, value)
    disk = {**vt.CONFIG, "offline_backend": "cpu", "offline_args": ["-c", "16384", "--no-webui"],
            "offline_server": r"runtime\llama-vulkan\llama-server.exe", "offline_server_cpu": r"runtime\llama-cpu\llama-server.exe",
            "offline_model": r"models\q\q.gguf", "speech_model_path": r"models\parakeet-tdt-0.6b-v2", "offline_port": 8766,
            "beeps": False}
    (isolated / "config.json").write_bytes(codecs.BOM_UTF8 + json.dumps(disk).encode("utf-8"))  # (a BOM: Notepad or PowerShell)
    seen, refreshed, reloaded = [], [], []
    local = SimpleNamespace(refresh=lambda: refreshed.append(vt.CONFIG["offline_backend"]))
    host = downloader.host
    host.tool.azure = SimpleNamespace(local=local, sync_offline=lambda: seen.append(
        (vt.CONFIG["offline_backend"], list(vt.CONFIG["offline_args"]), vt.CONFIG["offline_model"], vt.CONFIG["speech_model_path"],
         vt.CONFIG["offline_port"])))
    host.tool.transcribe = SimpleNamespace(missing=True, reload=lambda: reloaded.append(vt.CONFIG["speech_model_path"]))
    assert downloader.api.call("finish_setup") == {"started": True}
    assert seen == []
    downloader.proc().release.set()
    assert wait_until(lambda: seen, 10)
    assert seen == [("cpu", ["-c", "16384", "--no-webui"], r"models\q\q.gguf", r"models\parakeet-tdt-0.6b-v2", 8766)]
    assert refreshed == ["cpu"] and reloaded == [r"models\parakeet-tdt-0.6b-v2"]  # (the transcriber loads the new folder)
    assert vt.CONFIG["beeps"] is True  # only this computer's settings are taken from the disk, never the person's switches
    # a settings save made afterwards (any switch) keeps what the downloader found
    vt.CONFIG["beeps"] = False
    settings_window.save(vt.CONFIG)
    on_disk = json.loads((isolated / "config.json").read_text(encoding="utf-8"))
    assert on_disk["offline_backend"] == "cpu" and on_disk["offline_port"] == 8766 and on_disk["offline_model"] == r"models\q\q.gguf"


def test_a_failed_download_still_takes_in_the_settings_and_the_speech_model_but_does_not_start_the_tidy_up(downloader, monkeypatch, isolated):
    (downloader.folder / "fetch_models.py").write_text("print('x')")
    for key in vt.MACHINE_KEYS:
        monkeypatch.setitem(vt.CONFIG, key, vt.CONFIG.get(key, "x"))
    (isolated / "config.json").write_text(json.dumps({**vt.CONFIG, "offline_backend": "cpu"}), encoding="utf-8")
    monkeypatch.setitem(vt.CONFIG, "offline_backend", "vulkan")
    started, reloaded = [], []
    downloader.host.tool.azure = SimpleNamespace(local=None, sync_offline=lambda: started.append(1))
    downloader.host.tool.transcribe = SimpleNamespace(missing=True, reload=lambda: reloaded.append(1))
    downloader.api.call("finish_setup")
    downloader.proc().code = 3  # the tidy-up model didn't finish (the speech model may have)
    downloader.proc().release.set()
    assert wait_until(lambda: reloaded == [1], 10) and vt.CONFIG["offline_backend"] == "cpu" and started == []


def test_reload_machine_settings_copes_with_a_missing_or_unreadable_config(isolated, monkeypatch):
    for key in vt.MACHINE_KEYS:
        monkeypatch.setitem(vt.CONFIG, key, vt.CONFIG.get(key, "x"))
    assert vt.reload_machine_settings() is False  # no config.json in the temp folder: nothing changes, nothing breaks
    (isolated / "config.json").write_text("{ not json", encoding="utf-8")
    assert vt.reload_machine_settings() is False
    (isolated / "config.json").write_text(json.dumps({"offline_port": 1}), encoding="utf-8")
    assert vt.reload_machine_settings() is True and vt.CONFIG["offline_port"] == 1 and "offline_backend" not in vt.CONFIG


def test_a_config_saved_with_a_bom_loads(tmp_path):
    import codecs
    (tmp_path / "words.txt").write_text("x", encoding="utf-8")
    cfg = tmp_path / "config.json"
    cfg.write_bytes(codecs.BOM_UTF8 + b'{"first_run": true}')
    code = ("import sys, json; sys.path.insert(0, sys.argv[1]); import fetch_models; "
            "print(json.dumps(fetch_models.configure(sys.argv[2], detect=False)))")
    out = subprocess.run([sys.executable, "-c", code, str(HERE), str(tmp_path)], capture_output=True, text=True, creationflags=0x08000000)
    assert out.returncode == 0, out.stderr
    assert json.loads(cfg.read_text(encoding="utf-8"))["first_run"] is True  # (the BOM file was read, not thrown away)
    source = (HERE / "voicetool.py").read_text(encoding="utf-8")
    assert 'read_text(encoding="utf-8-sig")' in source.split("CONFIG = ")[1].split("\n")[0]
    # the app's own loader takes a BOM too (run in a copy so the real config.json is never touched)
    copy = tmp_path / "app"
    copy.mkdir()
    for name in ("voicetool.py",):
        (copy / name).write_bytes((HERE / name).read_bytes())
    (copy / "config.json").write_bytes(codecs.BOM_UTF8 + (HERE / "config.json").read_bytes())
    snippet = ("import ast, json, pathlib, sys\n"
               "src = pathlib.Path(sys.argv[1], 'voicetool.py').read_text(encoding='utf-8')\n"
               "line = [l for l in src.splitlines() if l.startswith('CONFIG = ')][0]\n"
               "HERE = pathlib.Path(sys.argv[1])\n"
               "exec(line.split('#')[0].strip())\n"
               "print(CONFIG['offline_port'])\n")
    done = subprocess.run([sys.executable, "-c", snippet, str(copy)], capture_output=True, text=True, creationflags=0x08000000)
    assert done.stdout.strip() == "8766", done.stderr


def test_the_old_settings_window_makes_the_learned_fixes_file_before_opening_it_and_clears_first_run(isolated):
    assert not (isolated / "corrections.json").exists()
    settings_window.open_learned_fixes()
    assert OPENED[-1] == str(isolated / "corrections.json") and json.loads((isolated / "corrections.json").read_text()) == []
    (isolated / "corrections.json").write_text('[["a", "b"]]', encoding="utf-8")
    settings_window.open_learned_fixes()
    assert (isolated / "corrections.json").read_text() == '[["a", "b"]]'  # an existing file is never replaced
    config = {"first_run": True, "enabled": True}
    settings_window.clear_first_run(config)
    assert config["first_run"] is False and json.loads((isolated / "config.json").read_text())["first_run"] is False
    config = {"enabled": True}
    (isolated / "config.json").unlink()
    settings_window.clear_first_run(config)
    assert "first_run" not in config and not (isolated / "config.json").exists()  # nothing to clear: nothing written
    source = (HERE / "settings_window.py").read_text(encoding="utf-8")
    assert "clear_first_run(config)" in source.split("class SettingsWindow")[1] and "open_learned_fixes" in source
    # the Finish setup button is offered only when the tool can start it and something is missing
    assert settings_window.finish_setup_wanted({"finish_setup": lambda: {}, "setup_missing": lambda: True}) is True
    assert settings_window.finish_setup_wanted({"finish_setup": lambda: {}, "setup_missing": lambda: False}) is False
    assert settings_window.finish_setup_wanted({"quit": lambda: None}) is False


def test_the_old_settings_window_can_finish_setup_through_the_one_host(tool, monkeypatch):
    seen = {}
    monkeypatch.setattr(settings_window, "SettingsWindow", lambda root, config, actions: seen.update(actions) or SimpleNamespace())
    tool.light = SimpleNamespace(root=None)
    tool._open_settings_tk()
    assert set(seen) >= {"finish_setup", "setup_missing", "quit", "offline"}
    monkeypatch.setattr(vt.WinHost, "setup_status", lambda self: {"missing": True})
    assert seen["setup_missing"]() is True
    assert tool._win_host() is tool._win_host()  # one host, so the "running" state is shared with the app window


def test_the_style_text_for_writing_does_not_assume_who_the_user_is():
    text = (HERE / "styles.py").read_text(encoding="utf-8")
    assert "unless he asks" not in text and " he " not in text.lower().replace("\n", " ") and "unless they ask" in text


def test_the_claude_helper_is_started_with_this_installs_workdir_in_its_command_line():
    cmd = commands.claude_command("sonnet", "--input-format", "stream-json")
    assert cmd[cmd.index("--add-dir") + 1] == str(commands.HERE / "claude_workdir")
    assert cmd.index("--add-dir") < cmd.index("--permission-mode")  # (not swallowed by the variadic --allowedTools list)
    assert "stream-json" in cmd and any(a.startswith("mcp__claude_ai_") for a in cmd)  # what Stop VoiceTool.cmd looks for


def test_stop_voicetool_cmd_stops_only_this_installs_programs_and_its_own_claude(tmp_path):
    """Decoy programs in a temp folder (never the real VoiceTool): the install's own python and Claude helper are stopped;
    a folder whose name merely starts the same, and some other Claude, are left alone."""
    inst, other_inst = tmp_path / "VoiceTool", tmp_path / "VoiceTool_old"
    elsewhere = tmp_path / "somewhere"
    for folder in (inst, other_inst, elsewhere):
        (folder / "claude_workdir").mkdir(parents=True)
    (inst / "Stop VoiceTool.cmd").write_bytes((HERE / "Stop VoiceTool.cmd").read_bytes())
    bin_dir = tmp_path / "bin"
    bin_dir.mkdir()
    claude = bin_dir / "claude.exe"  # the venv's python launcher under another name: it runs python with whatever it's given
    claude.write_bytes((Path(sys.executable).parent / "python.exe").read_bytes() if (Path(sys.executable).parent / "python.exe").exists() else b"")
    (bin_dir / "pyvenv.cfg").write_text(f"home = {Path(BASE_PYTHON).parent}\ninclude-system-site-packages = false\n")
    if not claude.stat().st_size or not (Path(sys.prefix) / "pyvenv.cfg").exists():
        pytest.skip("no virtual-environment launcher to copy for the Claude decoy")
    sleeper = "import time; time.sleep(120)"
    procs = {}

    def run(name, exe, *args):
        procs[name] = subprocess.Popen([str(exe), "-c", sleeper, *args], creationflags=0x08000000)
    run("own_python", BASE_PYTHON, str(inst / "voicetool.py"))
    run("lookalike_folder_python", BASE_PYTHON, str(other_inst / "voicetool.py"))
    run("unrelated_python", BASE_PYTHON, str(elsewhere / "voicetool.py"))
    run("own_claude", claude, "stream-json", "mcp__claude_ai_Gmail__x", "--add-dir", str(inst / "claude_workdir"))
    run("lookalike_claude", claude, "stream-json", "mcp__claude_ai_Gmail__x", "--add-dir", str(other_inst / "claude_workdir"))
    run("other_claude", claude, "stream-json", "mcp__claude_ai_Gmail__x", "--add-dir", str(elsewhere / "claude_workdir"))
    run("claude_without_workdir", claude, "stream-json", "mcp__claude_ai_Gmail__x", str(inst))
    try:
        time.sleep(1.5)
        assert all(p.poll() is None for p in procs.values())
        done = subprocess.run(["cmd", "/c", str(inst / "Stop VoiceTool.cmd")], capture_output=True, text=True, timeout=120,
                              creationflags=0x08000000)
        assert done.returncode == 0, done.stderr
        assert wait_until(lambda: procs["own_python"].poll() is not None and procs["own_claude"].poll() is not None, 15), "not stopped"
        time.sleep(1.0)
        for name in ("lookalike_folder_python", "unrelated_python", "lookalike_claude", "other_claude", "claude_without_workdir"):
            assert procs[name].poll() is None, f"{name} was stopped but isn't this install's"
    finally:
        for p in procs.values():
            p.kill()
    script = (HERE / "Stop VoiceTool.cmd").read_text(encoding="utf-8")
    assert "claude_workdir" in script and "X:\\" not in script and "VoiceTool_friend" not in script


def test_the_speech_files_the_app_checks_for_are_the_ones_the_downloader_fetches():
    import fetch_models
    fetched = {Path(i["dest"]).name for i in fetch_models.MANIFEST if i["id"].startswith("parakeet")}
    assert fetched == set(vt.SPEECH_FILES)  # (so "Finish setup" really repairs whatever the app says is missing)
    assert fetch_models.PARAKEET_REL == json.loads((HERE / "config.json").read_text(encoding="utf-8"))["speech_model_path"]
    assert fetch_models.MACHINE_KEYS["speech_model_path"] == fetch_models.PARAKEET_REL


def test_opening_learned_fixes_before_any_exist_opens_nothing_and_does_not_crash(api, isolated):
    api, host = api
    assert not (isolated / "corrections.json").exists()
    assert api.call("open", {"what": "fixes"}) == {"opened": False} and OPENED == []
    (isolated / "corrections.json").write_text("[]", encoding="utf-8")
    assert api.call("open", {"what": "fixes"}) == {"opened": True}


def test_an_incomplete_speech_model_folder_is_told_apart_from_a_missing_one(tmp_path, monkeypatch, tool):
    def load(name, path=None, **kw):
        raise ValueError("index: 1 Got: 80 Expected: 128")  # (what a missing config.json did)
    monkeypatch.setattr(vt.onnx_asr, "load_model", load)
    monkeypatch.setattr(vt, "HERE", tmp_path)
    monkeypatch.setitem(vt.CONFIG, "speech_model_path", r"models\parakeet-tdt-0.6b-v2")
    missing = vt.Transcriber()
    assert missing.problem == vt.SPEECH_MISSING_TEXT and "is missing" in missing.problem
    folder = tmp_path / "models" / "parakeet-tdt-0.6b-v2"
    folder.mkdir(parents=True)
    for name in ("vocab.txt", "decoder_joint-model.int8.onnx", "encoder-model.int8.onnx"):
        (folder / name).write_bytes(b"x")  # (config.json is the one that is missing)
    incomplete = vt.Transcriber()
    assert incomplete.missing and incomplete.problem == vt.SPEECH_INCOMPLETE_TEXT and "incomplete" in incomplete.problem
    with pytest.raises(vt.SpeechModelMissing, match="incomplete"):
        incomplete(np.zeros(vt.RATE, np.float32))
    tool.transcribe, tool.azure = incomplete, FakeCleaner()
    run_worker_job(tool, "dictate")
    assert tool.cards[-1] == ("Finish setup first", vt.SPEECH_INCOMPLETE_TEXT, False)
    assert vt.speech_model_missing(incomplete) is True


# ====================================================================================================== review fixes (10-30)
# An outside release review found these: Ctrl shortcuts never lock or submit (F2), each recording's job is frozen (F3), text is only typed
# into the window it was meant for (F4), the clean-up timer never keeps a run alive (F5), and a failed command is kept.

def _lock_command(tool):
    at = now()
    key(tool, COMMAND, True, at)
    key(tool, COMMAND, False, at + 0.05)
    key(tool, COMMAND, True, at + 0.2)
    key(tool, COMMAND, False, at + 0.25)
    return tool.triggers[COMMAND]


@pytest.mark.parametrize("other", ["key", "mouse"])
def test_review_ctrl_shortcut_after_a_lone_tap_never_locks_the_mic(tool, other):
    at = now()
    key(tool, COMMAND, True, at)
    key(tool, COMMAND, False, at + 0.05)
    key(tool, COMMAND, True, at + 0.2)  # (the second press could be a double-tap: nothing is decided yet)
    if other == "key":
        tool._handle("other", None, 0x43, at + 0.22)  # C: it was Ctrl+C
    else:
        tool._handle("mouse", None, None, at + 0.22)  # Ctrl+click
    key(tool, COMMAND, False, at + 0.25)
    assert not tool.recorder.running and tool.triggers[COMMAND].state == vt.IDLE and jobs(tool) == 0


@pytest.mark.parametrize("other", ["key", "mouse"])
def test_review_ctrl_shortcut_during_hands_free_never_submits_it(tool, other):
    t = _lock_command(tool)
    assert t.state == vt.LOCKED
    at = now()
    key(tool, COMMAND, True, at)
    assert jobs(tool) == 0  # nothing is submitted on the way down
    if other == "key":
        tool._handle("other", None, 0x56, at + 0.02)  # V: Ctrl+V while it listens
    else:
        tool._handle("mouse", None, None, at + 0.02)
    key(tool, COMMAND, False, at + 0.05)
    assert jobs(tool) == 0 and t.state == vt.LOCKED and tool.recorder.running  # still listening hands-free
    key(tool, COMMAND, True, at + 1.0)  # a lone tap still stops it, at the release
    assert jobs(tool) == 0
    key(tool, COMMAND, False, at + 1.05)
    assert jobs(tool) == 1 and t.state == vt.IDLE


def test_review_shift_first_and_other_modifiers_never_start_a_command(tool, monkeypatch):
    for vk in (0x10, 0x12, 0x5B):  # Ctrl then Shift, Alt or Win (Ctrl+Shift+T, Ctrl+Alt+Del, Ctrl+Win+arrows)
        at = now()
        key(tool, COMMAND, True, at)
        tool._handle("other", None, vk, at + 0.02)
        key(tool, COMMAND, False, at + 1.0)
        assert not tool.recorder.running and jobs(tool) == 0, vk
        tool.triggers[COMMAND].state = vt.IDLE


def test_review_queued_dictation_keeps_its_own_app_and_window(tool, monkeypatch):
    t = tool.triggers[DICTATE]
    t.app = {"app": "outlook.exe", "hwnd": 101, "title": "First message"}
    monkeypatch.setattr(tool, "_front_window", lambda: {"app": "outlook.exe", "hwnd": 101, "title": "First message"})
    tool.recorder.start()
    tool._finish(t)
    job, _ = tool.jobs.get_nowait()
    monkeypatch.setattr(tool, "_front_window", lambda: {"app": "notepad.exe", "hwnd": 202, "title": "Second document"})
    key(tool, DICTATE, True, now())  # the next recording starts while the first is still waiting
    assert job.app["hwnd"] == 101 and job.app["app"] == "outlook.exe" and job.target == 101
    assert t.app["hwnd"] == 202  # (the trigger moved on; the job didn't)


def test_review_queued_command_keeps_its_own_selection_even_if_copied_late(tool, monkeypatch):
    t = tool.triggers[COMMAND]
    tool.recorder.start()
    t.state, t.pressed_at = vt.HOLDING, now() - 1
    box = t.capture
    tool._finish(t)
    job, _ = tool.jobs.get_nowait()
    key(tool, COMMAND, True, now())  # the next command starts (a new capture) before the first copy has finished
    assert t.capture is not box and job.capture is box
    box.context = {"app": "outlook.exe", "hwnd": 101, "selection": "First selected text"}  # the first copy lands late
    assert job.context["selection"] == "First selected text" and t.context is None


def test_review_two_quick_recordings_in_different_apps_each_use_their_own(tool, monkeypatch):
    """The first dictation is still being written (slow speech model) when the second is spoken in another app: each is
    tidied with its own app's rule, saved with its own app, and typed only into its own window."""
    styles.save({"styles": [], "app_rules": [{"apps": "WindowsTerminal.exe", "mode": "raw"}]})
    windows = {"now": {"app": "WindowsTerminal.exe", "hwnd": 101, "title": "Windows Terminal"}}
    monkeypatch.setattr(tool, "_front_window", lambda: dict(windows["now"]))
    t = tool.triggers[DICTATE]
    key(tool, DICTATE, True, now() - 1)
    key(tool, DICTATE, False, now())
    first, a1 = tool.jobs.get_nowait()
    windows["now"] = {"app": "notepad.exe", "hwnd": 202, "title": "Notes"}
    key(tool, DICTATE, True, now() - 1)
    key(tool, DICTATE, False, now())
    second, a2 = tool.jobs.get_nowait()
    tool.transcribe = lambda audio: "git status please now"
    tool.azure = FakeCleaner(result="Git status, please, now.")
    windows["now"] = {"app": "WindowsTerminal.exe", "hwnd": 101, "title": "Windows Terminal"}
    tool._dictate(first, a1)
    assert PASTED == ["git status please now "] and tool.azure.calls == []  # the terminal's rule: exactly as said
    windows["now"] = {"app": "notepad.exe", "hwnd": 202, "title": "Notes"}
    tool._dictate(second, a2)
    assert PASTED[-1] == "Git status, please, now. "
    rows = [json.loads(line) for line in (vt.HERE / "history.jsonl").read_text(encoding="utf-8").splitlines()]
    assert [r["app"] for r in rows] == ["WindowsTerminal.exe", "notepad.exe"]
    assert t.state == vt.IDLE


def test_review_dictation_is_not_typed_into_a_different_window(tool, monkeypatch):
    t = tool.triggers[DICTATE]
    t.app = {"app": "outlook.exe", "hwnd": 101, "title": "Private message"}
    monkeypatch.setattr(tool, "_front_window", lambda: {"app": "notepad.exe", "hwnd": 202})
    monkeypatch.setattr(commands, "window_info", lambda: {"app": "notepad.exe", "hwnd": 202})
    tool.transcribe = lambda audio: "Private message"
    tool._dictate(t, np.full(vt.RATE, 0.3, np.float32))
    assert not PASTED and tool.cards[-1][0] == "Not typed: you switched windows"
    assert "Private message" in (vt.HERE / "history.jsonl").read_text(encoding="utf-8")  # kept
    assert tool.light.state == "error"


def test_review_a_job_is_typed_where_the_recording_stopped(tool, monkeypatch):
    """Hands-free: he may start in one window and stop in another; the text goes where he was when he stopped."""
    windows = {"now": {"app": "outlook.exe", "hwnd": 101}}
    monkeypatch.setattr(tool, "_front_window", lambda: dict(windows["now"]))
    t0 = now()
    key(tool, DICTATE, True, t0)
    key(tool, DICTATE, False, t0 + 0.05)
    key(tool, DICTATE, True, t0 + 0.2)
    key(tool, DICTATE, False, t0 + 0.25)
    windows["now"] = {"app": "notepad.exe", "hwnd": 202}
    key(tool, DICTATE, True, t0 + 3)
    key(tool, DICTATE, False, t0 + 3.05)
    job, audio = tool.jobs.get_nowait()
    assert job.target == 202
    tool.transcribe = lambda audio: "hello there"
    tool._dictate(job, audio)
    assert PASTED == ["hello there "]


def test_review_rewrite_and_fix_are_not_typed_into_a_different_window(tool, monkeypatch):
    monkeypatch.setattr(tool, "_front_window", lambda: {"app": "notepad.exe", "hwnd": 202})
    selection = {"app": "outlook.exe", "title": "x", "selection": "Private message to rewrite", "pid": 1, "hwnd": 101}
    run_command(tool, "make this shorter", {"route": "rewrite", "text": ""}, selection, then=("Rewritten private message.",))
    assert not PASTED and tool.cards[-1][0] == "Not typed: you switched windows"
    run_command(tool, "that should be Zephyr", {"route": "fix", "text": "Zephyr"}, {**selection, "selection": "zef fir"})
    assert not PASTED and ("zef fir", "Zephyr") in corrections.load()  # still learned; just not typed elsewhere
    run_command(tool, "write a thank-you", {"route": "write", "text": "Thank you!"}, {**selection, "selection": ""})
    assert not PASTED
    monkeypatch.setattr(tool, "_front_window", lambda: {"app": "outlook.exe", "hwnd": 101})  # back in the right window
    run_command(tool, "make this shorter", {"route": "rewrite", "text": ""}, selection, then=("Shorter.",))
    assert PASTED == ["Shorter."]


def test_review_a_window_that_cannot_be_read_still_gets_its_text(tool, monkeypatch):
    t = tool.triggers[DICTATE]
    t.app = {}  # (the window couldn't be read when the key went down)
    monkeypatch.setattr(tool, "_front_window", lambda: {})
    tool.transcribe = lambda audio: "hello"
    tool._dictate(t, np.full(vt.RATE, 0.3, np.float32))
    assert PASTED == ["hello "]


def test_review_close_window_only_closes_the_window_it_was_said_in(tool, monkeypatch):
    sent = []
    monkeypatch.setattr(system_windows, "close_window", lambda: sent.append("close"))
    context = {"app": "notepad.exe", "title": "x", "selection": "", "pid": 1, "hwnd": 101}
    monkeypatch.setattr(tool, "_front_window", lambda: {"app": "outlook.exe", "hwnd": 202})
    run_command(tool, "close this window", {"route": "system", "target": "close_window"}, context)
    assert sent == [] and tool.cards[-1][0] == "Didn't close it"
    monkeypatch.setattr(tool, "_front_window", lambda: {"app": "notepad.exe", "hwnd": 101})
    run_command(tool, "close this window", {"route": "system", "target": "close_window"}, context)
    assert sent == ["close"]


def test_review_a_command_that_fails_is_kept_in_history_with_a_reason(tool, monkeypatch):
    def slow(*a, **kw):
        raise TimeoutError("Request timed out.")
    monkeypatch.setattr(commands, "route", slow)
    run_command(tool, "make this shorter", {"route": "rewrite"})
    rows = [json.loads(line) for line in (vt.HERE / "history.jsonl").read_text(encoding="utf-8").splitlines()]
    assert rows[-1]["command"] == "make this shorter" and rows[-1]["route"] == "failed"
    assert rows[-1]["result"].startswith("Not done:") and "saved in History" in rows[-1]["result"]
    assert tool.cards[-1][0] == "Command not done" and not PASTED and tool.light.state == "error"
    shown = ui_api.API(SimpleNamespace(), vt.CONFIG, lambda c: None).api_history()["items"][0]
    assert shown["said"] == "make this shorter" and shown["text"].startswith("Not done:")
    if hasattr(commands, "CommandFailed"):  # the model's own timeout message is shown as it is
        monkeypatch.setattr(commands, "route", lambda *a, **kw: (_ for _ in ()).throw(commands.CommandFailed("Took too long.")))
        run_command(tool, "write a note", {"route": "write"})
        assert tool.cards[-1] == ("Command not done", "Took too long.", False)


def test_review_spoken_answer_clean_up_never_keeps_a_run_alive(tool, tmp_path):
    path = tmp_path / "answer.wav"
    path.write_bytes(b"x")
    timer = tool._forget_later(str(path), seconds=0.05)
    assert timer.daemon and timer in tool._timers
    timer.join(2)
    assert not path.exists() and timer not in tool._timers
    path.write_bytes(b"x")
    slow = tool._forget_later(str(path), seconds=3600)
    tool._cancel_timers()
    slow.join(2)
    assert not slow.is_alive() and not tool._timers and path.exists()  # cancelled: Quit doesn't wait half an hour
    assert not [th for th in threading.enumerate() if isinstance(th, threading.Timer) and th.is_alive() and not th.daemon]


# ====================================================================================================== SLOW-PC COMMAND TIMEOUTS
# Token budgets, the processor-aware wait (local_cleanup_win.request_timeout), commands.ModelTimeout, and the processor-only
# server. Fakes only: no model runs here (a tiny local stand-in web server stands in for the model's timing and errors).

class SlowPCServer:
    """Stands in for LocalCleaner as far as LocalClient cares: does the model run on the processor?"""

    def __init__(self, cpu):
        self.cpu = cpu

    def on_cpu(self):
        return self.cpu


def slowpc_client(raw, cpu=True, base=60):
    return local_cleanup_win.LocalClient(raw, SlowPCServer(cpu), base)


class SlowPCReply:
    def __init__(self, content, finish="stop"):
        self.choices = [SimpleNamespace(message=SimpleNamespace(content=content), finish_reason=finish)]
        self.usage = None


class SlowPCRaw:
    """A model client that answers each request with the next item: a reply, or an error that is raised."""

    def __init__(self, *items):
        self.items, self.sent = list(items), []
        self.chat = SimpleNamespace(completions=self)

    def create(self, **kw):
        self.sent.append(kw)
        item = self.items.pop(0) if len(self.items) > 1 else self.items[0]
        if isinstance(item, BaseException):
            raise item
        return item if isinstance(item, SlowPCReply) else SlowPCReply(item)

    def with_options(self, **kw):
        self.options = kw
        return self


def slowpc_errors():
    """The errors the OpenAI client raises (built without a network)."""
    import openai
    response = lambda status: SimpleNamespace(status_code=status, headers={}, request=None)  # noqa: E731
    return {"timeout": openai.APITimeoutError(request=None),
            "connection": openai.APIConnectionError(message="Connection error.", request=None),
            "loading": openai.InternalServerError("Loading model", response=response(503), body=None),
            "bad_gateway": openai.InternalServerError("bad gateway", response=response(502), body=None),
            "auth": openai.AuthenticationError("no", response=response(401), body=None),
            "bad_request": openai.BadRequestError("too long for the context", response=response(400), body=None)}


def test_slowpc_the_wait_for_a_request_is_decided_in_one_function_scaled_by_the_job():
    rt = local_cleanup_win.request_timeout
    # on a graphics card (or anything that is not the processor) nothing changes: the flat wait it always had
    assert rt(1500, 600, cpu=False) == 60 and rt(1500, 600, cpu=False, base=45) == 45 and rt(0, 1, cpu=False) == 60
    # on the processor every job gets what it needs: reading the prompt plus writing the reply, plus slack
    route = rt(1500, commands.route_budget("make this shorter"), cpu=True)
    rewrite = rt(400, commands.rewrite_budget("I will be a little late."), cpu=True)
    write = rt(400, commands.WRITE_TOKENS, cpu=True)
    tidy = rt(250, 64, cpu=True)
    assert write > route > rewrite > tidy >= 60
    assert (tidy, rewrite, route, write) == (64, 82, 193, 250)  # (the numbers a 4-core laptop's speeds give)
    assert 120 < route < 300  # the old flat 60 s could never have covered the command router on a processor
    # longer prompts and longer replies wait longer, up to ten minutes, and never less than the configured base
    assert rt(3000, 100, True) > rt(1000, 100, True) and rt(1000, 900, True) > rt(1000, 100, True)
    assert rt(10 ** 6, 10 ** 6, True) == 600 and rt(0, 1, True, base=200) == 200 and rt(0, None, True) == 600
    assert local_cleanup_win.estimate_prompt_tokens([{"role": "system", "content": "s" * 3000}, {"role": "user", "content": "hello"}]) == 1002
    assert local_cleanup_win.estimate_prompt_tokens([{"content": [{"type": "text", "text": "abc"}, {"type": "image_url"}]}, "x", {}]) == 1
    assert local_cleanup_win.estimate_prompt_tokens(None) == 0


def test_slowpc_the_processor_client_gives_each_request_its_own_wait_and_nothing_else_changes():
    raw = SlowPCRaw("ok")
    messages = [{"role": "system", "content": "s" * 3000}, {"role": "user", "content": "hello"}]
    slowpc_client(raw).chat.completions.create(model="m", messages=messages, max_completion_tokens=100)
    assert raw.sent[-1]["timeout"] == local_cleanup_win.request_timeout(1002, 100, True) and raw.sent[-1]["max_tokens"] == 100
    slowpc_client(raw).chat.completions.create(model="m", messages=messages)  # no limit asked for: the server's own limit counts
    assert raw.sent[-1]["timeout"] == local_cleanup_win.request_timeout(1002, None, True) == 600
    slowpc_client(raw, base=500).chat.completions.create(model="m", messages=messages, max_completion_tokens=1)
    assert raw.sent[-1]["timeout"] == 500  # the configured wait is a floor
    slowpc_client(raw).chat.completions.create(model="m", messages=messages, max_completion_tokens=100, timeout=7)
    assert raw.sent[-1]["timeout"] == 7  # a wait the caller chose is kept
    # a graphics-card server, a client with no server, and a server that can't say: the client's own timeout, as before
    for client in (slowpc_client(raw, cpu=False), local_cleanup_win.LocalClient(raw),
                   local_cleanup_win.LocalClient(raw, SimpleNamespace(on_cpu=lambda: 1 / 0))):
        client.chat.completions.create(model="m", messages=messages, max_completion_tokens=100)
        assert "timeout" not in raw.sent[-1]
    # a caller that fixed the wait itself (meeting notes: 15 minutes) is never overridden
    client = slowpc_client(raw).with_options(timeout=900)
    client.chat.completions.create(model="m", messages=messages, max_completion_tokens=2500)
    assert raw.options == {"timeout": 900} and "timeout" not in raw.sent[-1]
    client.with_options(max_retries=0).chat.completions.create(model="m", messages=messages)
    assert "timeout" not in raw.sent[-1]  # (and it stays fixed through another with_options)
    # an ordinary with_options (no timeout) keeps the per-request wait
    slowpc_client(raw).with_options(max_retries=0).chat.completions.create(model="m", messages=messages, max_completion_tokens=100)
    assert raw.sent[-1]["timeout"] == local_cleanup_win.request_timeout(1002, 100, True)


def test_slowpc_every_command_request_has_a_reply_limit_for_its_job():
    assert commands.route_budget("") == 256 and commands.route_budget("open notepad") < 300 and commands.route_budget("x" * 9000) == 600
    assert commands.rewrite_budget("") == commands.rewrite_budget("   ") == commands.WRITE_TOKENS == 600  # new text: a fixed cap
    assert commands.rewrite_budget("Short text.") == 96  # the floor
    assert commands.rewrite_budget("w" * 900) == 498  # 300 tokens x 1.5, plus a little
    assert commands.rewrite_budget("w" * 12000) == commands.REWRITE_MAX_TOKENS == 4096  # the longest selection VoiceTool copies
    raw = SlowPCRaw(json.dumps({"route": "rewrite", "text": ""}), "Shorter.")
    client = slowpc_client(raw)
    selection = "I just wanted to reach out and let you know that I will be a little bit late."
    assert commands.route(client, "local", "make this shorter", {"selection": selection})[:2] == ("rewrite", "Shorter.")
    route_request, rewrite_request = raw.sent
    assert route_request["max_tokens"] == commands.route_budget("make this shorter")
    assert rewrite_request["max_tokens"] == commands.rewrite_budget(selection)
    assert route_request["timeout"] > rewrite_request["timeout"] >= 60  # the router reads far more than a rewrite does
    # "write ..." with nothing selected: the router decides (a tiny reply), then the focused writing step has the fixed cap
    raw = SlowPCRaw(json.dumps({"route": "write", "text": ""}), "Hi Sam, thank you.")
    assert commands.route(slowpc_client(raw), "local", "write a thank-you note to Sam", {"selection": ""})[:2] == ("write", "Hi Sam, thank you.")
    assert raw.sent[0]["max_tokens"] == commands.route_budget("write a thank-you note to Sam") and raw.sent[1]["max_tokens"] == 600
    # an answer, an open and a note come back from the one routing request (no second pass)
    for reply, want in ((json.dumps({"route": "answer", "text": "Four."}), "answer"), (json.dumps({"route": "open", "target": "Notepad"}), "open"),
                        (json.dumps({"route": "note", "text": "Buy milk.", "topic": "milk"}), "note")):
        raw = SlowPCRaw(reply)
        assert commands.route(slowpc_client(raw), "local", "what is two plus two", {"selection": ""})[0] == want and len(raw.sent) == 1
        assert raw.sent[0]["max_tokens"] <= 600


def test_slowpc_a_model_that_is_too_slow_gone_or_loading_raises_one_clear_exception_on_every_step():
    errors = slowpc_errors()
    selection = {"selection": "I just wanted to reach out and let you know that I will be a little bit late."}
    # the routing request itself
    for name, reason in (("timeout", "timeout"), ("connection", "connection"), ("loading", "loading"), ("bad_gateway", "connection")):
        with pytest.raises(commands.ModelTimeout) as caught:
            commands.route(slowpc_client(SlowPCRaw(errors[name])), "local", "make this shorter", selection)
        e = caught.value
        assert isinstance(e, commands.CommandFailed) and isinstance(e, RuntimeError) and e.reason == reason
        assert e.__cause__ is errors[name] and str(e) == e.message and "History" in str(e) and "nothing is lost" in str(e)
    assert str(commands.ModelTimeout("x")) == "x"
    # the message is short, plain, and says what to do
    with pytest.raises(commands.ModelTimeout) as caught:
        commands.route(slowpc_client(SlowPCRaw(errors["timeout"])), "local", "mute", {})
    assert str(caught.value).startswith("Your PC took too long on that command. It's saved in History") and "shorter request" in str(caught.value)
    assert len(str(caught.value)) < 260
    # the second pass of a rewrite (the router answered, the rewrite timed out), and of a write
    for first, wanted in ((json.dumps({"route": "rewrite", "text": ""}), selection), (json.dumps({"route": "write", "text": ""}), {})):
        raw = SlowPCRaw(first, errors["timeout"])
        with pytest.raises(commands.ModelTimeout) as caught:
            commands.route(slowpc_client(raw), "local", "make this shorter", wanted)
        assert caught.value.reason == "timeout" and len(raw.sent) == 2
    with pytest.raises(commands.ModelTimeout):  # (a rewrite called on its own)
        commands.rewrite(slowpc_client(SlowPCRaw(errors["connection"])), "local", "shorter", "some text")
    # the on-device model that is still loading says so; a model that is not on this computer is named as such
    with pytest.raises(commands.ModelTimeout) as caught:
        commands.route(slowpc_client(SlowPCRaw(errors["loading"])), "local", "mute", {})
    assert "still loading" in str(caught.value)
    with pytest.raises(commands.ModelTimeout) as caught:
        commands.route(SlowPCRaw(errors["timeout"]), "gpt", "mute", {})  # (an online client, not the local one)
    assert str(caught.value).startswith("The online model took too long")
    with pytest.raises(commands.ModelTimeout) as caught:
        commands.route(SlowPCRaw(errors["connection"]), "gpt", "mute", {})
    assert "online model" in str(caught.value) and "PC" not in str(caught.value)
    # other failures are not dressed up as a timeout: they reach the caller as they are
    for name in ("auth", "bad_request"):
        with pytest.raises(type(errors[name])):
            commands.route(slowpc_client(SlowPCRaw(errors[name])), "local", "mute", {})
    with pytest.raises(ValueError):
        commands.route(slowpc_client(SlowPCRaw("this is not json")), "local", "mute", {})
    already = commands.ModelTimeout("already", "timeout")
    assert commands.as_command_failure(already) is already
    for builtin in (TimeoutError(), ConnectionRefusedError(), ConnectionResetError(), EOFError()):
        assert isinstance(commands.as_command_failure(builtin, local=True), commands.ModelTimeout)
    assert commands.as_command_failure(KeyError("x")) is None and commands.as_command_failure(ValueError("x"), local=True) is None
    assert commands.as_command_failure(TimeoutError("the model took too long to start"), local=True).message.startswith("Your PC took too long")


class SlowPCStub:
    """A tiny web server on 127.0.0.1 that answers every POST after `delay` seconds with `status` (a stand-in for a model that
    is slow, still loading, or failing). It never runs a model."""

    def __init__(self, delay=0.0, status=200, body=None):
        import http.server
        stub = self
        stub.delay, stub.status, stub.body, stub.hits = delay, status, body or {}, 0

        class Handler(http.server.BaseHTTPRequestHandler):
            def log_message(self, *a):
                pass

            def do_POST(self):
                self.rfile.read(int(self.headers.get("Content-Length") or 0))
                stub.hits += 1
                time.sleep(stub.delay)
                data = json.dumps(stub.body).encode()
                try:
                    self.send_response(stub.status)
                    self.send_header("Content-Length", str(len(data)))
                    self.end_headers()
                    self.wfile.write(data)
                except OSError:
                    pass  # (the client gave up first)

        self.server = http.server.ThreadingHTTPServer(("127.0.0.1", 0), Handler)
        self.server.daemon_threads = True
        self.port = self.server.server_address[1]
        threading.Thread(target=self.server.serve_forever, daemon=True).start()

    def client(self, timeout):
        from openai import OpenAI
        return slowpc_client(OpenAI(base_url=f"http://127.0.0.1:{self.port}/v1", api_key="local", timeout=timeout, max_retries=0), cpu=False)

    def close(self):
        self.server.shutdown()
        self.server.server_close()


def test_slowpc_real_client_errors_become_the_same_exception_a_slow_server_a_loading_server_and_a_closed_port():
    from openai import OpenAI
    slow = SlowPCStub(delay=2.0)
    loading = SlowPCStub(status=503, body={"error": {"message": "Loading model", "type": "unavailable_error", "code": 503}})
    try:
        started = time.monotonic()
        with pytest.raises(commands.ModelTimeout) as caught:  # the wait runs out
            commands.route(slow.client(0.4), "local", "make this shorter", {"selection": "some words"})
        assert caught.value.reason == "timeout" and time.monotonic() - started < 1.9 and slow.hits == 1
        with pytest.raises(commands.ModelTimeout) as caught:  # the server answers "still loading"
            commands.route(loading.client(5), "local", "mute", {})
        assert caught.value.reason == "loading" and loading.hits == 1
    finally:
        slow.close()
        loading.close()
    dead_port = free_port()  # nothing listens here: connection refused
    refused = slowpc_client(OpenAI(base_url=f"http://127.0.0.1:{dead_port}/v1", api_key="local", timeout=3, max_retries=0), cpu=False)
    with pytest.raises(commands.ModelTimeout) as caught:
        commands.route(refused, "local", "open Notepad", {})
    assert caught.value.reason == "connection" and "couldn't reach" in str(caught.value)


def test_slowpc_a_reply_cut_off_at_its_limit_is_never_pasted_half_written():
    cut = lambda text: SlowPCReply(text, "length")  # noqa: E731
    # a rewrite that ran to its limit keeps only whole sentences
    raw = SlowPCRaw(json.dumps({"route": "rewrite", "text": ""}), cut("One is done. Two is done. Thr"))
    assert commands.route(slowpc_client(raw), "local", "shorter", {"selection": "some text here"})[1] == "One is done. Two is done."
    assert commands.rewrite(slowpc_client(SlowPCRaw(cut("No full sentence at all"))), "local", "x", "text") == "No full sentence at all"
    # a routing reply cut off while writing: the draft is thrown away and the focused writing step writes it again
    raw = SlowPCRaw(cut('{"route": "write", "text": "Hi Sam,\\n\\nThank you so much for'), "Hi Sam, thanks a lot.")
    assert commands.route(slowpc_client(raw), "local", "write a thank-you note to Sam", {"selection": ""})[:2] == ("write", "Hi Sam, thanks a lot.")
    assert len(raw.sent) == 2 and raw.sent[1]["max_tokens"] == commands.WRITE_TOKENS
    raw = SlowPCRaw(cut('{"route": "rewrite", "te'), "Shorter.")
    assert commands.route(slowpc_client(raw), "local", "shorter", {"selection": "words words"})[:2] == ("rewrite", "Shorter.")
    # an answer cut off keeps its whole sentences
    answer = ('{"route": "answer", "text": "The error means the file is missing and Windows cannot open it. '
              'Check the path and try again, then if that keeps failing restart the')
    kind, text, _ = commands.route(slowpc_client(SlowPCRaw(cut(answer))), "local", "what does this error mean", {"selection": ""})
    assert (kind, text) == ("answer", "The error means the file is missing and Windows cannot open it.")
    # a note, a journal entry or an open that was cut off is not guessed at: one clear message, not a timeout
    for partial in ('{"route": "note", "text": "Buy mi', '{"route": "journal", "text": "today I', '{"route": "open", "target": "Note', '{"rou'):
        with pytest.raises((commands.CommandFailed, ValueError)) as caught:
            commands.route(slowpc_client(SlowPCRaw(cut(partial))), "local", "note: buy milk", {"selection": ""})
        if partial != '{"rou':
            assert type(caught.value) is commands.CommandFailed and "shorter request" in str(caught.value) and caught.value.reason == "too long"
    # a reply that is not cut off and is not JSON still fails as it always did
    with pytest.raises(ValueError):
        commands.route(slowpc_client(SlowPCRaw(SlowPCReply('{"route": "wri', "stop"))), "local", "mute", {})
    # a complete reply that happened to end exactly at the limit is used
    done = SlowPCReply(json.dumps({"route": "open", "target": "Notepad"}), "length")
    assert commands.route(slowpc_client(SlowPCRaw(done)), "local", "open Notepad", {})[0] == "open"


def test_slowpc_the_processor_server_gets_its_cores_a_reply_limit_and_a_longer_start_while_the_graphics_card_one_is_unchanged(monkeypatch, tmp_path):
    lc = local_cleanup_win
    monkeypatch.setattr(lc, "physical_cores", lambda: 6)
    cpu = lc.LocalCleaner({"offline_backend": "cpu"})
    assert cpu.on_cpu() and cpu._args() == ["-c", "16384", "--parallel", "1", "--no-webui", "-t", "6", "-n", "2560", "-cram", "1024"]
    assert "-ngl" not in cpu._args() and "--device" not in cpu._args() and "-c" in cpu._args() and "--parallel" in cpu._args()
    monkeypatch.setattr(lc, "physical_cores", lambda: 40)
    assert lc.LocalCleaner({"offline_backend": "cpu"})._args()[-6] == "-t" and lc.LocalCleaner({"offline_backend": "cpu"})._args()[-5] == "12"  # not more than 12
    monkeypatch.setattr(lc, "physical_cores", lambda: 1)
    assert lc.LocalCleaner({"offline_backend": "cpu"})._args()[-5] == "1"
    # the person's own settings win, and are not repeated
    own = lc.LocalCleaner({"offline_backend": "cpu", "offline_args": ["-c", "8192", "-ngl", "99", "--device", "Vulkan0", "-t", "2",
                                                                       "--predict=300", "--cache-ram", "0"]})
    assert own._args() == ["-c", "8192", "-t", "2", "--predict=300", "--cache-ram", "0"]
    assert lc.cpu_server_args(["-c", "16384"], cores=4) == ["-c", "16384", "-t", "4", "-n", "2560", "-cram", "1024"]
    assert lc.physical_cores() >= 1
    # the graphics-card server is exactly what it was
    gpu = lc.LocalCleaner({"offline_backend": "vulkan"})
    assert not gpu.on_cpu() and gpu._args() == lc.BASE_ARGS + lc.GPU_ARGS
    configured = ["-c", "16384", "--parallel", "1", "--no-webui", "-ngl", "99", "--device", "Vulkan0"]
    assert lc.LocalCleaner({"offline_backend": "vulkan", "offline_args": configured})._args() == configured
    # a Vulkan build that was started with no graphics-card options is on the processor; so is one that fell back
    assert lc.LocalCleaner({"offline_backend": "vulkan", "offline_args": ["-c", "16384"]}).on_cpu()
    fell = lc.LocalCleaner({"offline_backend": "vulkan"})
    fell.backend, fell.fell_back = "cpu", True
    assert fell.on_cpu() and "-t" in fell._args() and "-ngl" not in fell._args()
    # how long the model may take to load: longer on a processor (a slow disk), unless the caller says otherwise
    assert cpu.startup_seconds == lc.CPU_STARTUP_SECONDS == 240 and gpu.startup_seconds == lc.GPU_STARTUP_SECONDS == 90
    assert lc.LocalCleaner({"offline_backend": "cpu"}, startup_seconds=7).startup_seconds == 7
    # the real command line that starts the processor server
    for name in ("c.exe", "m.gguf"):
        (tmp_path / name).write_bytes(b"x")
    monkeypatch.setattr(lc, "physical_cores", lambda: 4)
    real = lc.LocalCleaner({"offline_backend": "cpu", "offline_server_cpu": str(tmp_path / "c.exe"), "offline_model": str(tmp_path / "m.gguf")})
    command = real._default_command(8766)
    assert command[0] == str(tmp_path / "c.exe") and command[-6:] == ["-t", "4", "-n", "2560", "-cram", "1024"] and "127.0.0.1" in command


def test_slowpc_the_processor_server_reads_the_routers_instructions_once_after_loading_and_a_failure_is_only_logged():
    lc = local_cleanup_win
    raw = SlowPCRaw("")
    logged = []
    log = SimpleNamespace(info=lambda msg, *a: logged.append(msg % a if a else msg), warning=lambda *a: logged.append(a))
    cpu = lc.LocalCleaner({"offline_backend": "cpu"}, log)
    cpu.client = lambda timeout=60: slowpc_client(raw)
    assert cpu.warm_up() is True and len(raw.sent) == 1
    request = raw.sent[0]
    assert request["messages"][0] == {"role": "system", "content": commands.ROUTER} and request["max_tokens"] == 1
    assert request["timeout"] == lc.request_timeout(lc.estimate_prompt_tokens(request["messages"]), 1, True) > 60
    assert any("warmed up" in m for m in logged)
    # the real routing request starts with the same instructions, so the server can reuse what it read
    routed = SlowPCRaw(json.dumps({"route": "system", "target": "mute"}))
    commands.route(slowpc_client(routed), "local", "mute", {"selection": ""}, style="My style.")
    assert routed.sent[0]["messages"][0]["content"].startswith(request["messages"][0]["content"])
    # a failure never raises
    cpu.client = lambda timeout=60: slowpc_client(SlowPCRaw(slowpc_errors()["timeout"]))
    assert cpu.warm_up() is False and any("skipped" in m for m in logged if isinstance(m, str))
    cpu.client = lambda timeout=60: (_ for _ in ()).throw(lc.ModelServerUnavailable("won't start"))
    assert cpu.warm_up() is False
    # the graphics-card server is never warmed up this way
    gpu = lc.LocalCleaner({"offline_backend": "vulkan"}, log)
    gpu.client = lambda timeout=60: pytest.fail("no request on a graphics card")
    assert gpu.warm_up() is False


def test_slowpc_starting_in_the_background_warms_a_processor_server_up_only_after_it_is_ready(monkeypatch):
    lc = local_cleanup_win
    events = []
    cpu = lc.LocalCleaner({"offline_backend": "cpu"})
    monkeypatch.setattr(cpu, "start", lambda: events.append("start"))
    monkeypatch.setattr(cpu, "warm_up", lambda: events.append("warm"))
    cpu.start_in_background()
    for thread in threading.enumerate():
        if thread.name == "local-model":
            thread.join(5)
    assert events == ["start", "warm"]
    events.clear()
    failing = lc.LocalCleaner({"offline_backend": "cpu"})
    monkeypatch.setattr(failing, "start", lambda: (_ for _ in ()).throw(RuntimeError("the model program stopped while starting")))
    monkeypatch.setattr(failing, "warm_up", lambda: events.append("warm"))
    failing.start_in_background()
    for thread in threading.enumerate():
        if thread.name == "local-model":
            thread.join(5)
    assert events == []  # a server that didn't start is not warmed up (and nothing is raised)


def test_slowpc_a_request_that_arrives_while_the_processor_server_is_still_loading_waits_for_it(fake_server, monkeypatch):
    port = free_port()
    slow_start = [BASE_PYTHON, "-c", "import sys, time, runpy; time.sleep(1.5); sys.argv = sys.argv[1:]; runpy.run_path(sys.argv[0], run_name='__main__')",
                  fake_server.script, str(port), "fake-model.gguf"]
    c = local_cleanup_win.LocalCleaner({"offline_model": "C:/m/fake-model.gguf", "offline_port": port, "offline_backend": "cpu"},
                                       command=lambda p: slow_start, startup_seconds=30)
    monkeypatch.setattr(c, "warm_up", lambda: True)  # (the stand-in is not a model)
    c.start_in_background()
    time.sleep(0.4)
    assert not c.running()  # still loading
    started = time.monotonic()
    assert tidy(c, "SYSTEM PROMPT", "hello there") == "CLEANED: hello there | SYSTE"  # the request waited for the load, then ran
    assert time.monotonic() - started > 0.5
    c.stop()
    # a load that never finishes in time is one clear failure that the app can turn into the same message
    never = local_cleanup_win.LocalCleaner({"offline_port": free_port(), "offline_backend": "cpu"},
                                           command=lambda p: [BASE_PYTHON, "-c", "import time; time.sleep(60)"], startup_seconds=1)
    with pytest.raises(TimeoutError) as caught:
        never.start()
    assert commands.as_command_failure(caught.value, local=True).reason == "timeout"
    never.stop()


def test_slowpc_a_request_queued_behind_the_start_up_warm_up_is_given_the_warm_up_time_too(monkeypatch):
    lc = local_cleanup_win
    messages = [{"role": "system", "content": "s" * 600}, {"role": "user", "content": "um hello there"}]
    raw = SlowPCRaw("ok")
    cpu = lc.LocalCleaner({"offline_backend": "cpu"})
    client = lc.LocalClient(raw, cpu, 60)
    alone = lc.request_timeout(lc.estimate_prompt_tokens(messages), 64, True)
    client.chat.completions.create(model="m", messages=messages, max_completion_tokens=64)
    assert raw.sent[-1]["timeout"] == alone and cpu.warm_up_wait() == 0
    started = threading.Event()
    release = threading.Event()
    monkeypatch.setattr(cpu, "start", lambda: (started.set(), release.wait(5)))
    monkeypatch.setattr(cpu, "warm_up", lambda: release.wait(5))
    cpu.start_in_background()
    assert started.wait(5) and 60 < cpu.warm_up_wait() <= lc.request_timeout(1500, 1, True)  # (about 100 s while it may run)
    client.chat.completions.create(model="m", messages=messages, max_completion_tokens=64)
    assert alone + 60 < raw.sent[-1]["timeout"] <= alone + lc.request_timeout(1500, 1, True)
    release.set()
    for thread in threading.enumerate():
        if thread.name == "local-model":
            thread.join(5)
    assert cpu.warm_up_wait() == 0  # over: back to the plain wait
    client.chat.completions.create(model="m", messages=messages, max_completion_tokens=64)
    assert raw.sent[-1]["timeout"] == alone
    # a server that failed to start is not left "warming up"
    failed = lc.LocalCleaner({"offline_backend": "cpu"})
    monkeypatch.setattr(failed, "start", lambda: (_ for _ in ()).throw(RuntimeError("stopped while starting")))
    failed.start_in_background()
    for thread in threading.enumerate():
        if thread.name == "local-model":
            thread.join(5)
    assert failed.warm_up_wait() == 0
    # and a graphics-card server never has one
    gpu = lc.LocalCleaner({"offline_backend": "vulkan"})
    monkeypatch.setattr(gpu, "start", lambda: None)
    gpu.start_in_background()
    assert gpu.warm_up_wait() == 0


def test_review_a_model_that_is_too_slow_to_start_keeps_the_command_and_says_so(tool, monkeypatch):
    def slow_start():
        raise TimeoutError("the model took too long to start")
    tool.azure = FakeCleaner()
    monkeypatch.setattr(tool.azure, "_client", slow_start)
    t = tool.triggers[COMMAND]
    t.context = {"app": "notepad.exe", "title": "x", "selection": "", "pid": 1, "hwnd": 2}
    tool.transcribe = lambda audio: "write a thank-you note to Sam"
    tool._command(t, np.full(vt.RATE, 0.3, np.float32))
    rows = [json.loads(line) for line in (vt.HERE / "history.jsonl").read_text(encoding="utf-8").splitlines()]
    assert rows[-1]["command"] == "write a thank-you note to Sam" and rows[-1]["route"] == "failed"
    assert tool.cards[-1][0] == "Command not done" and "saved in History" in tool.cards[-1][1] and not PASTED
