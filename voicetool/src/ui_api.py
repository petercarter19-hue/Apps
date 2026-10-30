"""What the VoiceTool app window can ask for. The window (ui/index.html) shows the page; the small "host" object does the
things that depend on this computer (opening files, pausing, quitting, start at login, downloading the tidy-up model).

The page calls `vt.call(name, args)`; the host passes that to `API.call(name, args)`, which returns plain JSON data.
Nothing here touches keys, the microphone, windows or the clipboard.
"""
import json
import re
from datetime import date
from pathlib import Path

import commands
import corrections
import styles

HERE = Path(__file__).resolve().parent
HISTORY = HERE / "history.jsonl"
WORDS = HERE / "words.txt"

# (section, config key, title, key badge, what it does) - one list for the Settings page
FEATURES = [
    ("Talking", "dictate_enabled", "Dictation", "Hold or double-tap Left Alt",
     "Hold, talk, let go: your words appear wherever the cursor is. Double-tap for hands-free; tap again to stop."),
    ("Talking", "cleanup", "Tidy-up", None,
     "Removes ums, fixes punctuation, applies your word list. Runs on this computer."),
    ("Talking", "show_light", "Recording light", None,
     "The small 'Listening' pill at the bottom of the screen."),
    ("Talking", "beeps", "Beeps", None,
     "Short sounds when it starts listening, locks hands-free, and finishes."),
    ("Talking", "offline_only", "Keep everything on this computer", None,
     "Tidy-up, commands and meeting notes all run on a small model on this computer, so nothing you say leaves it. "
     "Turning this off sends them to Azure OpenAI instead: an advanced option that needs an Azure endpoint in "
     "config.json and the Azure CLI signed in."),
    ("Talking", "speak_answers", "Spoken answers", None,
     "Reads answers out loud (questions, history, Claude), as well as showing them. Esc stops it."),
    ("Doing things", "command_enabled", "Command key", "Hold or double-tap Left Ctrl",
     "Hold, say what you want, let go. Rewrites what's selected, writes something, or answers. "
     "Double-tap for longer instructions; tap again to stop."),
    ("Doing things", "claude_actions", "Email and calendar (Claude)", None,
     "Lets the Command key find emails, check calendars and write reply drafts. Drafts only: you always press Send. "
     "Needs Claude Code installed and signed in on your own Claude plan. Off until you turn it on."),
    ("Doing things", "claude_keep_running", "Keep Claude ready", None,
     "Keeps one Claude session open so email and calendar requests are quicker. Only used when the switch above is on."),
    ("Meetings", "meeting_notes", "Meeting notes", "Icon menu › Start meeting notes",
     "Records you and the call, then writes a summary, decisions, action items and transcript to a Word document in "
     "your Meeting notes folder. It always asks first: make sure everyone on the call agrees to be recorded. "
     "Off until you turn it on."),
    ("Meetings", "meeting_auto_offer", "Offer when a call starts", None,
     "When Teams, Zoom or another call app starts using the mic, a card asks whether to take notes, and they finish "
     "by themselves when the call ends. Needs Meeting notes."),
    ("Meetings", "meeting_followup", "Follow-up email draft", None,
     "After the meeting, Claude finds it in your calendar and drafts a follow-up to the attendees. Draft only. "
     "Needs Meeting notes and Email and calendar (Claude)."),
    ("Learning", "use_style", "Write like me", None,
     "Drafts and 'write' commands follow your style guide (your own words, plus what it learned from your sent emails "
     "if Claude is on)."),
    ("Learning", "typed_fix", "Type a fix", "Shift + tap Left Ctrl",
     "Select a misheard word, then type what it should be. Fixed now and learned for next time."),
    ("Learning", "auto_learn", "Learn mishearings automatically", None,
     "When the tidy-up fixes the same mishearing into one of your words twice, it remembers it."),
    ("Learning", "save_recordings", "Save recordings", None,
     "Keeps each recording on this computer (it is never uploaded). Off by default."),
]
# Everything else is on unless the config says otherwise.
OFF_BY_DEFAULT = {"meeting_notes", "meeting_auto_offer", "meeting_followup", "claude_actions", "claude_keep_running",
                  "speak_answers", "save_recordings"}
CONSENT = {"meeting_notes"}  # turning one of these ON needs the person to confirm first (confirmed=True)

LANGUAGES = ["Spanish", "French", "German", "Italian", "Portuguese", "Dutch", "Japanese", "Chinese (Simplified)",
             "Korean", "Arabic", "Hindi"]


def read_history():
    """Every history line, oldest first. Bad lines are skipped."""
    rows = []
    try:
        for line in HISTORY.read_text(encoding="utf-8").splitlines():
            try:
                rows.append(json.loads(line))
            except json.JSONDecodeError:
                pass
    except FileNotFoundError:
        pass
    return rows


def _fixes_in(row):
    """How many learned fixes a dictation had (rows from before the counter have none)."""
    n = row.get("fixes", 0)
    return n if isinstance(n, int) and not isinstance(n, bool) and n > 0 else 0


def describe(row):
    """One history line in a shape the page can show."""
    if "meeting" in row:
        kind = "meeting"
        text = f"{row.get('meeting') or 'Meeting'} ({row.get('minutes') or '?'} min). Notes: {Path(row.get('notes') or '').name}"
        said = ""
    elif "command" in row:
        kind = "claude" if row.get("route") == "claude" else "command"
        text = row.get("result") or ""
        said = row.get("command") or ""
    else:
        kind = "dictate"
        text = row.get("text") or ""
        said = row.get("raw") or ""
    return {"at": row.get("at", ""), "kind": kind, "route": row.get("route", ""), "app": row.get("app") or "",
            "said": said, "text": text, "phone": row.get("source") == "phone", "fixes": _fixes_in(row)}


class API:
    def __init__(self, host, config, save_config):
        """host: the platform's helper (see the methods it's asked for below). config: the live settings dict.
        save_config: writes config.json."""
        self.host, self.config, self.save_config = host, config, save_config

    def call(self, name, args=None):
        method = getattr(self, "api_" + str(name), None)
        if method is None:
            raise ValueError(f"unknown request {name!r}")
        return method(**(args or {}))

    def _azure_set_up(self):
        return bool(str(self.config.get("azure_openai_base_url") or "").strip())

    def _claude_ready(self):
        """Email and calendar are on and Claude Code is installed and signed in (never runs a program when they're off)."""
        return bool(self.config.get("claude_actions", False)) and bool(self.host.claude_signed_in())

    # ---------- Home ----------

    def api_overview(self):
        rows = [describe(r) for r in read_history() if (r.get("text") or r.get("command") or r.get("result"))]
        today = date.today().isoformat()
        todays = [r for r in rows if r["at"].startswith(today)]
        claude_on = bool(self.config.get("claude_actions", False))
        return {
            "enabled": self.config.get("enabled", True),
            "computer": self.host.computer_name(),
            "claude_actions": claude_on,
            "claude_signed_in": bool(claude_on and self.host.claude_signed_in()),
            "style_learned": commands.style_learned(),
            "use_style": self.config.get("use_style", True),
            "first_run": bool(self.config.get("first_run", False)),
            "setup": getattr(self.host, "setup_status", lambda: {"missing": False, "running": False, "failed": False})(),
            "today": {"dictations": sum(r["kind"] == "dictate" for r in todays),
                      "commands": sum(r["kind"] != "dictate" for r in todays),
                      "words": sum(len(r["text"].split()) for r in todays if r["kind"] == "dictate"),
                      "fixes": sum(r["fixes"] for r in todays)},
            "recent": list(reversed(rows[-6:])),
            "fixes": len(corrections.load()),
            "words": len(self.api_words()["words"]),
        }

    def api_set_paused(self, paused):
        self.config["enabled"] = not paused
        self.save_config(self.config)
        self.host.paused(bool(paused))
        return {"enabled": self.config["enabled"]}

    def api_dismiss_welcome(self):
        """The Welcome card on Home was closed: it won't open again."""
        self.config["first_run"] = False
        self.save_config(self.config)
        return {"first_run": False}

    def api_finish_setup(self):
        """The "Finish setup" button on Home: download the tidy-up model (the host starts the downloader)."""
        start = getattr(self.host, "finish_setup", None)
        if start is None:
            return {"started": False, "message": "Setup can't be finished from here."}
        return start()

    # ---------- History ----------

    def api_history(self, query="", kind="all", limit=200):
        words = [w for w in re.split(r"\s+", (query or "").lower()) if w]
        out = []
        for row in reversed(read_history()):
            r = describe(row)
            if not (r["text"] or r["said"]):
                continue
            if kind != "all" and r["kind"] != kind:
                continue
            hay = f"{r['text']} {r['said']} {r['app']}".lower()
            if all(w in hay for w in words):
                out.append(r)
                if len(out) >= limit:
                    break
        return {"items": out, "total": len(out)}

    # ---------- Settings ----------

    def api_settings(self):
        model_here = getattr(self.host, "offline_available", lambda: True)()
        features = []
        for s, k, t, b, d in FEATURES:
            on = bool(self.config.get(k, k not in OFF_BY_DEFAULT))
            if k == "offline_only" and not model_here:
                d += " (The tidy-up model isn't downloaded yet: click Finish setup on the Home page.)"
            if k == "claude_actions" and on and not self.host.claude_signed_in():
                d += " (Claude Code isn't installed or signed in on this computer yet, so this does nothing for now.)"
            features.append({"section": s, "key": k, "title": t, "badge": b, "text": d, "on": on})
        current = self.config.get("translate_to", "")
        import voice_actions
        return {"features": features,
                "translate_to": current, "languages": LANGUAGES + ([current] if current and current not in LANGUAGES else []),
                "read_voice": self.config.get("read_voice", ""), "read_rate": self.config.get("read_rate", "normal"),
                "voices": self.host.list_voices(), "natural_voice": bool(self.config.get("tts_deployment")),
                "offline_available": bool(model_here), "azure_set_up": self._azure_set_up(),  # (for Home's "Runs on your PC" card)
                "notes_destination": "folder",
                "notes_folder": self.config.get("notes_folder") or str(voice_actions.default_notes_folder()),
                "note_destinations": [{"id": d, "label": voice_actions.NOTE_DESTINATIONS[d]}
                                      for d in self.host.note_destinations()],
                "startup": self.host.startup_enabled(), "enabled": self.config.get("enabled", True)}

    def api_set_setting(self, key, on, confirmed=False):
        if key not in {f[1] for f in FEATURES}:
            raise ValueError(f"unknown setting {key!r}")
        on = bool(on)
        if on and key in CONSENT and not confirmed:
            # Nothing is changed until the person has read the reminder and said yes (the page shows it, then asks again).
            import meetings
            return {"key": key, "on": bool(self.config.get(key, False)), "needs_confirm": meetings.CONSENT_REMINDER}
        if key == "offline_only" and not on and not self._azure_set_up():
            raise ValueError("Azure isn't set up (there is no Azure endpoint in config.json), so this has to stay on.")
        self.config[key] = on
        self.save_config(self.config)
        if key == "claude_keep_running" and not on:
            self.host.stop_claude()
        if key == "claude_actions" and not on:
            self.host.stop_claude()
        if hasattr(self.host, "setting_changed"):  # lets the host react (the on-device model's server starts or stops)
            self.host.setting_changed(key, on)
        return {"key": key, "on": self.config[key]}

    def api_set_translate(self, language):
        language = str(language or "").strip()
        if language and language not in LANGUAGES and language != self.config.get("translate_to"):
            raise ValueError("unknown language")  # (a language said aloud that isn't in the list can be kept as it is)
        self.config["translate_to"] = language
        self.save_config(self.config)
        return {"translate_to": language}

    def api_set_read_voice(self, voice="", rate="normal"):
        voices = {v["name"] for v in self.host.list_voices()}
        if voice and voice not in voices:
            raise ValueError("that voice isn't on this computer")
        if rate not in ("slow", "normal", "fast"):
            raise ValueError("unknown speed")
        self.config.update(read_voice=voice, read_rate=rate)
        self.save_config(self.config)
        return {"read_voice": voice, "read_rate": rate}

    def api_set_notes(self, destination="folder", folder=""):
        if destination not in self.host.note_destinations():
            raise ValueError("unknown destination")
        self.config["notes_destination"] = destination
        if folder:
            self.config["notes_folder"] = str(Path(folder).expanduser())
        self.save_config(self.config)
        return {"notes_destination": destination, "notes_folder": self.config.get("notes_folder", "")}

    def api_choose_notes_folder(self):
        path = self.host.choose_folder(self.config.get("notes_folder", ""))
        if path:
            self.config["notes_folder"] = path
            self.save_config(self.config)
        return {"notes_folder": self.config.get("notes_folder", "")}

    def api_open_notes(self):
        import voice_actions
        target = Path(self.config.get("notes_folder") or voice_actions.default_notes_folder())
        target.mkdir(parents=True, exist_ok=True)
        self.host.open_path(target)
        return {}

    def api_try_voice(self):
        self.host.read_aloud("Hi there. This is how VoiceTool will read to you. Press Escape to stop me.")
        return {}

    def api_set_startup(self, on):
        self.host.set_startup(bool(on))
        return {"startup": self.host.startup_enabled()}

    # ---------- Words, fixes and style ----------

    def api_words(self):
        try:
            lines = WORDS.read_text(encoding="utf-8").splitlines()
        except FileNotFoundError:
            lines = []
        return {"words": [w.strip() for w in lines if w.strip() and not w.startswith("#")]}

    def api_save_words(self, words):
        clean, seen = [], set()
        for w in words:
            w = " ".join(str(w).split())
            if w and w.lower() not in seen:
                seen.add(w.lower())
                clean.append(w)
        WORDS.write_text("# One word or name per line. The cleanup step makes sure these are spelled exactly like "
                         "this.\n" + "\n".join(clean) + "\n", encoding="utf-8")
        return {"words": clean}

    def api_fixes(self):
        try:
            data = json.loads(corrections.FILE.read_text(encoding="utf-8"))
        except (FileNotFoundError, json.JSONDecodeError):
            data = []
        if isinstance(data, dict):
            data = [{"heard": h, "meant": m} for h, m in data.items()]
        return {"fixes": [{"heard": d.get("heard", ""), "meant": d.get("meant", ""), "added": d.get("added", ""),
                           "source": d.get("source", "")} for d in reversed(data)]}

    def api_edit_word(self, old, new):
        new = " ".join(str(new).split())
        words = self.api_words()["words"]  # read fresh, so a word learned meanwhile isn't lost
        others = [w for w in words if w.lower() != str(old).strip().lower()]
        reason = ("Type a word, or use ✕ to remove it." if not new else
                  "That word is already on your list." if new.lower() in {w.lower() for w in others} else "")
        if reason:
            return {"saved": False, "reason": reason, "words": words}
        at = next((i for i, w in enumerate(words) if w.lower() == str(old).strip().lower()), len(words))
        return {"saved": True, "reason": "", **self.api_save_words(words[:at] + [new] + words[at + 1:])}

    def api_add_fix(self, heard, meant):
        ok = corrections.add(heard, meant, "app window")
        return {"added": ok, "reason": corrections.problem(heard, meant), **self.api_fixes()}

    def api_edit_fix(self, old_heard, heard, meant):
        ok = corrections.edit(old_heard, heard, meant, "edited")
        return {"saved": ok, "reason": corrections.problem(heard, meant), **self.api_fixes()}

    def api_delete_fix(self, heard):
        with corrections._lock:
            data = json.loads(corrections.FILE.read_text(encoding="utf-8"))
            data = [d for d in data if (d.get("heard") or "").lower() != heard.lower()]
            corrections.FILE.write_text(json.dumps(data, indent=2, ensure_ascii=False), encoding="utf-8")
        return self.api_fixes()

    def api_style(self):
        learned = commands.style_learned()
        text = commands.STYLE_FILE.read_text(encoding="utf-8") if commands.STYLE_FILE.exists() else (
            "# My writing style\n\nNot written yet. Until it is, VoiceTool uses this description:\n\n"
            "## In my own words (always wins)\n" + commands.OWNER_SAYS)
        # "Re-learn from my sent emails" needs Claude: the page hides the button unless Claude is on, installed and signed in.
        return {"learned": learned, "exists": commands.STYLE_FILE.exists(), "text": text, "use_style": self.config.get("use_style", True),
                "claude_available": self._claude_ready()}

    def api_relearn_style(self):
        self.host.refresh_style()
        return {"started": True}

    # ---------- Named styles and per-app rules ----------

    def api_styles(self):
        recent = []
        for row in reversed(read_history()):
            app = row.get("app")
            if app and app not in recent:
                recent.append(app)
            if len(recent) >= 12:
                break
        return {**styles.load(), "recent_apps": recent}

    def api_save_styles(self, data):
        return {**styles.save(data), "recent_apps": self.api_styles()["recent_apps"]}

    # ---------- Meetings (history "meeting" rows + a Word document in the Meeting notes folder) ----------

    def api_meetings(self):
        import meetings
        items = []
        for row in reversed(read_history()):
            if row.get("meeting") is not None and row.get("notes"):
                path = Path(row["notes"])
                items.append({"title": row.get("meeting") or "Meeting", "at": row.get("at", ""),
                              "minutes": row.get("minutes"), "path": str(path), "exists": path.exists()})
        return {"status": self.host.meeting_status(), "items": items[:100], "folder": str(meetings.notes_folder()),
                "enabled": bool(self.config.get("meeting_notes", False))}

    def api_meeting_status(self):
        return self.host.meeting_status()

    def api_start_meeting(self):
        """Asks the host to start: the host shows the "everyone agreed" card first, so nothing is recorded yet."""
        self.host.start_meeting()
        return {"started": True}

    def api_stop_meeting(self):
        self.host.stop_meeting()
        return {"stopped": True}

    def _meeting_doc(self, path):
        """Only Word files in the Meeting notes folder can be read or opened from the window."""
        import meetings
        path = Path(path).resolve()
        if path.parent != meetings.notes_folder().resolve() or path.suffix.lower() != ".docx" or not path.exists():
            raise ValueError("those notes aren't in the Meeting notes folder")
        return path

    def api_open_meeting(self, path):
        self.host.open_path(self._meeting_doc(path))
        return {"opened": True}

    def api_open_meeting_folder(self):
        self.host.open_meeting_folder()
        return {"opened": True}

    def api_meeting_notes(self, path):
        """The notes as simple Markdown, read from the Word document, to show inside the window."""
        from docx import Document
        out = []
        doc = Document(str(self._meeting_doc(path)))
        for block in doc.element.body.iterchildren():
            tag = block.tag.rsplit("}", 1)[-1]
            if tag == "p":
                para = next(p for p in doc.paragraphs if p._p is block)
                style, text = para.style.name, para.text.strip()
                if not text or style == "Title":
                    continue
                if style.startswith("Heading"):
                    out.append(f"\n## {text}")
                elif "List" in style:
                    out.append(f"- {text}")
                else:
                    out.append(text)
            elif tag == "tbl":
                table = next(t for t in doc.tables if t._tbl is block)
                for row in table.rows[1:]:
                    who, what, when = (c.text.strip() for c in row.cells[:3])
                    out.append(f"- [ ] {who}: {what}" + (f" ({when})" if when else ""))
        return {"notes": "\n".join(out).strip()}

    # ---------- Voice shortcuts ----------

    def api_shortcuts(self):
        import voice_actions
        return {"shortcuts": voice_actions.load_shortcuts()}

    def api_save_shortcuts(self, shortcuts):
        import voice_actions
        return {"shortcuts": voice_actions.save_shortcuts(shortcuts)}

    # ---------- Opening things ----------

    def api_open(self, what):
        paths = {"log": HERE / "voicetool.log", "folder": HERE, "journal": HERE / "journal", "words": WORDS,
                 "fixes": corrections.FILE, "style": commands.STYLE_FILE, "readme": HERE / "README.md"}
        if what not in paths:
            raise ValueError(f"can't open {what!r}")
        if not paths[what].exists() and what in ("style", "journal", "fixes"):  # (nothing learned yet: no file to open)
            return {"opened": False}
        self.host.open_path(paths[what])
        return {"opened": True}

    def api_copy(self, text):
        self.host.copy_text(str(text))
        return {"copied": True}

    def api_quit(self):
        self.host.quit()
        return {}
