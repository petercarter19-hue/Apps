"""Meeting notes (off unless turned on): record a call (your mic + everyone else through the PC's speakers), turn it
into text as it goes, and at the end write a Word document (summary, decisions, action items, transcript). A follow-up
email draft is only made when Claude actions are on too.

- Always asks first: "Before recording, make sure everyone on the call agrees to be recorded."
- Two tracks: "Me" is the microphone; "Others" is everything the computer is playing (on the PC, WASAPI loopback of
  every output; on the Mac, ScreenCaptureKit's system audio), so the other side is caught whichever device the call
  plays through.
- The platform parts live in meetings_windows.py and meetings_mac.py (same function names); this file is shared.
- Transcribed on this PC in pieces while the meeting runs, so the notes are ready soon after it ends.
- MeetingWatch notices when a call app (Teams, Zoom, Webex, Slack, Discord, or a browser on Meet/Zoom/Teams) starts
  using the microphone, so VoiceTool can offer to take notes, and when it stops, so the notes can finish themselves.
- Audio: your own track is kept only if "Save recordings" is on; the other people's track is always deleted once
  the notes are written.
- The notes are written by the model on this computer, so a long transcript is trimmed to its most recent part (about
  12,000 tokens) to fit.
"""
import json
import re
import sys
import threading
import time
import wave
from pathlib import Path

import numpy as np
import sounddevice as sd

CONSENT_REMINDER = "Before recording, make sure everyone on the call agrees to be recorded."
RATE = 16000
PIECE_SECONDS = 25  # transcribe in pieces of about this length, cut at a quiet moment


# ---------- recording ----------

class Track:
    """One side of the call: audio waiting to be transcribed, plus a WAV copy on disk."""

    def __init__(self, name, path):
        self.name, self.path = name, path
        self.pending = []  # float32 chunks not yet transcribed
        self.pending_start = 0.0  # seconds from the meeting start where `pending` begins
        self.samples = 0
        self.lock = threading.Lock()
        self.wav = wave.open(str(path), "wb")
        self.wav.setnchannels(1)
        self.wav.setsampwidth(2)
        self.wav.setframerate(RATE)

    def add(self, audio):
        audio = np.asarray(audio, np.float32).reshape(-1)
        with self.lock:
            self.pending.append(audio)
            self.samples += len(audio)
            self.wav.writeframes((np.clip(audio, -1, 1) * 32767).astype(np.int16).tobytes())

    def take(self, final=False):
        """The next piece to transcribe (start seconds, audio), cut at a quiet moment; None if not enough yet."""
        with self.lock:
            if not self.pending:
                return None
            audio = np.concatenate(self.pending)
            if not final and len(audio) < PIECE_SECONDS * RATE:
                return None
            cut = len(audio)
            if not final:  # quietest 0.1 s in the last 5 s, so a word isn't cut in half
                step, lo = RATE // 10, len(audio) - 5 * RATE
                window = audio[lo:lo + (5 * RATE // step) * step].reshape(-1, step)
                cut = lo + int(np.argmin((window ** 2).mean(axis=1))) * step + step // 2
            start = self.pending_start
            self.pending = [audio[cut:]] if cut < len(audio) else []
            self.pending_start = start + cut / RATE
            return start, audio[:cut]

    def close(self):
        with self.lock:
            self.wav.close()


class MeetingRecorder:
    def __init__(self, folder, transcribe, mic_device=None, min_peak=0.005, log=None, capture=True):
        """capture=False (tests only) opens neither the microphone nor the call's sound: audio is put straight onto
        recorder.me / recorder.others with .add(); everything else, including stop(), works the same."""
        self.folder = Path(folder)
        self.folder.mkdir(parents=True, exist_ok=True)
        self.transcribe, self.min_peak, self.log = transcribe, min_peak, log
        self.me = Track("Me", self.folder / "me.wav")
        self.others = Track("Others", self.folder / "others.wav")
        self.lines = []  # (start seconds, speaker, text)
        self.started_at = time.time()
        self.running = True
        self.level = 0.0
        self._mic = None
        self._threads = [threading.Thread(target=self._transcriber, daemon=True)]
        if capture:
            self._mic = sd.InputStream(samplerate=RATE, channels=1, dtype="float32", device=mic_device,
                                       callback=lambda data, n, t, s: self._from_mic(data))
            self._mic.start()
            self._threads.insert(0, threading.Thread(target=self._loopback, daemon=True))
        for th in self._threads:
            th.start()

    def _from_mic(self, data):
        self.me.add(data[:, 0])
        self.level = float(np.sqrt((data ** 2).mean()))

    def _loopback(self):
        """The other people on the call: the computer's own sound (meetings_mac / meetings_windows)."""
        try:
            platform().start_system_audio(self.others.add, self.log, lambda: self.running)
        except Exception as e:
            self._say("the call's sound isn't being recorded: %s", e)

    def _transcribe_piece(self, track, piece):
        start, audio = piece
        loud = np.flatnonzero(np.abs(audio) >= self.min_peak)
        if len(audio) < RATE // 2 or not len(loud):
            return
        try:
            import corrections
            text = corrections.apply(strip_unk(self.transcribe(audio)))
        except Exception as e:
            self._say("meeting piece not transcribed: %s", e)
            return
        if text:  # stamped when the person actually started talking in this piece
            self.lines.append((start + max(0, loud[0] - RATE // 10) / RATE, track.name, text))

    def _transcriber(self):
        while self.running:
            for track in (self.me, self.others):
                piece = track.take()
                if piece:
                    self._transcribe_piece(track, piece)
            time.sleep(2)

    def stop(self):
        """Stop recording, transcribe what's left, and return the transcript lines in time order."""
        self.running = False
        try:
            if self._mic is not None:
                self._mic.stop()
                self._mic.close()
        except Exception:
            pass
        for th in self._threads:
            th.join(timeout=10)
        for track in (self.me, self.others):
            piece = track.take(final=True)
            if piece:
                self._transcribe_piece(track, piece)
            track.close()
        self.ended_at = time.time()
        return sorted(self.lines)

    def abort(self):
        """Stop recording at once and write no notes (VoiceTool is closing in the middle of a meeting)."""
        self.running = False
        try:
            if self._mic is not None:
                self._mic.stop()
                self._mic.close()
        except Exception:
            pass
        for th in self._threads:
            th.join(timeout=3)
        for track in (self.me, self.others):
            track.close()
        self.ended_at = time.time()

    def _say(self, msg, *args):
        if self.log:
            self.log.info(msg, *args)


def strip_unk(text):
    """Room noise can come out of the speech model as "<unk>" markers; they're never words (the Mac does the same)."""
    return " ".join(str(text).replace("<unk>", " ").split())


def delete_other_side_audio(base):
    """The other people's audio is deleted when the notes are written; if VoiceTool was cut off before that, this finds
    what it left behind (at the next start-up) and deletes it. Your own track is left alone. Returns what it deleted."""
    gone = []
    base = Path(base)
    if base.exists():
        for others in base.glob("*/others.wav"):
            try:
                others.unlink()
                gone.append(str(others))
            except OSError:
                pass
    return gone


def transcript_text(lines):
    return "\n".join(f"[{int(s // 60):02d}:{int(s % 60):02d}] {who}: {text}" for s, who, text in lines)


# ---------- noticing a call ----------

CALL_WINDOWS = [(re.compile(r"^Zoom Meeting|^Zoom Webinar|Zoom Workplace.*Meeting", re.I), "Zoom"),
                (re.compile(r"^Meet [-–] |Google Meet", re.I), "Google Meet"),
                (re.compile(r"Webex Meeting|Cisco Webex Meetings", re.I), "Webex"),
                (re.compile(r"\bHuddle\b", re.I), "Slack")]



def platform():
    """This computer's half of meeting notes: meetings_mac on a Mac, meetings_windows on the PC."""
    if sys.platform == "darwin":
        import meetings_mac as p
    else:
        import meetings_windows as p
    return p


def call_in_progress():
    """'Teams', 'Zoom', ... if a call is going on now, else None."""
    return platform().call_in_progress()


class MeetingWatch:
    """Checks every few seconds; calls on_start(app) when a call begins and on_end() when it has ended."""

    def __init__(self, on_start, on_end, enabled, every=4, end_after=20, autostart=True):
        self.on_start, self.on_end, self.enabled = on_start, on_end, enabled
        self.every, self.end_after = every, end_after
        self.current, self.quiet_since = None, None
        if autostart:
            threading.Thread(target=self._loop, daemon=True).start()

    def check(self):
        """One look: has a call just started or ended?"""
        app = call_in_progress() if self.enabled() else None
        if app and not self.current:
            self.current, self.quiet_since = app, None
            self.on_start(app)
        elif app:
            self.quiet_since = None
        elif self.current:
            self.quiet_since = self.quiet_since or time.time()
            if time.time() - self.quiet_since >= self.end_after:  # muting briefly doesn't end the meeting
                self.current, self.quiet_since = None, None
                self.on_end()

    def _loop(self):
        while True:
            try:
                self.check()
            except Exception:
                pass
            time.sleep(self.every)


# ---------- the notes ----------

NOTES_PROMPT = """You write meeting notes from a transcript made by a voice tool.
Lines marked "Me" come from the user's microphone; lines marked "Others" are everyone else on the call, through their
speakers (possibly several people; use names only if they are actually said). The transcript is automatic, so fix
obvious mishearings from context and ignore small talk.

Reply with JSON only:
{"title": "short meeting title (from the calendar if given, else from the content)",
 "summary": "3-6 plain sentences: what the meeting was about and where things landed",
 "decisions": ["each decision made"],
 "action_items": [{"owner": "Me / a name / Others", "task": "what", "due": "when, or empty"}],
 "open_questions": ["anything left unresolved"]__FOLLOW__}"""

FOLLOW_UP_FIELD = """,
 "follow_up_email": "a short follow-up email body from me to the attendees: thanks, the key decisions, the action
   items with owners, next step. Plain text, no subject line, no signature unless my name is said in the meeting." """

# The notes are written by a 4B model with a 16k-token window: keep the transcript to about 12k tokens (3.5 characters
# each, to be safe), the most recent part, and say so.
MAX_TRANSCRIPT_CHARS = 42000
TRIMMED_NOTE = ("(This meeting was too long for the notes model, so the earlier part of the transcript is left out. "
                "These are the most recent lines.)\n")


LOCAL_NOTES_TIMEOUT = 900  # seconds the on-device model gets to write the notes


def transcript_only_notes(problem=""):
    """The "notes" for a meeting whose summary couldn't be written: the Word document then holds the transcript only."""
    why = f" ({problem})" if problem else ""
    return {"title": "Meeting transcript",
            "summary": f"The summary and action items couldn't be written{why}, so this document has the full transcript "
                       "only. It is kept below, word for word.",
            "decisions": [], "action_items": [], "open_questions": []}


def trim_transcript(lines, limit=MAX_TRANSCRIPT_CHARS):
    """(transcript text, True if the start was left out): the most recent lines that fit in `limit` characters."""
    text = transcript_text(lines)
    if len(text) <= limit:
        return text, False
    kept, size = [], 0
    for line in reversed(text.split("\n")):
        if size + len(line) + 1 > limit and kept:
            break
        kept.append(line[:limit])
        size += len(line) + 1
    return TRIMMED_NOTE + "\n".join(reversed(kept)), True


def make_notes(client, model, lines, started_at, ended_at, calendar_hint="", style="", local=False, followup=True):
    """The notes as a dict. local=True: the on-device model writes them, so the transcript is trimmed to fit.
    followup=False leaves the follow-up email out (it is only wanted when Claude will draft it)."""
    from usage import azure
    if local:
        transcript, _ = trim_transcript(lines)
    else:
        transcript = transcript_text(lines)
    transcript = transcript or "(nothing was heard)"
    system = NOTES_PROMPT.replace("__FOLLOW__", FOLLOW_UP_FIELD.rstrip() if followup else "")
    if followup and style:
        system += f"\n\nWrite the follow-up email in the user's style:\n{style}"
    user = (f"Meeting from {time.strftime('%A %d %B %Y, %I:%M %p', time.localtime(started_at))} to "
            f"{time.strftime('%I:%M %p', time.localtime(ended_at))}.\n{calendar_hint}\n\nTranscript:\n{transcript}")
    # A long meeting takes far longer than a dictation tidy-up, and a small model on a processor is slow: give the one
    # on this computer up to 15 minutes before the transcript-only document is written instead.
    client = client.with_options(timeout=LOCAL_NOTES_TIMEOUT if local else 180)
    extra = {"max_completion_tokens": 2500} if local else {}
    reply = client.chat.completions.create(model=model, reasoning_effort="low", response_format={"type": "json_object"},
                                           messages=[{"role": "system", "content": system}, {"role": "user", "content": user}],
                                           **extra)
    azure(reply, model, "meeting notes")
    notes = json.loads(reply.choices[0].message.content)
    notes.setdefault("title", "Meeting")
    if not followup:
        notes.pop("follow_up_email", None)
    return _american(notes)


def notes_folder():
    folder = platform().notes_base() / "Meeting notes"
    folder.mkdir(parents=True, exist_ok=True)
    return folder


def save_transcript_text(lines, started_at, ended_at, folder=None):
    """Last resort when the Word document can't be made: the transcript as a plain text file (in the Meeting notes folder,
    or if that fails, next to the program). Returns its path."""
    name = f"{time.strftime('%Y-%m-%d %H%M', time.localtime(started_at))} Meeting transcript.txt"
    text = (f"Meeting transcript, {time.strftime('%A %d %B %Y, %I:%M %p', time.localtime(started_at))}\n\n"
            + transcript_text(lines) + "\n")
    here = Path(__file__).resolve().parent
    candidates = [folder] if folder else []
    try:
        candidates.append(notes_folder())
    except Exception:
        pass
    candidates.append(here)
    last = None
    for target in candidates:
        try:
            Path(target).mkdir(parents=True, exist_ok=True)
            path = Path(target) / name
            path.write_text(text, encoding="utf-8")
            return path
        except OSError as e:
            last = e
    raise last or OSError("nowhere to save the transcript")


def write_document(notes, lines, started_at, ended_at, folder=None):
    """Save the notes as a Word document in the Meeting notes folder and return its path."""
    from docx import Document
    from docx.shared import Pt
    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(11)
    title = notes.get("title") or "Meeting"
    doc.add_heading(title, level=0)
    minutes = max(1, round((ended_at - started_at) / 60))
    doc.add_paragraph(f"{time.strftime('%A %d %B %Y, %I:%M %p', time.localtime(started_at))} · {minutes} minutes · "
                      "notes by VoiceTool")
    doc.add_heading("Summary", level=1)
    doc.add_paragraph(notes.get("summary") or "")
    for heading, key in (("Decisions", "decisions"), ("Open questions", "open_questions")):
        items = notes.get(key) or []
        if items:
            doc.add_heading(heading, level=1)
            for item in items:
                doc.add_paragraph(str(item), style="List Bullet")
    actions = notes.get("action_items") or []
    if actions:
        doc.add_heading("Action items", level=1)
        table = doc.add_table(rows=1, cols=3)
        table.style = "Light List Accent 1"
        for cell, text in zip(table.rows[0].cells, ("Who", "What", "When")):
            cell.text = text
        for a in actions:
            row = table.add_row().cells
            row[0].text, row[1].text, row[2].text = str(a.get("owner", "")), str(a.get("task", "")), str(a.get("due", ""))
    if notes.get("follow_up_email"):
        doc.add_heading("Follow-up email", level=1)
        doc.add_paragraph(notes["follow_up_email"])
    doc.add_heading("Transcript", level=1)
    for s, who, text in lines:
        p = doc.add_paragraph()
        p.add_run(f"[{int(s // 60):02d}:{int(s % 60):02d}] {who}: ").bold = True
        p.add_run(text)
    safe = re.sub(r'[\\/:*?"<>|]+', "", title).strip()[:60] or "Meeting"
    path = (folder or notes_folder()) / f"{time.strftime('%Y-%m-%d %H%M', time.localtime(started_at))} {safe}.docx"
    doc.save(path)
    return path


FOLLOW_UP_INSTRUCTION = """A meeting just ended. It ran from {start} to {end}. Find it in the user's Outlook and Google
calendars (an event overlapping that time) to get its title and attendees. Then create ONE Outlook draft (a draft only;
never send) to the attendees other than the user, subject "Follow-up: <meeting title>", with exactly this body:

{body}

If no calendar event matches, address the draft to the user themselves and say so. Finish by saying who the draft is
addressed to."""


def _american(value):
    """American spelling throughout the notes."""
    from spelling import americanize
    if isinstance(value, str):
        return americanize(value)
    if isinstance(value, list):
        return [_american(v) for v in value]
    if isinstance(value, dict):
        return {k: _american(v) for k, v in value.items()}
    return value
